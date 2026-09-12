import unittest
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from cuenta_formatos import (categorize_cuenta, clasificar_esencial, concepto_cuenta, es_no_presupuestar,
                              marcar_novedad_movimientos, movimiento_cuenta, parse_detalle_transacciones_sheet,
                              periodo_colilla, plataforma_inversion, posiciones_binance_desde_movimientos,
                              posiciones_desde_operaciones, resumen_operaciones_inversion,
                              simbolo_cotizacion)


class CuentaFormatosTests(unittest.TestCase):
    def _worksheet(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        data = BytesIO()
        workbook.save(data)
        data.seek(0)
        return load_workbook(data, data_only=True).active

    def test_periodo_colilla_distinguishes_primas_and_regular_pay_periods(self):
        self.assertEqual("Prima dic-2025", periodo_colilla(
            "Primera Quincena De Diciembre De 2025\n130 PRIMA LEGAL", "dic", "2025", "1a"))
        self.assertEqual("Prima jun-2026", periodo_colilla(
            "Segunda Quincena De Junio De 2026\n130 prima legal", "jun", "2026", "2a"))
        self.assertEqual("2a quincena jun-2026", periodo_colilla(
            "Segunda Quincena De Junio De 2026\n1 SUELDO", "jun", "2026", "2a"))

    def test_periodo_colilla_normalizes_cesantias(self):
        self.assertEqual("Cesantías ene-2026", periodo_colilla(
            "Segunda Quincena De Enero De 2026\n181 CESANTÍAS AÑO ANTERIOR", "ene", "2026", "2a"))

    def test_reads_detalle_transacciones_and_uses_transaction_type_for_sign(self):
        sheet = self._worksheet([
            ["Fecha", "Tipo de transacción", "Descripción", "Valor"],
            ["1 sept 2026", "Débito", "Pago uno", 750000],
            ["31 ago 2026", "Crédito", "Pago dos", -1323004],
        ])

        result = parse_detalle_transacciones_sheet(sheet)

        self.assertEqual([
            dict(fecha=date(2026, 9, 1), descripcion="Pago uno", valor=-750000.0),
            dict(fecha=date(2026, 8, 31), descripcion="Pago dos", valor=1323004.0),
        ], result)

    def test_returns_none_for_the_legacy_printed_statement_format(self):
        sheet = self._worksheet([
            ["DESDE", "HASTA"],
            ["2026/08/01", "2026/08/31"],
            ["01/08", "PAGO AUTOM TC VISA", None, None, -100000],
        ])

        self.assertIsNone(parse_detalle_transacciones_sheet(sheet))

    def test_reads_signed_movements_export_without_transaction_type(self):
        sheet = self._worksheet([
            ["Fecha", "Descripción", "Referencia", "Valor"],
            [date(2026, 9, 4), "ABONO INTERESES AHORROS", None, 4.14],
            [date(2026, 9, 4), "PAGO QR HELADOS YUR", "0047452842", -13500],
        ])

        result = parse_detalle_transacciones_sheet(sheet)

        self.assertEqual([
            dict(fecha=date(2026, 9, 4), descripcion="ABONO INTERESES AHORROS", valor=4.14),
            dict(fecha=date(2026, 9, 4), descripcion="PAGO QR HELADOS YUR", valor=-13500.0),
        ], result)

    def test_normalizes_verbose_descriptions_for_cross_format_deduplication(self):
        self.assertEqual(
            "PAGO DE NOMI HOSPITAL HPTU",
            concepto_cuenta("PAGO TRASLADO HECHO POR HOSPITAL HPTU CON NIT 890901826 POR CONCEPTO PGO NOMIN"),
        )
        self.assertEqual(
            "TRANSFERENCIA CTA SUC VIRTUAL 29280457394",
            concepto_cuenta("Traslado de fondos por SUCURSAL VIRTUAL al producto 292-804 57394"),
        )
        self.assertEqual(
            "TRANSFERENCIA ESPOSA 23077460411",
            concepto_cuenta("Transferencia de fondos por SUCURSAL VIRTUAL del producto 2 30-77460411"),
        )
        self.assertEqual(
            "TRANSFERENCIA MAMA 10312780933",
            concepto_cuenta("Traslado de fondos por SUCURSAL VIRTUAL al producto 103-127 80933"),
        )
        self.assertEqual(
            "PAGO LILI 10072477435",
            concepto_cuenta("Traslado de fondos por SUCURSAL VIRTUAL al producto 100-724 77435"),
        )
        self.assertEqual(
            "PAGO LILI 10072477435",
            concepto_cuenta("Traslado de fondos por SUCURSAL VIRTUAL al producto 100-742 77435"),
        )

    def test_marks_duplicates_between_files_and_existing_sheet_rows(self):
        movimientos = [
            dict(fecha=date(2026, 9, 1), descripcion="Pago", concepto="PAGO", valor=-100,
                 _archivo="extracto-a.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 1), descripcion="Pago", concepto="PAGO", valor=-100,
                 _archivo="extracto-b.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 2), descripcion="Ingreso", concepto="INGRESO", valor=200,
                 _archivo="extracto-b.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 1), descripcion="Pago", concepto="PAGO", valor=-100,
                 _archivo="extracto-a.xlsx", _sheet="Hoja 1"),
        ]
        ya_ingresos = {("02/09/2026", "INGRESO", 200)}

        marcar_novedad_movimientos(movimientos, ya_ingresos, set())

        self.assertTrue(movimientos[0]["_nuevo"])
        self.assertFalse(movimientos[0]["_duplicado_subida"])
        self.assertFalse(movimientos[1]["_nuevo"])
        self.assertTrue(movimientos[1]["_duplicado_subida"])
        self.assertFalse(movimientos[2]["_nuevo"])
        self.assertFalse(movimientos[2]["_duplicado_subida"])
        self.assertTrue(movimientos[3]["_nuevo"])
        self.assertFalse(movimientos[3]["_duplicado_subida"])

    def test_es_no_presupuestar_matches_the_no_presupuestar_suffix_only(self):
        self.assertTrue(es_no_presupuestar("Ajustes y Reversiones (no presupuestar)"))
        self.assertTrue(es_no_presupuestar("Pago Tarjeta de Crédito (no presupuestar)"))
        self.assertTrue(es_no_presupuestar("Intereses y Cargos Financieros (no presupuestar)"))
        self.assertFalse(es_no_presupuestar("Honorarios / Consultoría"))
        self.assertFalse(es_no_presupuestar(""))
        self.assertFalse(es_no_presupuestar(None))

    def test_deduplication_preserves_legitimate_identical_occurrences(self):
        movimientos = [
            dict(fecha=date(2026, 9, 4), descripcion="Pago", valor=-100,
                 _archivo="nuevo.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 4), descripcion="Pago", valor=-100,
                 _archivo="nuevo.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 4), descripcion="Pago", valor=-100,
                 _archivo="solapado.xlsx", _sheet="Hoja 1"),
            dict(fecha=date(2026, 9, 4), descripcion="Pago", valor=-100,
                 _archivo="solapado.xlsx", _sheet="Hoja 1"),
        ]

        marcar_novedad_movimientos(movimientos, [], [("04/09/2026", "Pago", 100)])

        self.assertFalse(movimientos[0]["_nuevo"])
        self.assertTrue(movimientos[1]["_nuevo"])
        self.assertFalse(movimientos[2]["_nuevo"])
        self.assertFalse(movimientos[2]["_duplicado_subida"])
        self.assertTrue(movimientos[3]["_duplicado_subida"])

    def test_categorize_cuenta_is_idempotent_on_an_already_canonical_concept(self):
        # Regresión: categorize_cuenta() usaba a re-normalizar el concepto ya
        # canónico con concepto_cuenta() antes de buscarlo en la tabla, lo que
        # mandaba "PAGO DE NOMI HOSPITAL HPTU" (ya canónico) a la rama
        # "PAGO DE PROV HOSPITAL HPTU" porque no contiene el texto crudo
        # "PGO NOMIN" que dispara esa rama en concepto_cuenta().
        concepto = "PAGO DE NOMI HOSPITAL HPTU"
        primera = categorize_cuenta(concepto, es_ingreso=True)
        segunda = categorize_cuenta(concepto, es_ingreso=True)
        self.assertEqual(primera, segunda)
        categoria, _ = primera
        self.assertTrue(es_no_presupuestar(categoria))

    def test_movimiento_cuenta_marks_payroll_deposit_as_no_presupuestar(self):
        movimiento = movimiento_cuenta(
            date(2026, 9, 1),
            "PAGO TRASLADO HECHO POR HOSPITAL HPTU CON NIT 890901826 POR CONCEPTO PGO NOMIN",
            2500000,
        )
        self.assertEqual("PAGO DE NOMI HOSPITAL HPTU", movimiento["concepto"])
        self.assertTrue(movimiento["no_presupuestar"])

    def test_movimiento_cuenta_marks_real_income_as_presupuestable(self):
        movimiento = movimiento_cuenta(date(2026, 9, 1), "PAGO DE NOMINA UDEA UNIVERSIDAD DE ANTIOQUIA", 1500000)
        self.assertFalse(movimiento["no_presupuestar"])
        self.assertEqual("Honorarios / Consultoría", movimiento["categoria"])

    def test_movimiento_cuenta_marks_automatic_card_payment_as_no_presupuestar_expense(self):
        movimiento = movimiento_cuenta(date(2026, 9, 2), "PAGO AUTOM TC VISA", -350000)
        self.assertTrue(movimiento["no_presupuestar"])

    def test_movimiento_cuenta_marks_cash_withdrawal_and_other_bank_transfer_as_no_presupuestar(self):
        # Regresión: auditoría 2025 encontró estos dos sin regla, cayendo en
        # "Otros" — un retiro de efectivo y una transferencia a otro banco
        # no son gasto de consumo, solo cambian de forma la misma plata.
        retiro = movimiento_cuenta(date(2026, 9, 2), "RETIRO TARJETA EN SUCURSAL", -500000)
        self.assertTrue(retiro["no_presupuestar"])
        traslado = movimiento_cuenta(date(2026, 9, 2), "TRASLADO VIRTUAL OTROS BANCOS", -1000000)
        self.assertTrue(traslado["no_presupuestar"])

    def test_movimiento_cuenta_normalizes_atm_withdrawals_regardless_of_location_suffix(self):
        # Regresión: auditoría 2026-09 encontró ~10 retiros de cajero cayendo
        # en "Otros" porque cada cajero puntual ("ATM HOSP PABLO", "SUC
        # HOSPITAL PA", "CAJERO TESORO E"...) quedaba como concepto distinto.
        for descripcion in ("RETIRO CAJERO", "RETIRO CAJERO  ATM HOSP PABLO",
                            "RETIRO CAJERO  SUC HOSPITAL PA", "RETIRO CAJERO  CAJERO TESORO E"):
            movimiento = movimiento_cuenta(date(2026, 9, 2), descripcion, -200000)
            self.assertEqual("RETIRO CAJERO", movimiento["concepto"])
            self.assertTrue(movimiento["no_presupuestar"])

    def test_plataforma_inversion_routes_by_currency(self):
        self.assertEqual(("Trii", "pesos"), plataforma_inversion("TRANSFERENCIA A TRII SAS"))
        self.assertEqual(("Acciones y Valores", "pesos"),
                          plataforma_inversion("Traslado de fondos por SUCURSAL VIRTUAL ACCIONES Y VALORES"))
        self.assertEqual(("Binance", "dolares"), plataforma_inversion("COMPRA BINANCE COLOMBIA"))
        self.assertEqual(("Plenti", "dolares"), plataforma_inversion("TRANSFERENCIA PLENTI SAS"))
        self.assertEqual(("Hapi", "dolares"), plataforma_inversion("PAGO HAPI INVEST"))
        self.assertEqual(("Hapi", "dolares"), plataforma_inversion("PAGO PSE MONO COLOMBIA SAS"))
        self.assertEqual(("Interactive Brokers", "dolares"),
                         plataforma_inversion("PAGO PSE SOLUCIONES DE PAGOS"))
        self.assertIsNone(plataforma_inversion("MERCADO PAGO"))

    def test_posiciones_desde_operaciones_keeps_brokers_separate_and_uses_average_cost(self):
        operaciones = [
            # El PDF viene descendente y puede listar la venta antes de la
            # compra del mismo día; la posición cerrada no debe reaparecer.
            dict(date=date(2025, 12, 31), symbol="VOO", type="SELL", quantity=1, price=110,
                 commission=1, current_price=120),
            dict(date=date(2025, 12, 31), symbol="VOO", type="BUY", quantity=1, price=100,
                 commission=1, current_price=120),
            dict(date=date(2026, 1, 1), symbol="ADBE", type="BUY", quantity=2, price=100,
                 commission=2, current_price=120),
            dict(date=date(2026, 1, 2), symbol="ADBE", type="BUY", quantity=1, price=130,
                 commission=1, current_price=120),
            dict(date=date(2026, 1, 3), symbol="ADBE", type="SELL", quantity=1, price=125,
                 commission=1, current_price=120),
        ]
        posiciones = posiciones_desde_operaciones(operaciones, "IBKR")
        self.assertEqual(1, len(posiciones))
        self.assertEqual("IBKR - ADBE", posiciones[0]["ticker"])
        self.assertAlmostEqual(2, posiciones[0]["cantidad"])
        self.assertAlmostEqual(111, posiciones[0]["precio_compra"])
        self.assertEqual(120, posiciones[0]["precio_actual"])

    def test_resumen_operaciones_inversion_keeps_short_and_closed_positions(self):
        operaciones = [
            dict(date=date(2026, 1, 1), symbol="WMT", type="SHORT", quantity=3, price=120,
                 commission=3, current_price=100),
            dict(date=date(2026, 1, 2), symbol="WMT", type="COVER", quantity=1, price=110,
                 commission=1, current_price=100),
            dict(date=date(2026, 1, 3), symbol="CLOSED", type="BUY", quantity=2, price=50,
                 commission=2, current_price=60),
            dict(date=date(2026, 1, 4), symbol="CLOSED", type="SELL", quantity=2, price=60,
                 commission=2, current_price=60),
        ]

        posiciones, historial = resumen_operaciones_inversion(operaciones, "Interactive Brokers")

        self.assertEqual(1, len(posiciones))
        self.assertEqual("Interactive Brokers - WMT", posiciones[0]["ticker"])
        self.assertEqual(-2, posiciones[0]["cantidad"])
        self.assertAlmostEqual(119, posiciones[0]["precio_compra"])
        self.assertEqual("Abierta corta", next(h for h in historial if h["activo"] == "WMT")["estado"])
        self.assertEqual("Cerrada", next(h for h in historial if h["activo"] == "CLOSED")["estado"])
        self.assertAlmostEqual(16, sum(h["resultado_realizado"] for h in historial
                                       if h["activo"] == "CLOSED"))

    def test_posiciones_binance_nets_internal_transfers(self):
        movimientos = [
            {"Moneda": "BTC", "Cambiar": "0.01"},
            {"Moneda": "BTC", "Cambiar": "-0.004"},
            {"Moneda": "BTC", "Cambiar": "0.004"},
            {"Moneda": "USDT", "Cambiar": "20"},
            {"Moneda": "USDT", "Cambiar": "-20"},
            {"Moneda": "COP", "Cambiar": "500000"},
        ]

        posiciones = posiciones_binance_desde_movimientos(movimientos)

        self.assertEqual(1, len(posiciones))
        self.assertEqual("Binance - BTC", posiciones[0]["ticker"])
        self.assertAlmostEqual(0.01, posiciones[0]["cantidad"])
        self.assertEqual("Cripto", posiciones[0]["tipo"])

    def test_simbolo_cotizacion_routes_markets_without_exposing_portfolio_data(self):
        self.assertEqual("NUCO.CL", simbolo_cotizacion("Acciones y Valores - NUCO", "pesos", "Acción"))
        self.assertEqual("BTC-USD", simbolo_cotizacion("Binance - BTC", "dolares", "Cripto"))
        self.assertEqual("BRK-B", simbolo_cotizacion("IBKR - BRK-B", "dolares", "Acción"))
        self.assertIsNone(simbolo_cotizacion("IBKR - Efectivo/Margen", "dolares", "Otro"))
        self.assertIsNone(simbolo_cotizacion("Plenti", "dolares", "Fondo de Inversión"))

    def test_movimiento_cuenta_flags_investment_platform_transfers(self):
        m_pesos = movimiento_cuenta(date(2026, 9, 1), "TRANSFERENCIA A TRII SAS", -500000)
        self.assertEqual("Inversiones", m_pesos["categoria"])
        self.assertEqual(("Trii", "pesos"), m_pesos["aporte_inversion"])

        m_dolares = movimiento_cuenta(date(2026, 9, 1), "COMPRA BINANCE COLOMBIA", -300000)
        self.assertEqual(("Binance", "dolares"), m_dolares["aporte_inversion"])

        # un ingreso (plata que vuelve del broker) no es un aporte
        m_ingreso = movimiento_cuenta(date(2026, 9, 1), "TRANSFERENCIA A TRII SAS", 500000)
        self.assertIsNone(m_ingreso["aporte_inversion"])
        self.assertEqual(("Trii", "pesos"), m_ingreso["flujo_inversion"])

        # un gasto normal no dispara nada
        m_normal = movimiento_cuenta(date(2026, 9, 1), "EXITO WOW ENVIGADO", -80000)
        self.assertIsNone(m_normal["aporte_inversion"])

    def test_clasificar_esencial(self):
        self.assertEqual("Esencial", clasificar_esencial("Salud"))
        self.assertEqual("Esencial", clasificar_esencial("Mercado y Supermercado"))
        self.assertEqual("No esencial", clasificar_esencial("Entretenimiento"))
        self.assertEqual("No esencial", clasificar_esencial("Restaurantes y Domicilios"))
        self.assertEqual("No consumo (ahorro/inversión)", clasificar_esencial("Inversiones"))
        self.assertEqual("No consumo (ahorro/inversión)", clasificar_esencial("Ahorro"))
        self.assertEqual("Sin clasificar", clasificar_esencial("Una categoría inventada"))


if __name__ == "__main__":
    unittest.main()
