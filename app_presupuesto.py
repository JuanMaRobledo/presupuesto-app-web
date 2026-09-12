#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
App para el presupuesto personal. La base de datos es un Google Sheet — una
copia en vivo de Presupuesto_Juan_Robledo.xlsx, con las mismas hojas y
fórmulas — así que todo lo que subís acá queda guardado ahí mismo, para
siempre, sin tener que subir ni descargar ningún archivo.

Cuatro secciones:
  - Dashboard: vista de los datos ya cargados (los números vienen de las
    fórmulas del Sheet — la app solo los lee y los grafica).
  - Colillas de pago (PDF): sube colillas del Hospital Pablo Tobón Uribe.
  - Extractos de tarjeta (Excel): sube extractos detallados de Bancolombia.
  - Cuenta de ahorros (Excel): sube extractos o detalles de transacciones.

Ver README.md para cómo configurar la conexión al Google Sheet
(credenciales de cuenta de servicio + ID del Sheet, como "secrets" de
Streamlit).
"""
import csv
import re
import time
from datetime import date, datetime
from io import StringIO

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots
from streamlit.errors import StreamlitSecretNotFoundError

try:
    import pdfplumber
except ImportError:
    st.error("Falta instalar pdfplumber. En una terminal corre:  pip3 install pdfplumber")
    st.stop()
try:
    from openpyxl import load_workbook
except ImportError:
    st.error("Falta instalar openpyxl. En una terminal corre:  pip3 install openpyxl")
    st.stop()
try:
    import yfinance as yf
except ImportError:
    yf = None

import sheets_backend as db
from cuenta_formatos import (CATEGORIAS_OTROS_INGRESOS, PLATAFORMAS_INVERSION_DOLARES, PLATAFORMAS_INVERSION_PESOS,
                              clasificar_esencial, concepto_cuenta, es_no_presupuestar, marcar_novedad_movimientos,
                              movimiento_cuenta, parse_detalle_transacciones_sheet, periodo_colilla,
                              posiciones_binance_desde_movimientos, resumen_operaciones_inversion,
                              simbolo_cotizacion)

# Paleta de acento para todos los gráficos Plotly: arranca con el mismo azul
# de las tarjetas de stats (.stat-num / st.metric), en vez de la paleta por
# defecto de Plotly, para que los gráficos se sientan parte de la misma app.
px.defaults.color_discrete_sequence = [
    "#1d4ed8", "#d97706", "#0d9488", "#dc2626", "#7c3aed",
    "#65a30d", "#0891b2", "#be185d", "#4b5563",
]

st.set_page_config(page_title="Presupuesto", page_icon="💰", layout="wide")

st.markdown("""
<style>
:root {
  --acc-bg: #f5f7fa;
  --acc-border: #e3e7ee;
  --acc-num: #1d4ed8;
  --acc-label: #667085;
  --acc-sidebar-bg: #1a1f36;
  --acc-sidebar-text: #e2e5f1;
  --acc-sidebar-rule: #333a56;
}
@media (prefers-color-scheme: dark) {
  :root {
    --acc-bg: #1e2530;
    --acc-border: #333c4d;
    --acc-num: #7aa2f7;
    --acc-label: #9aa5b1;
    --acc-sidebar-bg: #0d1117;
    --acc-sidebar-text: #d7dde5;
    --acc-sidebar-rule: #2a3241;
  }
}
/* Tarjetas propias (.stat-card) y st.metric() nativo comparten el mismo
   estilo — antes eran dos sistemas visuales distintos para lo mismo. */
.stat-card, div[data-testid="stMetric"] {background:var(--acc-bg); border-radius:10px; padding:14px 18px; border:1px solid var(--acc-border);}
.stat-num, div[data-testid="stMetricValue"] p {font-size:1.5rem; font-weight:700; color:var(--acc-num);}
.stat-label, div[data-testid="stMetricLabel"] p {font-size:0.8rem; color:var(--acc-label); text-transform:uppercase; letter-spacing:.03em;}
/* Streamlit trunca con "..." los valores/etiquetas de st.metric() cuando
   la columna es angosta (pantallas chicas, muchas columnas por fila) —
   con cifras en pesos de 9-10 dígitos eso corta el número a la mitad. El
   elemento real es un <label> (no un <div>) con un <p> anidado dos
   niveles adentro (dentro de stMarkdownContainer), así que el selector
   tiene que valer para cualquier etiqueta, sin fijar "div", y alcanzar
   ese <p> como descendiente, no como hijo directo. */
[data-testid="stMetricValue"], [data-testid="stMetricValue"] p,
[data-testid="stMetricLabel"], [data-testid="stMetricLabel"] p,
[data-testid="stMetricDelta"], [data-testid="stMetricDelta"] p {
    white-space: normal !important; overflow: visible !important; text-overflow: clip !important;
    overflow-wrap: break-word;
}
@media (max-width: 1200px) {
    [data-testid="stMetricValue"] p {font-size: 1.2rem;}
}
.pill-ok {background:#dcfce7; color:#166534; padding:2px 10px; border-radius:999px; font-size:0.8rem;}
.pill-skip {background:#fef3c7; color:#92400e; padding:2px 10px; border-radius:999px; font-size:0.8rem;}
section[data-testid="stSidebar"] {background:var(--acc-sidebar-bg);}
section[data-testid="stSidebar"] .stRadio label {color:var(--acc-sidebar-text) !important; font-size:0.95rem;}
section[data-testid="stSidebar"] h1, section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3, section[data-testid="stSidebar"] p {color:var(--acc-sidebar-text) !important;}
section[data-testid="stSidebar"] hr {border-color:var(--acc-sidebar-rule);}
</style>
""", unsafe_allow_html=True)


def fmt_moneda(v):
    """Formatea un valor en pesos con el signo negativo antes del '$', no
    entre el '$' y el número (evita '$-27,696,152')."""
    return f"-${abs(v):,.0f}" if v < 0 else f"${v:,.0f}"


# ===========================================================================
# CONEXIÓN AL GOOGLE SHEET (la base de datos)
# ===========================================================================
try:
    sheets_configurado = "gcp_service_account" in st.secrets and "sheet_id" in st.secrets
except StreamlitSecretNotFoundError:
    sheets_configurado = False

if not sheets_configurado:
    st.title("💰 Presupuesto")
    st.error(
        "Todavía no está configurada la conexión al Google Sheet. "
        "Faltan `sheet_id` y/o `[gcp_service_account]` en los secrets de Streamlit."
    )
    st.info("Mirá la sección **'Configurar la conexión a Google Sheets'** del README.md "
            "de este repositorio para el paso a paso.")
    st.stop()

try:
    db.check_connection()
except Exception as e:
    st.title("💰 Presupuesto")
    st.error(f"No pude conectarme al Google Sheet: {e}")
    st.info("Revisá que compartiste la hoja con el correo de la cuenta de servicio "
            "(como Editor) y que `sheet_id` en los secrets es el ID correcto "
            "(la parte de la URL entre `/d/` y `/edit`).")
    st.stop()

# ===========================================================================
# LÓGICA — COLILLAS DE PAGO (PDF)
# ===========================================================================
MESES = {"Enero": "ene", "Febrero": "feb", "Marzo": "mar", "Abril": "abr", "Mayo": "may", "Junio": "jun",
         "Julio": "jul", "Agosto": "ago", "Septiembre": "sep", "Octubre": "oct", "Noviembre": "nov",
         "Diciembre": "dic"}

CONCEPTOS = {
    "1":    ("Salario Base", "Sueldo", "dia"),
    "45":   ("Recargos y Horas Extra", "Recargo Nocturno 35%", "hor"),
    "52":   ("Recargos y Horas Extra", "Recargo Nocturno Dom/Fest {pct}%", "hor"),
    "66":   ("Recargos y Horas Extra", "Dominical Trabajado {pct}%", "hor"),
    "85":   ("Recargos y Horas Extra", "Horas Adicionales {pct}%", "dia"),
    "775":  ("Formación Continua", "Formación Continua", "dia"),
    "2195": ("Ahorro", "Cuenta A.F.C. Banco Colpatria", None),
    "2355": ("Fondo de Empleados", "Aporte Social Fondo Empleados", None),
    "2360": ("Ahorro", "Ahorro Permanente Fondo Empl.", None),
    "2370": ("Seguros", "FE Salud Sura", None),
    "2450": ("Deuda (Préstamo Fondo Empleados)", "Préstamo Fondo Empleados", None),
    "2806": ("Seguros", "FE Plan C Sura", None),
    "3008": ("Impuestos", "Retención en la Fuente (Método 2)", None),
    "3010": ("Aportes de Ley", "Aporte Salud (EPS Sura)", None),
    "3020": ("Aportes de Ley", "Aporte Pensión (Protección)", None),
    "3023": ("Aportes de Ley", "Aporte Fondo de Solidaridad (Protección)", None),
    "6521": ("Seguros", "Póliza Automóvil - Fondo Empl.", None),
    "6524": ("Transporte", "Parqueadero Carro", None),
    # Cesantías: cada enero, el Hospital emite además de la colilla quincenal
    # normal un comprobante de liquidación de cesantías del año anterior —
    # mismo formato de "colilla" (misma fecha, "Totales:", "Neto a Pagar:"),
    # así que sin distinguirlo el parser lo toma como si fuera otra quincena
    # más (ver detección por "CESANTIAS" en el texto, en parse_colilla). NO
    # es sueldo: el devengo de cesantías (181) se consigna íntegro al fondo
    # el mismo día (descuento 3180, mismo valor) y no queda como plata
    # disponible — aplicar_colillas() no guarda esos dos en Colillas de Pago
    # ni en ningún otro lado, quedan fuera del presupuesto. Los intereses
    # (191) sí se pagan — son ingreso real y aplicar_colillas() los manda
    # solos a Otros Ingresos, categoría "Cesantías".
    "181":  ("Cesantías (no presupuestar)", "Cesantías Año Anterior", None),
    "191":  ("Cesantías (Ingreso)", "Intereses de Cesantías Año Anterior", None),
    "3180": ("Cesantías (no presupuestar)", "Consignación Cesantías a Fondo", None),
    # Estos sí son ingreso real (se pagan efectivamente, sin una consignación
    # que los anule como pasa con cesantías) — solo les faltaba categoría.
    "130":  ("Prima de Servicios", "Prima Legal", None),
    "145":  ("Vacaciones y Licencias", "Vacaciones", None),
    "951":  ("Vacaciones y Licencias", "Ajuste de Vacaciones en Tiempo", None),
    "15":   ("Vacaciones y Licencias", "Incapacidad", None),
    "20":   ("Vacaciones y Licencias", "Gasto de Incapacidad", None),
    "736":  ("Vacaciones y Licencias", "Tiempo Para Ti", None),
    "26":   ("Vacaciones y Licencias", "Licencia de Paternidad", None),
    # Descuentos reales (ya se venían categorizando a mano en el Sheet).
    "3510": ("Seguros", "Servicios Hospitalarios 2", None),
    "2410": ("Ahorro", "Cuota Voluntaria Fondo Empleados", None),
    "2905": ("Ahorro", "Ahorro Navideño", None),
}

PALABRAS_CLAVE_DEVENGO = [
    (("PRIMA",), "Prima de Servicios", "Prima de Servicios"),
    (("BONO", "BONIFICAC"), "Bonificación", "Bonificación"),
]


def money_colilla(s):
    return float(s.replace(",", ""))


def parse_colilla(fileobj, nombre):
    with pdfplumber.open(fileobj) as pdf:
        if len(pdf.pages) != 1:
            st.warning(f"Aviso: {nombre} tiene {len(pdf.pages)} páginas; solo se lee la primera.")
        page = pdf.pages[0]
        text = page.extract_text() or ""
        words = page.extract_words(use_text_flow=False, keep_blank_chars=False)

    m = re.search(r"Fecha:\s*(\d{4}-\d{2}-\d{2})", text)
    if not m:
        raise ValueError(f"No encontré la fecha del comprobante en {nombre}. "
                          "¿Es una colilla del Hospital Pablo Tobón Uribe en este formato?")
    fecha_pago = datetime.strptime(m.group(1), "%Y-%m-%d").date()

    m = re.search(r"(Primera|Segunda) Quincena De (\w+) De (\d{4})", text)
    if not m:
        raise ValueError(f"No encontré 'Primera/Segunda Quincena De ... De ...' en {nombre}.")
    quincena_num = "1a" if m.group(1) == "Primera" else "2a"
    mes_nombre = m.group(2)
    anio = m.group(3)
    if mes_nombre not in MESES:
        raise ValueError(f"Mes desconocido '{mes_nombre}' en {nombre}.")
    # El comprobante de Prima (junio/diciembre) y el de liquidación de
    # Cesantías (enero) usan el mismo encabezado "Primera/Segunda Quincena
    # De..." que una quincena normal — el Hospital reusa la plantilla — así
    # que el título solo no alcanza para distinguirlos. Se detectan por su
    # devengo característico ("130 PRIMA LEGAL" / cualquier línea con
    # "CESANTIAS", ver CONCEPTOS 181/191/3180): si aparece, no es una
    # quincena más, y hay que etiquetarlo aparte para que no choque con la
    # quincena real de ese mismo medio mes (y quede sin cargar, tomada como
    # duplicado — pasó con la Prima de diciembre-2025 y con la liquidación
    # de cesantías de enero-2026).
    periodo = periodo_colilla(text, MESES[mes_nombre], anio, quincena_num)

    m = re.search(r"Totales:\s*\$\s*([\d,]+\.\d{2})\s*\$\s*([\d,]+\.\d{2})", text)
    if not m:
        raise ValueError(f"No encontré la línea de 'Totales:' en {nombre}.")
    total_devengos = money_colilla(m.group(1))
    total_descuentos = money_colilla(m.group(2))

    m = re.search(r"Neto a Pagar:\s*\$\s*([\d,]+\.\d{2})", text)
    if not m:
        raise ValueError(f"No encontré 'Neto a Pagar:' en {nombre}.")
    neto = money_colilla(m.group(1))

    rows = {}
    for w in words:
        key = round(w["top"])
        rows.setdefault(key, []).append(w)

    devengos, descuentos, revisar = [], [], []
    for top in sorted(rows):
        ws_row = sorted(rows[top], key=lambda w: w["x0"])
        if not ws_row:
            continue
        first = ws_row[0]
        if not re.match(r"^\d+$", first["text"]):
            continue
        code = first["text"]
        raw_concept = " ".join(w["text"] for w in ws_row[1:] if w["x0"] < 220)
        unidad_words = [w for w in ws_row if 220 <= w["x0"] < 270]
        unidades = None
        if unidad_words:
            try:
                unidades = float(unidad_words[0]["text"])
            except ValueError:
                unidades = None
        money_words = [w for w in ws_row if w["x0"] >= 300 and re.match(r"^[\d,]+\.\d{2}$", w["text"])]
        if not money_words:
            continue
        mw = money_words[0]
        valor = money_colilla(mw["text"])
        columna = "devengo" if mw["x0"] < 400 else ("descuento" if mw["x0"] < 500 else "saldo_prestamo")
        if columna == "saldo_prestamo":
            continue

        if code in CONCEPTOS:
            categoria, nombre_base, unidad_tipo = CONCEPTOS[code]
            pct_match = re.search(r"\((\d+)%", raw_concept)
            pct = pct_match.group(1) if pct_match else ""
            nombre_item = nombre_base.format(pct=pct) if "{pct}" in nombre_base else nombre_base
            if unidad_tipo == "hor" and unidades is not None:
                nombre_item += f" ({unidades:g} h)"
            elif unidad_tipo == "dia" and unidades is not None:
                nombre_item += f" ({unidades:g} días)"
        else:
            upper_concept = raw_concept.upper()
            match_kw = next((cat_nombre for kws, categoria_kw, cat_nombre in PALABRAS_CLAVE_DEVENGO
                              if any(kw in upper_concept for kw in kws)), None)
            categoria_kw = next((categoria_kw for kws, categoria_kw, _ in PALABRAS_CLAVE_DEVENGO
                                  if any(kw in upper_concept for kw in kws)), None)
            if match_kw and columna == "devengo":
                categoria = categoria_kw
                nombre_item = match_kw
            else:
                categoria = ""
                nombre_item = f"[REVISAR] {raw_concept} (código {code})"
                revisar.append((code, raw_concept, valor))

        item = dict(quincena=periodo, concepto=nombre_item, categoria=categoria, valor=valor, codigo=code)
        (devengos if columna == "devengo" else descuentos).append(item)

    return dict(fecha=fecha_pago, periodo=periodo, total_devengos=total_devengos,
                total_descuentos=total_descuentos, neto=neto, devengos=devengos,
                descuentos=descuentos, revisar=revisar)


# ===========================================================================
# LÓGICA — EXTRACTOS DE TARJETA (Bancolombia .xlsx)
# ===========================================================================
VISA = dict(label="Visa ****7497", sheet="Egresos - Tarjeta Visa 7497")
MASTERCARD = dict(label="Mastercard ****5922", sheet="Egresos - Mastercard 5922")
# Los últimos 4 dígitos cambian cada vez que el banco reemite la tarjeta física
# (se mantiene la misma cuenta/cupo). Se registran acá los números vistos hasta
# ahora para reconocerlos directo; si aparece un número nuevo que no está en
# esta lista, se resuelve por el nombre del archivo (ver resolve_tarjeta) en
# vez de fallar — así una tarjeta reemitida no bloquea la carga.
TARJETAS = {
    "7497": VISA, "5003": VISA,
    "5922": MASTERCARD, "2223": MASTERCARD,
}


def resolve_tarjeta(digitos, nombre_archivo=""):
    if digitos in TARJETAS:
        return TARJETAS[digitos]
    low = nombre_archivo.lower()
    if "visa" in low:
        return VISA
    if "master" in low:
        return MASTERCARD
    return None
MESES_ABR = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7,
             "ago": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12}
_MESES_3LETRAS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

# Categorías reales de devengo/descuento que puede producir parse_colilla()
# (ver CONCEPTOS más abajo) — más completas que el desplegable de la hoja
# 'Colillas de Pago', que quedó corto (armado antes de que existieran
# Prima de Servicios/Vacaciones y Licencias/Bonificación/Cesantías).
DEVENGOS_CATEGORIAS_MANUAL = ["Salario Base", "Recargos y Horas Extra", "Formación Continua",
                               "Prima de Servicios", "Vacaciones y Licencias", "Bonificación",
                               "Cesantías (no presupuestar)", "Cesantías (Ingreso)"]
DESCUENTOS_CATEGORIAS_MANUAL = ["Ahorro", "Fondo de Empleados", "Seguros", "Deuda (Préstamo Fondo Empleados)",
                                 "Deuda (Leasing Habitacional)", "Impuestos", "Aportes de Ley", "Transporte",
                                 "Cesantías (no presupuestar)"]

MERCHANT_CATEGORY = {
    "INTERESES CORRIENTES": ("Intereses y Cargos Financieros (no presupuestar)", "No", ""),
    "ABONO DEBITO AUTOMATICO": ("Pago Tarjeta de Crédito (no presupuestar)", "No", "Pago automático de la tarjeta, no es gasto nuevo"),
    "REVERSION DE ABONO": ("Ajustes y Reversiones (no presupuestar)", "No", "Reversión de un abono anterior"),
    "DLO*GOOGLE SKETCHBOOK": ("Tecnología y Suscripciones", "No", ""),
    "SOC DE MEJORAS PUBLICA": ("Vivienda y Servicios", "No", ""),
    "MERCADOPAGO COLOMBIA L": ("Compras Online / Varios", "No", "Verificar qué se compró"),
    "EDS ZULY": ("Transporte", "No", "Combustible"),
    "RAPPI COLOMBIA*DL": ("Restaurantes y Domicilios", "No", ""),
    "NEWREST HOSPITAL UNIVE": ("Restaurantes y Domicilios", "No", "Cafetería del hospital"),
    "COMFAMA 1190 CIS CITY": ("Entretenimiento", "No", "Verificar servicio específico de Comfama"),
    "AIRBNB * HMF5Z5TMPJ": ("Viajes", "No", ""),
    "PARMESSANO REST DELICA": ("Restaurantes y Domicilios", "No", ""),
    "TERPEL MAYORAL": ("Transporte", "No", "Combustible"),
    "aliexpress": ("Compras Online / Varios", "No", ""),
    "LAVAPRES CAMPESTRE DRI": ("Transporte", "No", "Lavado de carro"),
    "CREPES Y WAFFLES CAMPE": ("Restaurantes y Domicilios", "No", ""),
    "EXITO WOW ENVIGADO": ("Mercado y Supermercado", "No", ""),
    "DROGUERIA EX ENVIGADO": ("Salud", "No", ""),
    "MGP*YR Bookingcom": ("Viajes", "No", ""),
    "PAYU PEEWAH": ("Otros", "No", "Verificar comercio"),
    "HOTEL BOGOTA PLAZA": ("Viajes", "No", ""),
    "AIRE DE ROMERO": ("Otros", "No", "Verificar comercio"),
    "CTRO CLINICO Y D INV S": ("Salud", "No", ""),
    "DLO*GOOGLE GOOGLE ONE": ("Tecnología y Suscripciones", "No", ""),
    "MP*TIENDADELM*TIENDADE": ("Compras Online / Varios", "No", ""),
    "PRICESMART": ("Mercado y Supermercado", "No", ""),
    "TBL* ARENA ALFA EDUCAC": ("Educación y Profesional", "No", ""),
    "PAGO ELECTRONICO FLYPASS": ("Transporte", "No", "Peajes"),
    "HOSP PABLO TOBON U": ("Restaurantes y Domicilios", "No", "Posible cafetería/tienda del hospital, verificar"),
    "EURO MURANO": ("Restaurantes y Domicilios", "No", ""),
    "SUMUP*TUVET CLINICA VE": ("Mascotas", "No", "Veterinaria"),
    "MERCADO PAGO*MELIMAS": ("Compras Online / Varios", "No", "Mercado Libre"),
    "MOVISTAR PAGOSEPAYCO": ("Tecnología y Suscripciones", "No", "Plan celular"),
    "AVIANCA SACQR59H": ("Viajes", "No", "Tiquete aéreo"),
    "MULTIPLEX VIVA ENVIGAD": ("Entretenimiento", "No", "Cine"),
    "HOMECENTER VTAS A DIST": ("Vivienda y Servicios", "No", ""),
    "MERCADO PAGO": ("Compras Online / Varios", "No", "Verificar qué se compró"),
    "OLIVENZA COCINA MEDITE": ("Restaurantes y Domicilios", "No", ""),
    # Agregados en auditoría 2026-09: comercios que aparecían en "Otros"
    # con montos altos/repetidos — el match es por texto EXACTO tal como
    # lo trunca cada banco, así que una variante de corte distinta del
    # mismo comercio (p. ej. "PRICESMART" vs "PRICESMART AMERICAS") no
    # matchea sola, hay que agregar cada una.
    "PRICESMART AMERICAS": ("Mercado y Supermercado", "No", ""),
    "EDS LA MONTANA": ("Transporte", "No", "Combustible"),
    "EDS TEXACO PUNTO CERO": ("Transporte", "No", "Combustible"),
    "AUTOAMERICA INDUSTRIAL": ("Transporte", "No", "Repuestos/accesorios de carro"),
    "CIA SURAMERICANA DE SE": ("Seguros", "No", ""),
    "SURAMERICANA SEGUROS D": ("Seguros", "No", ""),
    "CREPES Y WAFFLES TESOR": ("Restaurantes y Domicilios", "No", ""),
    "ITADAKI RAMEN SAS": ("Restaurantes y Domicilios", "No", ""),
    "AVIANCA": ("Viajes", "No", "Tiquete aéreo"),
    "MERCPAGO*JETSMART": ("Viajes", "No", "Tiquete aéreo"),
    "SNCF-VOYAGEURS": ("Viajes", "No", "Tren (Francia)"),
    "RENFE VIRTUAL INTERNET": ("Viajes", "No", "Tren (España)"),
    "LA TIQUETERA": ("Viajes", "No", "Agencia de viajes/tiquetes"),
    "AIRBNB * HM5A2S3AXE": ("Viajes", "No", ""),
    "AIRBNB * HMS28NYNJY": ("Viajes", "No", ""),
    "MERCPAGO*MERCADOLIBRE": ("Compras Online / Varios", "No", ""),
    "WOMPI*ACEM": ("Educación y Profesional", "No", "Posible membresía/asociación médica, verificar"),
    "GOOGLE *PLAY YOUTUBE*D": ("Tecnología y Suscripciones", "No", "Suscripción YouTube/Google"),
    "DONATELLA TRATTORIA": ("Restaurantes y Domicilios", "No", ""),
    "KAMIL HOSPITAL": ("Restaurantes y Domicilios", "No", "Posible cafetería/tienda del hospital, verificar"),
    "NOVAVENTA": ("Compras Online / Varios", "No", ""),
    "CREDIMAPFRE": ("Seguros", "No", "Seguro del carro"),
    "OMA BARRA MED HSPTAL P": ("Restaurantes y Domicilios", "No", "Café OMA"),
    "AVIANCA SACD3875": ("Viajes", "No", "Tiquete aéreo"),
    "VIN Y GRETTA PHTU": ("Restaurantes y Domicilios", "No", ""),
    "SUBWAY HOSPITAL PABLO": ("Restaurantes y Domicilios", "No", ""),
    "UBER RIDES*DL": ("Transporte", "No", ""),
    "MONTOLIVO PASTA EXPRES": ("Restaurantes y Domicilios", "No", ""),
    "ANTONIOS GELATO SAS": ("Restaurantes y Domicilios", "No", ""),
}


def categorize(merchant):
    return MERCHANT_CATEGORY.get(merchant, ("Otros", "No", "Comercio no reconocido, clasificar manualmente"))


def money_extracto(s):
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    s = str(s).strip()
    if s == "":
        return None
    neg = s.startswith("-") or s.endswith("-")  # el formato nuevo de Bancolombia pone el '-' al final
    s = s.strip("-").lstrip("+")
    lc, lp = s.rfind(","), s.rfind(".")
    if lc == -1 and lp == -1:
        val = float(s)
    elif lc == -1:
        val = float(s.replace(".", "")) if s.count(".") > 1 or len(s.split(".")[-1]) == 3 else float(s)
    elif lp == -1:
        val = float(s.replace(",", "")) if s.count(",") > 1 or len(s.split(",")[-1]) == 3 else float(s.replace(",", "."))
    elif lc > lp:
        val = float(s.replace(".", "").replace(",", "."))
    else:
        val = float(s.replace(",", ""))
    return -val if neg else val


def parse_date_ddmmyyyy(s):
    d, m, y = s.split("/")
    return date(int(y), int(m), int(d))


def find_label_value(rows, label, value_col_offset=1):
    for r in rows:
        for i, cell in enumerate(r):
            if isinstance(cell, str) and cell.strip().rstrip(":").lower() == label.lower():
                if i + value_col_offset < len(r):
                    return r[i + value_col_offset]
    return None


def parse_periodo_fecha(s):
    m = re.match(r"(\d{1,2})\s+([A-Za-zñÑ]+)\.?\s+(\d{4})", s.strip())
    if not m:
        return None
    dia, mes_txt, anio = m.groups()
    mes = MESES_ABR.get(mes_txt.lower().rstrip("."))
    if not mes:
        return None
    return date(int(anio), mes, int(dia))


def parse_limite_fecha(s):
    m = re.match(r"([A-Za-zñÑ]+)\.?\s+(\d{1,2}),\s+(\d{4})", s.strip())
    if not m:
        return None
    mes_txt, dia, anio = m.groups()
    mes = MESES_ABR.get(mes_txt.lower().rstrip("."))
    if not mes:
        return None
    return date(int(anio), mes, int(dia))


def parse_extracto_sheet(ws, nombre_archivo=""):
    """Bancolombia cambió el formato del extracto detallado en algún momento
    de 2025 — hay archivos viejos y nuevos circulando. Se detecta cuál es y
    se despacha al parser correspondiente."""
    rows = list(ws.iter_rows(values_only=True))
    header_rows = rows[:29]
    if find_label_value(header_rows, "Información de la Tarjeta"):
        return parse_extracto_sheet_v1(ws, rows, header_rows, nombre_archivo)
    return parse_extracto_sheet_v2(ws, rows, nombre_archivo)


def parse_extracto_sheet_v1(ws, rows, header_rows, nombre_archivo=""):
    """Formato original: bloque 'Información de la Tarjeta' con
    etiqueta:valor en la misma fila, un solo listado continuo de
    movimientos."""
    tarjeta_cell = find_label_value(header_rows, "Información de la Tarjeta")
    if not tarjeta_cell:
        raise ValueError("No encontré 'Información de la Tarjeta' — ¿es este el formato correcto?")
    m = re.search(r"(\d{4})$", str(tarjeta_cell))
    if not m:
        raise ValueError(f"No pude leer los últimos 4 dígitos de la tarjeta en '{tarjeta_cell}'.")
    digitos = m.group(1)
    tarjeta = resolve_tarjeta(digitos, nombre_archivo)
    if not tarjeta:
        raise ValueError(f"No reconozco la tarjeta terminada en {digitos} y el nombre del archivo no dice "
                          "'Visa' ni 'Mastercard'. Avísame para agregarla.")

    periodo_fin_txt = None
    for r in header_rows:
        if r[0] and str(r[0]).strip().rstrip(":").lower() == "periodo facturado":
            periodo_fin_txt = r[2]
            break
    if not periodo_fin_txt:
        raise ValueError("No encontré 'Periodo facturado'.")
    fecha_corte = parse_periodo_fecha(str(periodo_fin_txt))
    if not fecha_corte:
        raise ValueError(f"No pude leer la fecha de cierre de '{periodo_fin_txt}'.")
    periodo_extracto = f"{fecha_corte.year:04d}-{fecha_corte.month:02d}"

    limite_txt = find_label_value(header_rows, "Pagar antes de")
    fecha_limite = parse_limite_fecha(str(limite_txt)) if limite_txt else None

    moneda_txt = find_label_value(header_rows, "Moneda")
    moneda = "USD" if moneda_txt and "DOLAR" in str(moneda_txt).upper() else "COP"

    stmt = dict(
        periodo=periodo_extracto,
        fecha_corte=fecha_corte,
        fecha_limite=fecha_limite,
        cupo_total=money_extracto(find_label_value(header_rows, "Cupo total")),
        cupo_disponible=money_extracto(find_label_value(header_rows, "Cupo disponible")),
        saldo_anterior=money_extracto(find_label_value(header_rows, "+ Saldo anterior")),
        pago_minimo=money_extracto(find_label_value(header_rows, "Pago mínimo")),
        pago_total=money_extracto(find_label_value(header_rows, "Pago total")),
        notas="",
    )

    txns = []
    i = 0
    while i < len(rows):
        r = rows[i]
        if r and r[0] in ("Movimientos durante el periodo", "Movimientos antes del periodo"):
            i += 2
            continue
        if len(r) > 1 and r[1] and isinstance(r[1], str) and re.match(r"\d{2}/\d{2}/\d{4}", r[1]):
            merchant = r[2]
            valor_total = money_extracto(r[3])
            cuotas = r[4] if r[4] else "1/1"
            valor_periodo = money_extracto(r[5])
            saldo_pend = money_extracto(r[8]) if len(r) > 8 else None
            cat, reemb, nota = categorize(merchant)
            txns.append(dict(
                periodo_extracto=periodo_extracto, tarjeta=tarjeta["label"], sheet=tarjeta["sheet"],
                moneda=moneda, fecha_compra=parse_date_ddmmyyyy(r[1]), comercio=merchant, cuotas=cuotas,
                valor_total=valor_total, valor_periodo=valor_periodo, saldo_pendiente=saldo_pend,
                categoria=cat, reembolsable=reemb, nota=nota,
            ))
        i += 1

    return dict(tarjeta=tarjeta, moneda=moneda, statement=stmt, txns=txns)


def _find_valor_bajo(rows, label, limite=40):
    """Para el formato nuevo: busca 'label' como encabezado de una
    mini-tabla (fila de encabezados, fila de valores justo debajo, misma
    columna) y devuelve (valor, índice de fila del encabezado)."""
    label_norm = label.strip().lower()
    for i in range(min(limite, len(rows))):
        r = rows[i]
        if not r:
            continue
        for j, cell in enumerate(r):
            if isinstance(cell, str) and cell.strip().lower() == label_norm:
                if i + 1 < len(rows) and rows[i + 1] and j < len(rows[i + 1]):
                    return rows[i + 1][j], i
    return None, None


def parse_extracto_sheet_v2(ws, rows, nombre_archivo=""):
    """Formato nuevo de Bancolombia (visto desde mediados de 2025 en
    adelante): trae 'Información Cliente:' / 'Resumen de la Tarjeta' en vez
    de 'Información de la Tarjeta', con tablas encabezado-arriba/valor-abajo
    en vez de etiqueta:valor en la misma fila. Repite el encabezado por
    página impresa (como el extracto de cuenta de ahorros), así que en vez
    de delimitar un bloque de movimientos se escanea toda la hoja."""
    header_rows = rows[:29]

    tarjeta_cell, _ = _find_valor_bajo(header_rows, "Tarjeta")
    if not tarjeta_cell:
        raise ValueError("No encontré la columna 'Tarjeta' — ¿es este el formato correcto?")
    m = re.search(r"(\d{4})$", str(tarjeta_cell))
    if not m:
        raise ValueError(f"No pude leer los últimos 4 dígitos de la tarjeta en '{tarjeta_cell}'.")
    digitos = m.group(1)
    tarjeta = resolve_tarjeta(digitos, nombre_archivo)
    if not tarjeta:
        raise ValueError(f"No reconozco la tarjeta terminada en {digitos} y el nombre del archivo no dice "
                          "'Visa' ni 'Mastercard'. Avísame para agregarla.")

    cupo_total, fila_cupo = _find_valor_bajo(header_rows, "Cupo Total")
    cupo_disponible, fila_disp = _find_valor_bajo(header_rows, "Disponible Total")

    fecha_corte = fecha_limite = None
    if fila_cupo is not None and fila_cupo + 1 < len(rows):
        datos = rows[fila_cupo + 1]
        if len(datos) > 3 and datos[3]:
            fecha_corte = datetime.strptime(str(datos[3]).strip(), "%d/%m/%Y").date()
    if fila_disp is not None and fila_disp + 1 < len(rows):
        datos = rows[fila_disp + 1]
        if len(datos) > 2 and datos[2]:
            fecha_limite = datetime.strptime(str(datos[2]).strip(), "%d/%m/%Y").date()
    if not fecha_corte:
        raise ValueError("No pude leer 'Período Facturado Hasta'.")
    periodo_extracto = f"{fecha_corte.year:04d}-{fecha_corte.month:02d}"

    moneda = "USD" if ws.title.strip().upper() == "DOLARES" else "COP"

    stmt = dict(
        periodo=periodo_extracto,
        fecha_corte=fecha_corte,
        fecha_limite=fecha_limite,
        cupo_total=money_extracto(cupo_total),
        cupo_disponible=money_extracto(cupo_disponible),
        saldo_anterior=money_extracto(find_label_value(header_rows, "Saldo Anterior")),
        pago_minimo=money_extracto(find_label_value(header_rows, "= Pago mínimo")),
        pago_total=money_extracto(find_label_value(header_rows, "= Pagos total")),
        notas="",
    )

    txns = []
    for r in rows:
        if not r or len(r) < 3 or not r[1] or not isinstance(r[1], str):
            continue
        if not re.match(r"^\d{2}/\d{2}/\d{4}$", r[1].strip()):
            continue
        merchant = r[2]
        if not merchant:
            continue
        valor_total = money_extracto(r[3]) if len(r) > 3 else None
        valor_periodo = money_extracto(r[6]) if len(r) > 6 and r[6] not in (None, "") else valor_total
        saldo_pend = money_extracto(r[7]) if len(r) > 7 else None
        cuotas = r[8] if len(r) > 8 and r[8] else "1/1"
        cat, reemb, nota = categorize(merchant)
        txns.append(dict(
            periodo_extracto=periodo_extracto, tarjeta=tarjeta["label"], sheet=tarjeta["sheet"],
            moneda=moneda, fecha_compra=parse_date_ddmmyyyy(r[1].strip()), comercio=merchant, cuotas=cuotas,
            valor_total=valor_total, valor_periodo=valor_periodo, saldo_pendiente=saldo_pend,
            categoria=cat, reembolsable=reemb, nota=nota,
        ))

    return dict(tarjeta=tarjeta, moneda=moneda, statement=stmt, txns=txns)


# ===========================================================================
# LÓGICA — EXTRACTO DE CUENTA DE AHORROS (Bancolombia .xlsx)
# ===========================================================================
# La categorización de movimientos de cuenta (CUENTA_CATEGORY_INGRESO/EGRESO,
# categorize_cuenta, movimiento_cuenta, es_no_presupuestar) vive en
# cuenta_formatos.py — es lógica pura sin dependencias de Streamlit ni de
# openpyxl, así que puede tener tests unitarios directos (ver
# tests/test_cuenta_formatos.py). Acá solo queda el parseo específico de los
# dos formatos de extracto .xlsx.
def parse_extracto_cuenta_sheet(ws):
    """Lee el extracto impreso o la exportación Detalle de transacciones."""
    detalle = parse_detalle_transacciones_sheet(ws)
    if detalle is not None:
        return [movimiento_cuenta(**movimiento) for movimiento in detalle]

    rows = list(ws.iter_rows(values_only=True))

    desde = None
    for i, r in enumerate(rows):
        if r and r[0] == "DESDE" and len(r) > 1 and r[1] == "HASTA":
            val = rows[i + 1][0] if i + 1 < len(rows) else None
            if val:
                desde = datetime.strptime(str(val).strip(), "%Y/%m/%d").date()
            break
    if not desde:
        raise ValueError("No encontré 'DESDE/HASTA' — ¿es este el formato de extracto de cuenta de Bancolombia?")

    current_year = desde.year
    prev_month = desde.month
    txns = []
    for r in rows:
        if not r or not r[0] or not isinstance(r[0], str):
            continue
        m = re.match(r"^(\d{1,2})/(\d{2})$", r[0].strip())
        if not m:
            continue
        dia, mes = int(m.group(1)), int(m.group(2))
        if mes < prev_month:
            current_year += 1
        prev_month = mes
        try:
            fecha = date(current_year, mes, dia)
        except ValueError:
            continue
        descripcion = str(r[1]).strip() if len(r) > 1 and r[1] else ""
        if not descripcion:
            continue
        valor = money_extracto(r[4]) if len(r) > 4 and r[4] not in (None, "") else None
        if valor is None:
            continue
        txns.append(movimiento_cuenta(fecha, descripcion, valor))
    return txns


# ===========================================================================
# ESCRITURA EN EL GOOGLE SHEET
# ===========================================================================
def aplicar_colillas(nuevas):
    dev_row = db.first_blank_row("colillas_devengos")
    ded_row = db.first_blank_row("colillas_descuentos")
    for r in nuevas:
        if r["periodo"].startswith("Cesantías"):
            # Liquidación anual de cesantías del fondo, no una quincena (ver
            # comentario junto a CONCEPTOS 181/191/3180): el capital se
            # consigna al fondo el mismo día y no tiene ningún efecto en el
            # flujo de caja, así que no se guarda ni en devengos ni en
            # descuentos. Se deja igual una fila en el resumen (en $0) solo
            # para que "ya_cargadas" (más arriba, en la subida de PDFs)
            # reconozca este período si se sube por error una segunda vez.
            rr = db.first_blank_row("colillas_resumen")
            db.write_block_rows("colillas_resumen", rr, [[r["fecha"].isoformat(), r["periodo"], 0, 0]])
            interes = sum(it["valor"] for it in r["devengos"] if it.get("codigo") == "191")
            if interes:
                fila_oi = db.first_blank_row("otros_ingresos")
                db.write_block_rows("otros_ingresos", fila_oi,
                                     [[r["fecha"].isoformat(), "Intereses de Cesantías", "Cesantías", interes,
                                       f"Liquidación de cesantías {r['periodo']} — no incluye el capital, que "
                                       "se consigna al fondo el mismo día y no es ingreso"]])
            continue

        rr = db.first_blank_row("colillas_resumen")
        # El "Totales:" impreso en el comprobante puede incluir ítems
        # "(no presupuestar)" — p. ej. una liquidación de cesantías, que
        # llega en enero con el mismo formato que una colilla quincenal
        # normal (misma "Fecha:"/"Totales:"/"Neto a Pagar:"), así que el
        # parser la toma como si fuera otra quincena más (ver CONCEPTOS,
        # códigos 181/191/3180). No es sueldo: la cesantías se consigna al
        # fondo el mismo día por el mismo valor. Se restan acá antes de
        # guardar para que "Ingresos Brutos Totales" no la cuente.
        total_devengos = r["total_devengos"] - sum(
            it["valor"] for it in r["devengos"] if es_no_presupuestar(it["categoria"]))
        total_descuentos = r["total_descuentos"] - sum(
            it["valor"] for it in r["descuentos"] if es_no_presupuestar(it["categoria"]))
        db.write_block_rows("colillas_resumen", rr,
                             [[r["fecha"].isoformat(), r["periodo"], total_devengos, total_descuentos]])

        if r["devengos"]:
            rows = [[it["quincena"], it["concepto"], it["categoria"], it["valor"]] for it in r["devengos"]]
            db.write_block_rows("colillas_devengos", dev_row, rows)
            dev_row += len(rows)

        if r["descuentos"]:
            rows = [[it["quincena"], it["concepto"], it["categoria"], it["valor"]] for it in r["descuentos"]]
            db.write_block_rows("colillas_descuentos", ded_row, rows)
            ded_row += len(rows)
    db.clear_read_cache()


def _fila_detalle_tarjeta(t):
    saldo_pendiente = t["saldo_pendiente"] if t["saldo_pendiente"] is not None else ""
    return [db.as_text(t["periodo_extracto"]), t["fecha_compra"].isoformat(), t["comercio"], t["moneda"],
            db.as_text(t["cuotas"]), t["valor_total"], t["valor_periodo"], saldo_pendiente,
            t["categoria"], t["reembolsable"], t["nota"]]


def aplicar_extractos(nuevos):
    por_hoja = {}
    for r in nuevos:
        por_hoja.setdefault(r["tarjeta"]["label"], []).append(r)

    for tarjeta_label, entradas in por_hoja.items():
        blocks = db.TARJETA_BLOCKS[tarjeta_label]
        det_row = db.first_blank_row(blocks["detalle"])
        det_usd_row = db.first_blank_row(blocks["detalle_usd"]) if "detalle_usd" in blocks else None
        vistos_periodo = set()
        for entrada in entradas:
            periodo = entrada["statement"]["periodo"]
            if periodo not in vistos_periodo:
                sr = db.first_blank_row(blocks["resumen"])
                # Un mismo extracto de Mastercard trae dos hojas (PESOS y
                # DOLARES) para el mismo período — cada una llega acá como
                # su propia "entrada", con su propio "statement". Los datos
                # generales de cupo/fechas y el "pago total" en pesos van en
                # las columnas de siempre (tomados de la entrada en COP); si
                # además hay una entrada en USD para este período, su propio
                # "pago total" (el saldo a pagar en dólares) va aparte, en
                # la columna 11 ("Saldo a pagar USD") — solo existe esa
                # columna en hojas con bloque separado para USD (Mastercard),
                # nunca se mezcla con el de pesos.
                entradas_periodo = [e for e in entradas if e["statement"]["periodo"] == periodo]
                cop_entrada = next((e for e in entradas_periodo if e["moneda"] == "COP"), None)
                usd_entrada = next((e for e in entradas_periodo if e["moneda"] == "USD"), None)
                # periodo/fecha_corte/fecha_limite son iguales sin importar la moneda de la
                # hoja, así que se toman de cualquier entrada disponible.
                s_fechas = (cop_entrada or usd_entrada)["statement"]
                segments = [
                    (1, [db.as_text(s_fechas["periodo"]), s_fechas["fecha_corte"].isoformat(),
                         s_fechas["fecha_limite"].isoformat() if s_fechas["fecha_limite"] else ""]),
                ]
                if cop_entrada is not None:
                    s = cop_entrada["statement"]
                    segments.append((4, [s["cupo_total"], s["cupo_disponible"]]))
                    segments.append((7, [s["saldo_anterior"]]))
                    segments.append((9, [s["pago_minimo"], s["pago_total"]]))
                if "detalle_usd" in blocks and usd_entrada is not None:
                    segments.append((11, [usd_entrada["statement"]["pago_total"]]))
                db.write_row_segments(blocks["resumen"], sr, segments)
                if cop_entrada is None and usd_entrada is not None:
                    st.warning(f"⚠️ {tarjeta_label} {periodo}: solo encontré la hoja en dólares — el cupo, "
                               "saldo anterior y pago total en pesos quedaron en blanco. Si el extracto trae "
                               "también la hoja en pesos, subila para completarlos.")
                vistos_periodo.add(periodo)

            # COP y otras monedas van a bloques separados — nunca se suman entre sí.
            cop_txns = [t for t in entrada["txns"] if t["moneda"] == "COP"]
            otras_txns = [t for t in entrada["txns"] if t["moneda"] != "COP"]

            if cop_txns:
                rows = [_fila_detalle_tarjeta(t) for t in cop_txns]
                db.write_block_rows(blocks["detalle"], det_row, rows)
                det_row += len(rows)

            if otras_txns:
                if "detalle_usd" not in blocks:
                    raise ValueError(f"{tarjeta_label} tiene movimientos en {otras_txns[0]['moneda']} pero "
                                      f"todavía no tiene un bloque separado para esa moneda — avisame para crearlo.")
                rows = [_fila_detalle_tarjeta(t) for t in otras_txns]
                db.write_block_rows(blocks["detalle_usd"], det_usd_row, rows)
                det_usd_row += len(rows)
    db.clear_read_cache()


def aplicar_movimientos_cuenta(nuevos_ingresos, nuevos_egresos):
    if nuevos_ingresos:
        row = db.first_blank_row("otros_ingresos")
        rows = [[t["fecha"].isoformat(), t["descripcion"], t["categoria"], t["valor"], t["nota"] or ""]
                for t in nuevos_ingresos]
        db.write_block_rows("otros_ingresos", row, rows)

    if nuevos_egresos:
        row = db.first_blank_row("efectivo_detalle")
        rows = [[db.as_text(t["fecha"].strftime("%Y-%m")), t["fecha"].isoformat(), t["descripcion"], "COP",
                 db.as_text("1/1"), abs(t["valor"]), abs(t["valor"]), 0,
                 t["categoria"], "No", t["nota"] or ""] for t in nuevos_egresos]
        db.write_block_rows("efectivo_detalle", row, rows)

    # Guarda ambos lados del flujo de inversión. Los depósitos al broker son
    # positivos; los retiros que vuelven a la 1031 son negativos, para que el
    # flujo neto no confunda un retorno de capital con ingreso nuevo.
    movimientos_inversion = [t for t in nuevos_ingresos + nuevos_egresos if t.get("flujo_inversion")]
    for moneda in ("pesos", "dolares"):
        movimientos = [t for t in movimientos_inversion if t["flujo_inversion"][1] == moneda]
        filas = [[t["fecha"].isoformat(), t["flujo_inversion"][0],
                  abs(t["valor"]) if t["valor"] < 0 else -abs(t["valor"]),
                  "Depósito detectado en extracto de cuenta" if t["valor"] < 0
                  else "Retiro detectado en extracto de cuenta"] for t in movimientos]
        db.agregar_aportes_inversion(moneda, filas)

    db.clear_read_cache()
    for moneda in ("pesos", "dolares"):
        _tomar_snapshot_cartera(moneda)


MESES_NOMBRE = {1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
                 7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre"}


def _extraer_anio_mes(periodo):
    """De un 'Periodo' de la tabla de movimientos (puede venir como
    'dd/mm/yyyy', 'yyyy-mm-dd' o como '1a quincena jul-2026') saca (año, mes)
    para poder filtrar el Dashboard por esas temporalidades. (None, None) si
    no se pudo reconocer el formato.

    'yyyy-mm-dd' aparece en 'Fecha Compra' cuando el extracto subido trae la
    fecha ya en ese formato (algunos extractos de 2026 en adelante) — antes
    de este chequeo, esas filas no matcheaban ningún patrón y quedaban
    afuera de todo cálculo por mes/año (Flujo de Efectivo, Estado de
    Resultados, etc.) sin ningún aviso, descuadrando el saldo calculado."""
    if not periodo:
        return None, None
    s = str(periodo)
    m = re.match(r"^\d{1,2}/(\d{1,2})/(\d{4})$", s)
    if m:
        return int(m.group(2)), int(m.group(1))
    m = re.match(r"^(\d{4})-(\d{1,2})-\d{1,2}$", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"^(\d{4})-(\d{1,2})$", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.search(r"([a-záéíóúñ]{3,4})\.?-(\d{4})", s.lower())
    if m:
        mes_txt, anio_txt = m.groups()
        mes_num = MESES_ABR.get(mes_txt.rstrip("."))
        if mes_num:
            return int(anio_txt), mes_num
    return None, None


def _shift_mes(mes_str, delta):
    """Suma `delta` meses a un 'YYYY-MM' (delta puede ser negativo). Ej:
    _shift_mes('2026-08', 1) -> '2026-09'; _shift_mes('2026-01', -1) -> '2025-12'."""
    anio, mes = int(mes_str[:4]), int(mes_str[5:7])
    total = anio * 12 + (mes - 1) + delta
    return f"{total // 12}-{total % 12 + 1:02d}"


def _fijar_eje_periodos(fig, valores):
    """Muestra etiquetas YYYY-MM como períodos, no como el día 1 de una fecha."""
    periodos = list(dict.fromkeys(str(valor) for valor in valores if pd.notna(valor)))
    fig.update_xaxes(type="category", categoryorder="array", categoryarray=periodos)
    return fig


def _gasto_real_categoria_mes(mes_str):
    """Gasto real por categoría de un mes puntual, en la misma convención
    'efectivo real' que usa la columna 'Gasto Real' de la hoja
    Presupuesto: el efectivo cuenta en su propio mes, las tarjetas en el
    mes siguiente al extracto (cuando se pagan)."""
    mes_corte_tarjetas = _shift_mes(mes_str, -1)
    totales = {}
    for block_key, mes_objetivo in [("efectivo_detalle", mes_str), ("visa_detalle", mes_corte_tarjetas),
                                      ("mc_detalle", mes_corte_tarjetas)]:
        df_b = db.read_egreso_detalle(block_key)
        if df_b.empty:
            continue
        df_b = df_b[(df_b["Moneda"] == "COP") & (df_b["Presupuestar"] == "Sí")
                    & (df_b["Periodo Extracto"] == mes_objetivo)]
        for cat, valor in df_b.groupby("Categoría")["Valor Cargado Este Periodo"].sum().items():
            totales[cat] = totales.get(cat, 0.0) + valor
    return totales


def _descuentos_categoria_mes(mes_str):
    """Descuentos de nómina por categoría de un mes puntual (por quincena
    de pago — mismo criterio que usa la hoja Presupuesto)."""
    df_desc = db.read_colillas_descuentos()
    if df_desc.empty:
        return {}
    am = [_extraer_anio_mes(q) for q in df_desc["Quincena"]]
    anio_obj, mes_obj = int(mes_str[:4]), int(mes_str[5:7])
    mask = [a == anio_obj and m == mes_obj for a, m in am]
    return df_desc.loc[mask].groupby("Categoría")["Valor"].sum().to_dict()


def _render_egreso_tab(block_key, titulo):
    """Cuerpo común de las pestañas del Dashboard que muestran una hoja de
    egresos (Efectivo, Visa, Mastercard) tal como está en el Sheet, con
    búsqueda y filtros de categoría/año/mes. El filtro de Año/Mes y los
    gráficos siempre se basan en 'Fecha Compra' (el día real de la compra),
    no en 'Periodo Extracto' (el corte de la tarjeta puede incluir compras
    de hasta dos meses calendario distintos). Para las tarjetas, el
    selector de abajo deja elegir en cambio 'efectivo real' — el mes en
    que se paga esa compra —, que sí se basa en 'Periodo Extracto'."""
    st.markdown(f"#### {titulo}")
    es_tarjeta = block_key != "efectivo_detalle"
    if es_tarjeta:
        st.caption("El filtro de Año/Mes de abajo es por la fecha real de la compra (columna 'Fecha Compra'), "
                   "no por 'Periodo Extracto' del corte — un corte puede incluir compras de hasta dos meses "
                   "calendario. El selector de abajo deja ver en cambio el mes en que se paga cada compra.")
    df = db.read_egreso_detalle(block_key)
    if df.empty:
        st.info(f"Todavía no hay movimientos cargados en {titulo}.")
        return

    # Mes de COMPRA (fecha real, columna 'Fecha Compra') — es la base del
    # filtro de Año/Mes y de los gráficos en modo consumo/devengado.
    am_compra = [_extraer_anio_mes(f) for f in df["Fecha Compra"]]
    df = df.assign(_anio=[a for a, _ in am_compra], _mes=[m for _, m in am_compra])
    if es_tarjeta:
        # Mes del CORTE ('Periodo Extracto') — solo se usa para "efectivo real".
        am_extracto = [_extraer_anio_mes(p) for p in df["Periodo Extracto"]]
        df = df.assign(_anio_extracto=[a for a, _ in am_extracto], _mes_extracto=[m for _, m in am_extracto])

    ocultar_no_presup = st.checkbox(
        "Ocultar movimientos de conciliación (no presupuestar)", value=True, key=f"ocultar_np_{block_key}",
        help="Movimientos como pagos automáticos de tarjeta, transferencias entre tus propias cuentas o "
             "intereses se cargan igual para poder conciliar el saldo real de la cuenta, pero no cuentan en "
             "tus totales de ingreso/gasto. Desmarcá esto para verlos también.")

    vista_gasto = None
    if es_tarjeta:
        vista_gasto = st.radio(
            "Vista de gasto", ["🛍️ Consumo (mes en que se compra)", "💳 Efectivo real (mes en que se paga)"],
            horizontal=True, key=f"vista_gasto_{block_key}")
    efectivo_mode = bool(vista_gasto) and vista_gasto.startswith("💳")

    busqueda = st.text_input("🔍 Buscar en comercio / concepto", "", key=f"busqueda_{block_key}")
    categorias = ["(todas)"] + sorted([c for c in df["Categoría"].unique().tolist() if isinstance(c, str) and c])
    anios = ["(todos)"] + sorted({str(a) for a in df["_anio"] if a}, reverse=True)
    meses = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]

    c1, c2, c3 = st.columns(3)
    cat_sel = c1.selectbox("Categoría", categorias, key=f"cat_{block_key}")
    anio_sel = c2.selectbox("Año", anios, key=f"anio_{block_key}")
    mes_sel = c3.selectbox("Mes", meses, key=f"mes_{block_key}")

    df_f = df
    if ocultar_no_presup:
        df_f = df_f[df_f["Presupuestar"] == "Sí"]
    if busqueda:
        df_f = df_f[df_f["Comercio / Concepto"].str.contains(busqueda, case=False, na=False)]
    if cat_sel != "(todas)":
        df_f = df_f[df_f["Categoría"] == cat_sel]

    if efectivo_mode and anio_sel != "(todos)" and mes_sel != "(todos)":
        # El Año/Mes elegido es el mes en que se PAGA — el corte que cae ahí
        # es el del mes anterior, así que filtramos por Periodo Extracto.
        corte = _shift_mes(f"{int(anio_sel)}-{int(mes_sel[:2]):02d}", -1)
        anio_corte, mes_corte_num = int(corte[:4]), int(corte[5:7])
        df_f = df_f[(df_f["_anio_extracto"] == anio_corte) & (df_f["_mes_extracto"] == mes_corte_num)]
        st.caption(f"Mostrando el corte de **{MESES_NOMBRE[mes_corte_num]} {anio_corte}** — es el que se paga "
                   f"en {mes_sel[5:]} {anio_sel}.")
    else:
        if anio_sel != "(todos)":
            df_f = df_f[df_f["_anio"] == int(anio_sel)]
        if mes_sel != "(todos)":
            df_f = df_f[df_f["_mes"] == int(mes_sel[:2])]

    total_f = df_f["Valor Cargado Este Periodo"].sum()
    st.caption(f"{len(df_f):,} de {len(df):,} movimientos — suma cargada este periodo: {fmt_moneda(total_f)}")
    cols_mostrar = [c for c in df_f.columns if not c.startswith("_")]
    st.dataframe(df_f[cols_mostrar], hide_index=True, use_container_width=True, height=600)

    df_f_cop = df_f[df_f["Moneda"] == "COP"] if "Moneda" in df_f.columns else df_f
    if not df_f_cop.empty:
        if efectivo_mode:
            periodo = (df_f_cop["_anio_extracto"].astype(int).astype(str) + "-"
                      + df_f_cop["_mes_extracto"].map(lambda m: f"{m:02d}"))
            df_f_cop = df_f_cop.assign(_periodo=periodo.map(lambda p: _shift_mes(p, 1)))
            etiqueta_periodo = "mes de pago"
        else:
            periodo = (df_f_cop["_anio"].astype(int).astype(str) + "-"
                      + df_f_cop["_mes"].map(lambda m: f"{m:02d}"))
            df_f_cop = df_f_cop.assign(_periodo=periodo)
            etiqueta_periodo = "mes de compra"

        st.markdown(f"**Tendencia por período** (según el filtro de arriba, por {etiqueta_periodo})")
        df_periodo = (df_f_cop.groupby("_periodo")["Valor Cargado Este Periodo"].sum()
                      .reset_index().sort_values("_periodo").rename(columns={"_periodo": "Período"}))
        fig_periodo = px.bar(df_periodo, x="Período", y="Valor Cargado Este Periodo")
        fig_periodo.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=300, xaxis_title="")
        _fijar_eje_periodos(fig_periodo, df_periodo["Período"])
        st.plotly_chart(fig_periodo, use_container_width=True)

        st.markdown(f"**Gasto por categoría y por mes** (según el filtro de arriba, por {etiqueta_periodo})")
        df_cat_mes = (df_f_cop.groupby(["_periodo", "Categoría"])["Valor Cargado Este Periodo"]
                      .sum().reset_index().rename(columns={"_periodo": "Período"}))
        fig_cat_mes = px.bar(df_cat_mes, x="Período", y="Valor Cargado Este Periodo", color="Categoría",
                             barmode="stack")
        fig_cat_mes.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=420, xaxis_title="")
        _fijar_eje_periodos(fig_cat_mes, df_cat_mes["Período"])
        st.plotly_chart(fig_cat_mes, use_container_width=True)


def render_categorias():
    st.markdown("#### Gasto real por categoría")
    st.caption("Vista **devengado**: total histórico por categoría, cada compra contada en su propia "
               "fecha (no en el mes en que pagaste la tarjeta). Para ver el flujo de caja real mes a "
               "mes, con las tarjetas atribuidas al mes en que se pagan, usá Balance Mensual o Evolución.")
    df_cat = db.read_resumen_categorias()
    if not df_cat.empty:
        fig = px.bar(df_cat, x="Gasto Real", y="Categoría", orientation="h")
        fig.update_layout(yaxis=dict(autorange="reversed"), margin=dict(l=0, r=0, t=10, b=0),
                          height=max(420, 32 * len(df_cat)))
        st.plotly_chart(fig, use_container_width=True)
        st.dataframe(df_cat, hide_index=True, use_container_width=True)
    else:
        st.info("Todavía no hay gastos cargados para graficar por categoría.")


def render_esenciales():
    st.markdown("#### Gasto esencial vs. no esencial")
    st.caption("Incluye solo gasto de consumo real en pesos (compras de Visa/Mastercard en COP y Egresos - "
               "Efectivo, sin los movimientos '(no presupuestar)' de conciliación). No incluye compras en "
               "USD, descuentos de nómina (Impuestos, Salud, Pensión, etc. — esos están en Balance Mensual → "
               "Descuentos de nómina) ni Ahorro/Inversiones, porque no son consumo. "
               "La clasificación de cada categoría se lee de la hoja **'Categorías Esenciales'** del Sheet — "
               "editala ahí directamente para corregirla, la app la respeta.")

    clasificacion_sheet = db.read_clasificacion_esencial()

    def _clasificar(categoria):
        return clasificacion_sheet.get(categoria) or clasificar_esencial(categoria)

    df_mov_es = db.read_movimientos_recientes()
    if not df_mov_es.empty:
        df_gasto = df_mov_es[
            (df_mov_es["Valor"] < 0)
            & (df_mov_es["Moneda"] == "COP")
            & (df_mov_es["Presupuestar"] == "Sí")
            & (df_mov_es["Fuente"] != "Colilla (descuento)")
        ].copy()
        df_gasto["Valor"] = df_gasto["Valor"].abs()
        df_gasto["Clasificación"] = df_gasto["Categoría"].map(_clasificar)
        df_gasto = df_gasto[df_gasto["Clasificación"] != "No consumo (ahorro/inversión)"]

        if not df_gasto.empty:
            anios_meses_es = [_extraer_anio_mes(p) for p in df_gasto["Periodo"]]
            df_gasto = df_gasto.assign(_anio=[a for a, _ in anios_meses_es], _mes=[m for _, m in anios_meses_es])

            anios_es = ["(todos)"] + sorted({str(a) for a in df_gasto["_anio"] if a}, reverse=True)
            meses_es = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]
            c1, c2 = st.columns(2)
            anio_sel_es = c1.selectbox("Año", anios_es, key="anio_esenciales")
            mes_sel_es = c2.selectbox("Mes", meses_es, key="mes_esenciales")

            df_f = df_gasto
            if anio_sel_es != "(todos)":
                df_f = df_f[df_f["_anio"] == int(anio_sel_es)]
            if mes_sel_es != "(todos)":
                df_f = df_f[df_f["_mes"] == int(mes_sel_es[:2])]

            if not df_f.empty:
                resumen_clas = df_f.groupby("Clasificación")["Valor"].sum().reset_index()
                total_gasto = resumen_clas["Valor"].sum()

                c3, c4, c5 = st.columns(3)
                for col, clasif in zip((c3, c4, c5), ("Esencial", "No esencial", "Sin clasificar")):
                    valor_clasif = resumen_clas.loc[resumen_clas["Clasificación"] == clasif, "Valor"].sum()
                    pct = (valor_clasif / total_gasto * 100) if total_gasto else 0
                    col.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(valor_clasif)}</div>'
                                 f'<div class="stat-label">{clasif} ({pct:.0f}%)</div></div>',
                                 unsafe_allow_html=True)

                fig_es = px.pie(resumen_clas, names="Clasificación", values="Valor", hole=0.4)
                fig_es.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=320)
                st.plotly_chart(fig_es, use_container_width=True)

                st.markdown("**Detalle por categoría**")
                detalle_cat = (df_f.groupby(["Clasificación", "Categoría"])["Valor"].sum()
                               .reset_index().sort_values(["Clasificación", "Valor"], ascending=[True, False]))
                st.dataframe(detalle_cat, hide_index=True, use_container_width=True)

                if "Sin clasificar" in resumen_clas["Clasificación"].values:
                    sin_clasificar_cats = sorted(df_f.loc[df_f["Clasificación"] == "Sin clasificar",
                                                           "Categoría"].unique())
                    st.warning("Categorías sin clasificar todavía (avisame si son esenciales o no): "
                               + ", ".join(sin_clasificar_cats))
            else:
                st.info("No hay gasto para ese año/mes.")
        else:
            st.info("Todavía no hay gasto de consumo para clasificar.")
    else:
        st.info("Todavía no hay movimientos cargados.")


def render_evolucion():
    st.markdown("#### Evolución mensual")
    st.caption("Vista **efectivo real**: el gasto de tarjeta de cada mes es el que efectivamente se "
               "pagó ese mes (no el de las compras hechas ese mes) — Visa/Mastercard se cuentan un mes "
               "después de la fecha de corte, cuando se paga. Efectivo ya sale en su propio mes.")
    df_mes = db.read_resumen_mensual()
    if not df_mes.empty:
        mes_col = df_mes.columns[0]
        anios_disponibles = sorted({str(m)[:4] for m in df_mes[mes_col] if str(m)[:4].isdigit()}, reverse=True)
        anio_ev = st.selectbox("Año", ["(todos)"] + anios_disponibles, key="anio_evolucion")
        df_mes_f = df_mes if anio_ev == "(todos)" else df_mes[df_mes[mes_col].astype(str).str.startswith(anio_ev)]
        num_cols = [c for c in df_mes.columns[1:] if c]
        df_long = df_mes_f.melt(id_vars=[mes_col], value_vars=num_cols, var_name="Concepto", value_name="Valor")
        fig2 = px.line(df_long, x=mes_col, y="Valor", color="Concepto", markers=True)
        fig2.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=380)
        _fijar_eje_periodos(fig2, df_mes_f[mes_col])
        st.plotly_chart(fig2, use_container_width=True)
        with st.expander("Ver tabla"):
            st.dataframe(df_mes_f, hide_index=True, use_container_width=True)
    else:
        st.info("Todavía no hay suficientes meses cargados para la evolución mensual.")


def render_anio_vs_anio():
    st.markdown("#### Año vs. Año")
    st.caption("Vista **efectivo real** (igual que Evolución): compará el mismo mes entre distintos años "
               "para ver si vas mejor o peor que antes, no solo si vas mejor o peor que el mes pasado.")
    df_mes = db.read_resumen_mensual()
    if df_mes.empty:
        st.info("Todavía no hay suficientes meses cargados para comparar años.")
        return

    hoy_str = date.today().strftime("%Y-%m")
    df_mes = df_mes[df_mes["Mes"] <= hoy_str].copy()
    if df_mes.empty:
        st.info("Todavía no hay suficientes meses cargados para comparar años.")
        return
    df_mes["_anio"] = df_mes["Mes"].str[:4].astype(int)
    df_mes["_mes_num"] = df_mes["Mes"].str[5:7].astype(int)

    metricas = [c for c in df_mes.columns if c not in ("Mes", "_anio", "_mes_num")]
    metrica = st.selectbox("Métrica a comparar", metricas, key="metrica_anio_vs_anio")

    anios_disp = sorted(df_mes["_anio"].unique())
    anios_sel = st.multiselect("Años a comparar", anios_disp,
                                default=anios_disp[-2:] if len(anios_disp) >= 2 else anios_disp,
                                key="anios_anio_vs_anio")
    if not anios_sel:
        st.info("Elegí al menos un año.")
        return

    df_f = df_mes[df_mes["_anio"].isin(anios_sel)].copy()
    df_f["Año"] = df_f["_anio"].astype(str)
    df_f["Mes_nombre"] = df_f["_mes_num"].map(MESES_NOMBRE)

    fig = px.bar(df_f, x="Mes_nombre", y=metrica, color="Año", barmode="group",
                 category_orders={"Mes_nombre": [MESES_NOMBRE[m] for m in range(1, 13)]})
    fig.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=380, xaxis_title=None)
    st.plotly_chart(fig, use_container_width=True)

    totales = df_f.groupby("_anio")[metrica].sum().sort_index()
    filas_tot = []
    anterior = None
    for anio, total in totales.items():
        variacion = f"{(total - anterior) / abs(anterior) * 100:+.0f}%" if anterior else "-"
        filas_tot.append({"Año": anio, f"Total {metrica}": f"{fmt_moneda(total)}", "Variación": variacion})
        anterior = total
    st.dataframe(pd.DataFrame(filas_tot), hide_index=True, use_container_width=True)
    st.caption("La variación compara cada año contra el anterior de esta misma tabla (no necesariamente el "
               "año calendario inmediatamente anterior, si no elegiste años consecutivos).")


def render_movimientos():
    st.markdown("#### Movimientos")
    st.caption("Vista **devengado**: cada movimiento en su propia fecha (fecha de compra en tarjetas, "
               "no fecha de pago).")
    df_mov = db.read_movimientos_recientes()
    if not df_mov.empty:
        anios_meses = [_extraer_anio_mes(p) for p in df_mov["Periodo"]]
        df_mov = df_mov.assign(_anio=[a for a, _ in anios_meses], _mes=[m for _, m in anios_meses])

        ocultar_np_mov = st.checkbox(
            "Ocultar movimientos de conciliación (no presupuestar)", value=True, key="ocultar_np_movimientos",
            help="Algunos movimientos bancarios (pago automático de tarjeta, transferencias entre tus "
                 "propias cuentas, intereses...) se cargan igual que cualquier otro para que la cuenta "
                 "cuadre contra el extracto del banco, pero no cuentan en tus totales de ingreso/gasto — "
                 "por eso, por ejemplo, tu nómina puede aparecer acá dos veces: una vez como Colilla "
                 "(sí cuenta) y otra desde el extracto de la cuenta (no cuenta, solo conciliación). "
                 "Desmarcá esto para ver también esos movimientos de conciliación.")

        busqueda = st.text_input("🔍 Buscar en concepto", "", key="busqueda_movimientos")

        fuentes = ["(todas)"] + sorted(df_mov["Fuente"].unique().tolist())
        categorias = ["(todas)"] + sorted([c for c in df_mov["Categoría"].unique().tolist() if isinstance(c, str) and c])
        anios = ["(todos)"] + sorted({str(a) for a in df_mov["_anio"] if a}, reverse=True)
        meses = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]

        c1, c2, c3, c4 = st.columns(4)
        f_sel = c1.selectbox("Fuente", fuentes)
        cat_sel = c2.selectbox("Categoría", categorias)
        anio_sel = c3.selectbox("Año", anios, key="anio_movimientos")
        mes_sel = c4.selectbox("Mes", meses, key="mes_movimientos")

        df_f = df_mov
        if ocultar_np_mov:
            df_f = df_f[df_f["Presupuestar"] == "Sí"]
        if busqueda:
            df_f = df_f[df_f["Concepto"].str.contains(busqueda, case=False, na=False)]
        if f_sel != "(todas)":
            df_f = df_f[df_f["Fuente"] == f_sel]
        if cat_sel != "(todas)":
            df_f = df_f[df_f["Categoría"] == cat_sel]
        if anio_sel != "(todos)":
            df_f = df_f[df_f["_anio"] == int(anio_sel)]
        if mes_sel != "(todos)":
            df_f = df_f[df_f["_mes"] == int(mes_sel[:2])]

        # OJO: nunca sumar "Valor" mezclando monedas (los USD de Mastercard
        # se irían directo al total en pesos). Se suma por moneda, aparte.
        totales_moneda = df_f.groupby("Moneda")["Valor"].sum()
        resumen_totales = " · ".join(
            f"{'US$' if moneda == 'USD' else '$'} {valor:,.2f}" if moneda == "USD" else f"{fmt_moneda(valor)}"
            for moneda, valor in totales_moneda.items()
        )
        st.caption(f"{len(df_f):,} de {len(df_mov):,} movimientos — suma: {resumen_totales}")
        st.dataframe(df_f.drop(columns=["_anio", "_mes"]), hide_index=True, use_container_width=True, height=600)

        # -- Gráfico: top categorías de gasto (según el filtro de arriba) --
        st.markdown("**Top categorías de gasto (según el filtro de arriba)**")
        df_gasto_cop = df_f[(df_f["Moneda"] == "COP") & (df_f["Valor"] < 0)].copy()
        if not df_gasto_cop.empty:
            df_gasto_cop["Valor"] = df_gasto_cop["Valor"].abs()
            top_cat = (df_gasto_cop.groupby("Categoría")["Valor"].sum()
                       .reset_index().sort_values("Valor", ascending=False).head(12))
            fig_top_cat = px.bar(top_cat, x="Valor", y="Categoría", orientation="h")
            fig_top_cat.update_layout(yaxis=dict(autorange="reversed"), margin=dict(l=0, r=0, t=10, b=0),
                                      height=max(320, 30 * len(top_cat)))
            st.plotly_chart(fig_top_cat, use_container_width=True)
        else:
            st.info("No hay gasto en pesos para graficar con el filtro actual.")
    else:
        st.info("Todavía no hay movimientos cargados.")


def render_balance_mensual():
    st.markdown("#### Balance Mensual")
    st.caption("Vista **efectivo real** — todo lo que se paga e ingresa efectivamente este mes: ingreso bruto, egreso de "
               "tarjetas/efectivo y descuentos de nómina categorizados (también son egreso). "
               "El egreso de tarjeta se cuenta en el mes en que realmente se paga, no en el mes del "
               "extracto (p. ej. la Visa con corte 15 de agosto se paga el 2 de septiembre, así que "
               "cuenta como egreso de septiembre) — el efectivo sí se cuenta en su propio mes, porque ya "
               "es lo que efectivamente salió de la cuenta de ahorros. "
               "No se suman los meses entre sí — para ver la evolución de varios meses usá la pestaña Evolución.")

    hoy = date.today()
    anios_balance = list(range(2023, 2033))
    c1, c2 = st.columns(2)
    anio_bal = c1.selectbox("Año", anios_balance,
                             index=anios_balance.index(hoy.year) if hoy.year in anios_balance else 0,
                             key="anio_balance")
    mes_bal_num = c2.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                format_func=lambda m: MESES_NOMBRE[m], key="mes_balance")
    mes_bal_nombre = MESES_NOMBRE[mes_bal_num]

    if st.session_state.get("_balance_mes_aplicado") != (anio_bal, mes_bal_nombre):
        try:
            try:
                db.set_balance_mensual_mes(anio_bal, mes_bal_nombre)
            except Exception:
                time.sleep(2)  # reintento único — la API de Sheets a veces falla de forma transitoria
                db.set_balance_mensual_mes(anio_bal, mes_bal_nombre)
            db.clear_read_cache()
            st.session_state["_balance_mes_aplicado"] = (anio_bal, mes_bal_nombre)
        except Exception as e:
            st.error(f"No pude cambiar el mes en la hoja 'Balance Mensual': {e}. Probá refrescar la página.")
            st.stop()

    datos = db.read_balance_mensual()

    vista_gasto = st.radio(
        "Vista de gasto", ["💳 Efectivo real (mes en que se paga)", "🛍️ Consumo (mes en que se compra)"],
        horizontal=True, key="vista_gasto_balance",
        help="Efectivo real: el gasto de tarjeta cuenta en el mes en que se paga (igual que la explicación de "
             "arriba). Consumo: el mismo mes, pero contando cada compra en el mes en que efectivamente se hizo, "
             "sin importar cuándo se paga la tarjeta.")

    def _egresos_consumo_mes(mes_str):
        """Egresos de este mismo mes calendario, contando cada compra de
        tarjeta en el mes en que se hizo de verdad (columna 'Fecha Compra',
        no 'Periodo Extracto' — el corte de la tarjeta puede incluir
        compras de hasta dos meses calendario distintos) — a diferencia de
        read_balance_mensual(), que viene desplazado por la fórmula de la
        hoja 'Balance Mensual' según el corte, no la fecha de compra."""
        anio_obj, mes_obj = int(mes_str[:4]), int(mes_str[5:7])
        total = 0.0
        filas = []
        for block_key, nombre in [("efectivo_detalle", "Efectivo"), ("visa_detalle", "Visa"),
                                    ("mc_detalle", "Mastercard")]:
            df_b = db.read_egreso_detalle(block_key)
            valor = 0.0
            if not df_b.empty:
                df_b = df_b[(df_b["Moneda"] == "COP") & (df_b["Presupuestar"] == "Sí")]
                am = [_extraer_anio_mes(f) for f in df_b["Fecha Compra"]]
                mask = [a == anio_obj and m == mes_obj for a, m in am]
                valor = df_b.loc[mask, "Valor Cargado Este Periodo"].sum()
            filas.append({"Método": nombre, "Valor": valor})
            total += valor
        return total, pd.DataFrame(filas)

    if vista_gasto.startswith("🛍️"):
        egresos_mostrar, tabla_metodo_mostrar = _egresos_consumo_mes(f"{anio_bal}-{mes_bal_num:02d}")
        balance_mostrar = datos["ingresos_brutos"] - egresos_mostrar - datos["descuentos_nomina"]
        etiqueta_egresos = "Gasto de consumo (tarjetas + efectivo)"
    else:
        egresos_mostrar = datos["egresos_tarjetas_efectivo"]
        tabla_metodo_mostrar = datos["egresos_por_metodo"]
        balance_mostrar = datos["balance"]
        etiqueta_egresos = "Egresos (tarjetas + efectivo)"

    c3, c4, c5, c6 = st.columns(4)
    c3.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(datos["ingresos_brutos"])}</div>'
                f'<div class="stat-label">Ingresos brutos</div></div>', unsafe_allow_html=True)
    c4.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(egresos_mostrar)}</div>'
                f'<div class="stat-label">{etiqueta_egresos}</div></div>', unsafe_allow_html=True)
    c5.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(datos["descuentos_nomina"])}</div>'
                f'<div class="stat-label">Descuentos de nómina</div></div>', unsafe_allow_html=True)
    c6.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(balance_mostrar)}</div>'
                f'<div class="stat-label">Balance del mes</div></div>', unsafe_allow_html=True)

    col_a, col_b, col_c = st.columns(3)
    with col_a:
        st.markdown("**Egresos por método de pago**")
        st.dataframe(tabla_metodo_mostrar, hide_index=True, use_container_width=True)
    with col_b:
        st.markdown("**Descuentos de nómina por categoría**")
        st.dataframe(datos["descuentos_por_categoria"], hide_index=True, use_container_width=True)
    with col_c:
        st.markdown("**Ingresos por fuente**")
        st.dataframe(datos["ingresos_por_fuente"], hide_index=True, use_container_width=True)


def _form_actualizar_deuda(sheet_name, titulo, key_sufijo):
    datos = db.read_deuda_entidad(sheet_name)
    faltan = datos["saldo_actual"] is None
    if faltan:
        st.warning(f"Todavía no hay datos guardados para '{titulo}'.")
    st.caption("Ninguna de las 4 deudas con cuota fija tiene extracto automático en la app — actualizá "
               "'Saldo Actual', 'Tasa E.A.' y 'Cuota Mensual' cada vez que te llegue un extracto nuevo. "
               "Los campos numéricos en 0 y el N° de Obligación vacío no se guardan (para no pisar lo que "
               "ya había con un cero).")
    with st.form(f"form_deuda_{key_sufijo}"):
        st.markdown("**Lo esencial** — esto es lo que hace que aparezca en la tabla de arriba")
        c1, c2, c3 = st.columns(3)
        saldo_actual = c1.number_input("Saldo Actual", value=float(datos["saldo_actual"] or 0),
                                        step=10000.0, format="%.0f", key=f"saldo_{key_sufijo}")
        tasa_ea = c2.number_input("Tasa E.A. (%)", value=float((datos["tasa_ea"] or 0) * 100),
                                   step=0.1, format="%.2f", key=f"tasa_{key_sufijo}")
        cuota_mensual = c3.number_input("Cuota Mensual", value=float(datos["cuota_mensual"] or 0),
                                         step=10000.0, format="%.0f", key=f"cuota_{key_sufijo}")

        with st.expander("Datos adicionales (opcional, para el pronóstico de pagos completo)"):
            c4, c5, c6 = st.columns(3)
            plazo_meses = c4.number_input("Plazo (meses)", value=int(datos["plazo_meses"] or 0), step=1,
                                           key=f"plazo_{key_sufijo}")
            numero_cuota_actual = c5.number_input("N° Cuota Actual",
                                                   value=int(datos["numero_cuota_actual"] or 0), step=1,
                                                   key=f"ncuota_{key_sufijo}")
            numero_obligacion = c6.text_input("N° Obligación", value=str(datos["numero_obligacion"] or ""),
                                               key=f"nobl_{key_sufijo}")
            c7, c8, c9 = st.columns(3)
            fecha_saldo = c7.date_input("Fecha del Saldo", value=datos["fecha_saldo"] or date.today(),
                                         key=f"fsaldo_{key_sufijo}")
            proxima_fecha_pago = c8.date_input("Próxima Fecha de Pago",
                                                value=datos["proxima_fecha_pago"] or date.today(),
                                                key=f"fprox_{key_sufijo}")
            fecha_desembolso = c9.date_input("Fecha Desembolso",
                                              value=datos["fecha_desembolso"] or date.today(),
                                              key=f"fdes_{key_sufijo}")
            monto_inicial = st.number_input("Monto Inicial", value=float(datos["monto_inicial"] or 0),
                                             step=10000.0, format="%.0f", key=f"minicial_{key_sufijo}")
            nota = st.text_area("Nota", value=datos["nota"] or "", key=f"nota_{key_sufijo}")

        guardar = st.form_submit_button("💾 Guardar", type="primary")

    if guardar:
        valores = {"fecha_saldo": fecha_saldo, "proxima_fecha_pago": proxima_fecha_pago,
                   "fecha_desembolso": fecha_desembolso, "nota": nota}
        for campo, valor in [("saldo_actual", saldo_actual), ("tasa_ea", tasa_ea / 100),
                              ("cuota_mensual", cuota_mensual), ("plazo_meses", plazo_meses),
                              ("numero_cuota_actual", numero_cuota_actual), ("monto_inicial", monto_inicial)]:
            if valor:
                valores[campo] = valor
        if numero_obligacion:
            valores["numero_obligacion"] = numero_obligacion
        try:
            db.set_deuda_entidad(sheet_name, valores)
            db.clear_read_cache()
            st.success(f"Datos de '{titulo}' guardados.")
            st.rerun()
        except Exception as e:
            st.error(f"No pude guardar: {e}")


def render_deudas():
    st.title("🏦 Deudas y créditos")
    st.caption("Los 4 créditos con cuota fija (educativo, consumo, leasing, fondo de empleados) más el saldo "
               "a pagar de Visa y Mastercard en pesos — las tarjetas no tienen tasa/cuota/plazo fijo, así que "
               "esas columnas quedan vacías para ellas. La deuda de Mastercard en dólares se muestra aparte "
               "más abajo — nunca se suma con el total en pesos.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Saldo actual y cuota mensual de cada deuda, con gráficos de participación por entidad y de "
            "cuota mensual por entidad. Es un estado de hoy — no cambia con ningún selector de mes, porque "
            "un saldo pendiente no es algo que se acumule mes a mes como un gasto. Al final hay un formulario "
            "por cada una de las 4 deudas con cuota fija — ninguna tiene extracto automático en la app, se "
            "actualizan a mano cada vez que llega un estado de cuenta nuevo.")
    df_deudas = db.read_deudas_resumen()
    if not df_deudas.empty:
        c1, c2 = st.columns(2)
        c1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(df_deudas["Saldo Actual"].sum())}</div>'
                    f'<div class="stat-label">Saldo total pendiente (pesos)</div></div>', unsafe_allow_html=True)
        c2.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(df_deudas["Cuota Mensual"].sum())}</div>'
                    f'<div class="stat-label">Cuota mensual total</div></div>', unsafe_allow_html=True)

        deuda_usd = db.read_deuda_tarjeta_usd()
        if deuda_usd:
            entidad_usd, nota_usd, valor_usd = deuda_usd
            st.markdown(f'<div class="stat-card"><div class="stat-num">US$ {valor_usd:,.2f}</div>'
                        f'<div class="stat-label">{entidad_usd} — deuda en dólares (aparte)</div></div>',
                        unsafe_allow_html=True)
            st.caption(nota_usd)

        df_con_saldo = df_deudas[df_deudas["Saldo Actual"] > 0]
        if not df_con_saldo.empty:
            fig3 = px.pie(df_con_saldo, names="Entidad", values="Saldo Actual", hole=0.4)
            fig3.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=320)
            st.plotly_chart(fig3, use_container_width=True)

        df_show = df_deudas.copy()
        df_show["Saldo Actual"] = df_show["Saldo Actual"].map(lambda v: "-" if pd.isna(v) else f"{fmt_moneda(v)}")
        df_show["Cuota Mensual"] = df_show["Cuota Mensual"].map(lambda v: "-" if pd.isna(v) else f"{fmt_moneda(v)}")
        df_show["Tasa E.A."] = df_show["Tasa E.A."].map(lambda v: "-" if pd.isna(v) else f"{v * 100:,.2f}%")
        df_show["% Pagado"] = df_show["% Pagado"].map(lambda v: "-" if pd.isna(v) else f"{v * 100:,.1f}%")
        df_show["Meses Restantes Est."] = df_show["Meses Restantes Est."].map(
            lambda v: "-" if pd.isna(v) else f"{v:,.0f}")
        df_show["Fecha Est. de Pago Total"] = df_show["Fecha Est. de Pago Total"].fillna("-")
        st.dataframe(df_show, hide_index=True, use_container_width=True)
    else:
        st.info("Todavía no hay deudas cargadas.")
    # -- Gráfico nuevo: cuota mensual por entidad (además de la torta de saldos) --
    df_con_cuota = df_deudas.dropna(subset=["Cuota Mensual"])
    df_con_cuota = df_con_cuota[df_con_cuota["Cuota Mensual"] > 0]
    if not df_con_cuota.empty:
        st.markdown("**Cuota mensual por entidad**")
        fig_cuotas = px.bar(df_con_cuota.sort_values("Cuota Mensual", ascending=False),
                            x="Entidad", y="Cuota Mensual")
        fig_cuotas.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=320)
        st.plotly_chart(fig_cuotas, use_container_width=True)

    st.divider()
    st.markdown("#### ✏️ Actualizar deudas manualmente")
    entidades_deuda = [
        (db.SHEET_DEUDA_BANCOLOMBIA, "Bancolombia (Crédito Educativo)", "bancolombia"),
        (db.SHEET_DEUDA_SUFI, "Sufi (Crédito de Consumo)", "sufi"),
        (db.SHEET_DEUDA_SCOTIABANK, "Scotiabank Colpatria (Leasing Habitacional)", "scotiabank"),
        (db.SHEET_DEUDA_FONDO_EMPLEADOS, "Fondo de Empleados", "fondo_empleados"),
    ]
    tabs_deuda = st.tabs([titulo for _, titulo, _ in entidades_deuda])
    for tab, (sheet_name, titulo, key_sufijo) in zip(tabs_deuda, entidades_deuda):
        with tab:
            _form_actualizar_deuda(sheet_name, titulo, key_sufijo)


def _form_agregar_aporte(moneda, plataformas, key_sufijo):
    with st.expander("➕ Agregar un aporte o retiro manualmente"):
        with st.form(f"form_agregar_aporte_{key_sufijo}"):
            fecha_ap = st.date_input("Fecha", value=date.today(), key=f"fecha_aporte_{key_sufijo}")
            plataforma_ap = st.selectbox("Plataforma", plataformas, key=f"plataforma_aporte_{key_sufijo}")
            tipo_ap = st.radio("Tipo", ["Depósito", "Retiro"], horizontal=True, key=f"tipo_aporte_{key_sufijo}")
            valor_ap = st.number_input("Monto transferido (COP)", min_value=0.0, step=10000.0, format="%.0f",
                                        key=f"valor_aporte_{key_sufijo}")
            notas_ap = st.text_input("Notas (opcional)", key=f"notas_aporte_{key_sufijo}")
            guardar_ap = st.form_submit_button("💾 Guardar", type="primary")
        if guardar_ap:
            if valor_ap <= 0:
                st.error("El valor tiene que ser mayor que cero.")
            else:
                monto = valor_ap if tipo_ap == "Depósito" else -valor_ap
                block_key = f"aportes_inversion_{moneda}"
                try:
                    fila = db.first_blank_row(block_key)
                    db.write_block_rows(block_key, fila,
                                         [[fecha_ap.isoformat(), plataforma_ap, monto, notas_ap.strip()]])
                    db.clear_read_cache()
                    _tomar_snapshot_cartera(moneda)
                    st.success(f"{tipo_ap} de {fmt_moneda(valor_ap)} agregado.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"No pude guardar: {e}")


def _num_o_cero(v):
    return float(v) if pd.notna(v) else 0.0


def _tomar_snapshot_cartera(moneda):
    """Guarda un snapshot de hoy (valor de costo, valor actual, aportes
    netos) para la moneda dada — la base del gráfico de Crecimiento y
    Rentabilidad. Se llama cada vez que cambian posiciones o aportes; no
    hace falta que el usuario haga nada aparte."""
    df_pos = db.read_posiciones_inversion(moneda)
    cantidad = pd.to_numeric(df_pos["Cantidad"], errors="coerce").fillna(0.0) if not df_pos.empty else pd.Series(dtype=float)
    precio_compra = pd.to_numeric(df_pos["Precio Compra Promedio"], errors="coerce").fillna(0.0) if not df_pos.empty else pd.Series(dtype=float)
    precio_actual = pd.to_numeric(df_pos["Precio Actual"], errors="coerce").fillna(0.0) if not df_pos.empty else pd.Series(dtype=float)
    valor_costo = float((cantidad.abs() * precio_compra).sum())
    valor_actual = float((cantidad * precio_actual).sum())
    df_ap = db.read_aportes_inversion(moneda)
    aportes_netos = float(pd.to_numeric(df_ap["Monto Transferido (COP)"], errors="coerce").fillna(0.0).sum()
                          if not df_ap.empty else 0.0)
    if valor_costo == 0 and valor_actual == 0 and aportes_netos == 0:
        return
    try:
        db.guardar_snapshot_cartera(moneda, date.today(), valor_costo, valor_actual, aportes_netos)
    except Exception:
        pass  # el snapshot es un plus informativo — nunca debe romper la acción que lo disparó


def _form_editar_aportes(moneda, key_sufijo):
    columnas = ["Fecha", "Plataforma", "Monto Transferido (COP)", "Notas"]
    df_ap = db.read_aportes_inversion(moneda)
    if df_ap.empty:
        df_ap = pd.DataFrame(columns=columnas)
    with st.expander(f"✏️ Editar o eliminar aportes/retiros ({len(df_ap)})"):
        st.caption("Corregí cualquier celda, o borrá una fila con el ícono que aparece a la izquierda al pasar "
                   "el mouse. Un depósito lleva monto positivo; un retiro, negativo. Usá el '+' de abajo para "
                   "agregar una fila nueva sin pasar por el formulario de arriba.")
        df_editado = st.data_editor(
            df_ap, hide_index=True, use_container_width=True, num_rows="dynamic",
            height=min(420, 80 + 35 * (len(df_ap) + 1)), key=f"editor_aportes_{key_sufijo}",
            column_config={"Monto Transferido (COP)": st.column_config.NumberColumn(format="$%d")})
        if st.button("💾 Guardar cambios", key=f"btn_guardar_aportes_{key_sufijo}"):
            filas = []
            for _, fila in df_editado.iterrows():
                fecha = fila.get("Fecha")
                if fecha in (None, "") or (isinstance(fecha, float) and pd.isna(fecha)):
                    continue
                filas.append([str(fecha), str(fila.get("Plataforma") or ""),
                              _num_o_cero(fila.get("Monto Transferido (COP)")), str(fila.get("Notas") or "")])
            try:
                db.set_aportes_inversion(moneda, filas)
                db.clear_read_cache()
                _tomar_snapshot_cartera(moneda)
                st.success(f"{len(filas)} aporte(s)/retiro(s) guardado(s).")
                st.rerun()
            except db.SinEspacioError as e:
                st.error(str(e))
            except Exception as e:
                st.error(f"No pude guardar: {e}")


def _form_editar_historial():
    df_hist = db.read_historial_inversion()
    with st.expander(f"✏️ Editar o eliminar operaciones del historial ({len(df_hist)})"):
        st.caption("Corregí una operación mal importada o borrala con el ícono a la izquierda de la fila. "
                   "Operación: BUY (compra), SELL (venta), SHORT (venta en corto), COVER (cobertura de corto), "
                   "DIVIDEND/INTEREST (dividendo o interés recibido, no afecta cantidad ni cuenta como "
                   "compraventa). Cambiar cantidad/precio/comisión NO recalcula 'Resultado Realizado' solo — "
                   "ajustalo a mano si corresponde.")
        df_editado = st.data_editor(
            df_hist, hide_index=True, use_container_width=True, num_rows="dynamic",
            height=min(500, 80 + 35 * (len(df_hist) + 1)), key="editor_historial_inversion",
            column_config={
                "Operación": st.column_config.SelectboxColumn(
                    options=["BUY", "SELL", "SHORT", "COVER", "DIVIDEND", "INTEREST"]),
                "Moneda": st.column_config.SelectboxColumn(options=["COP", "USD"]),
                "Cantidad": st.column_config.NumberColumn(format="%.6f"),
                "Precio": st.column_config.NumberColumn(format="%.4f"),
                "Comisión": st.column_config.NumberColumn(format="%.4f"),
                "Resultado Realizado": st.column_config.NumberColumn(format="%.2f"),
            })
        if st.button("💾 Guardar cambios en el historial", key="btn_guardar_historial"):
            filas = []
            for _, fila in df_editado.iterrows():
                fecha = fila.get("Fecha")
                activo = str(fila.get("Activo") or "").strip()
                if not activo or fecha in (None, "") or (isinstance(fecha, float) and pd.isna(fecha)):
                    continue
                filas.append([str(fecha), str(fila.get("Plataforma") or ""), str(fila.get("Moneda") or ""),
                              activo, str(fila.get("Operación") or ""), _num_o_cero(fila.get("Cantidad")),
                              _num_o_cero(fila.get("Precio")), _num_o_cero(fila.get("Comisión")),
                              _num_o_cero(fila.get("Resultado Realizado")), str(fila.get("Fuente") or "")])
            try:
                db.set_historial_inversion(filas)
                db.clear_read_cache()
                st.success(f"{len(filas)} operación(es) guardada(s).")
                st.rerun()
            except Exception as e:
                st.error(f"No pude guardar: {e}")


def _form_agregar_dividendo():
    with st.expander("💵 Agregar un dividendo o interés recibido"):
        st.caption("Ingreso pasivo real (no viene de vender nada) — se guarda aparte del resultado realizado "
                   "de compraventas, en 'Historial de posiciones y cuenta de margen' más arriba.")
        with st.form("form_agregar_dividendo"):
            c1, c2 = st.columns(2)
            fecha_div = c1.date_input("Fecha", value=date.today(), key="fecha_dividendo")
            tipo_div = c2.selectbox("Tipo", ["DIVIDEND", "INTEREST"], key="tipo_dividendo",
                                    format_func=lambda t: "Dividendo" if t == "DIVIDEND" else "Interés")
            c3, c4 = st.columns(2)
            plataforma_div = c3.text_input("Plataforma (ej. IBKR, Hapi)", key="plataforma_dividendo")
            moneda_div = c4.selectbox("Moneda", ["USD", "COP"], key="moneda_dividendo")
            activo_div = st.text_input("Activo (ticker; dejalo vacío si es interés general de la cuenta)",
                                       key="activo_dividendo")
            valor_div = st.number_input("Valor recibido", min_value=0.0, step=1.0, format="%.2f",
                                        key="valor_dividendo")
            guardar_div = st.form_submit_button("💾 Guardar", type="primary")
        if guardar_div:
            if valor_div <= 0:
                st.error("El valor tiene que ser mayor que cero.")
            elif not plataforma_div.strip():
                st.error("Escribí la plataforma.")
            else:
                etiqueta = "Dividendo" if tipo_div == "DIVIDEND" else "Interés"
                fila = [fecha_div.isoformat(), plataforma_div.strip(), moneda_div,
                        activo_div.strip() or "(general)", tipo_div, 0, 0, 0, valor_div, "Manual"]
                agregadas = db.agregar_historial_inversion([fila])
                db.clear_read_cache()
                if agregadas:
                    st.success(f"{etiqueta} de {valor_div:,.2f} {moneda_div} agregado.")
                else:
                    st.warning("Ya había un registro idéntico (misma fecha/plataforma/activo/tipo/valor) — "
                              "no se agregó de nuevo.")
                st.rerun()


def _plataforma_visible(nombre):
    """Nombre comercial mostrado sin alterar las claves históricas guardadas."""
    return "Trii / Acciones y Valores" if str(nombre).strip() == "Acciones y Valores" else nombre


def _deduplicar_posiciones_importadas(posiciones):
    """Consolida posiciones repetidas cuando dos reportes se solapan."""
    unicas = {}
    for posicion in posiciones:
        ticker = str(posicion.get("ticker") or "").strip()
        if not ticker:
            continue
        anterior = unicas.get(ticker)
        if anterior is None:
            unicas[ticker] = posicion.copy()
            continue
        combinada = anterior.copy()
        for campo in ("tipo", "moneda", "cantidad", "precio_compra"):
            if posicion.get(campo) not in (None, "", 0, 0.0):
                combinada[campo] = posicion[campo]
        if posicion.get("cotizacion_reportada") or not combinada.get("precio_actual"):
            combinada["precio_actual"] = posicion.get("precio_actual", combinada.get("precio_actual", 0))
        combinada["cotizacion_reportada"] = bool(
            combinada.get("cotizacion_reportada") or posicion.get("cotizacion_reportada"))
        unicas[ticker] = combinada
    return list(unicas.values())


def _fecha_compacta(valor):
    return datetime.strptime(str(valor), "%Y%m%d").date()


def _fecha_espanol(valor):
    meses = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
             "jul": 7, "ago": 8, "sept": 9, "sep": 9, "oct": 10, "nov": 11, "dic": 12}
    coincidencia = re.search(r"(\d{1,2})\s+([a-z]+)\s+(\d{4})", str(valor).lower())
    if not coincidencia or coincidencia.group(2) not in meses:
        raise ValueError(f"Fecha no reconocida: {valor}")
    return date(int(coincidencia.group(3)), meses[coincidencia.group(2)], int(coincidencia.group(1)))


def _contenido_csv(archivo):
    contenido = archivo.getvalue() if hasattr(archivo, "getvalue") else archivo.read()
    texto = contenido.decode("utf-8-sig") if isinstance(contenido, bytes) else str(contenido)
    lector = csv.DictReader(StringIO(texto))
    return list(lector), set(lector.fieldnames or [])


def _parse_portfolio_csv(archivo):
    filas, _ = _contenido_csv(archivo)
    tipos = {str(f.get("Transaction Type") or "").upper() for f in filas}
    simbolos = {str(f.get("Symbol") or "").upper() for f in filas}
    if tipos & {"SHORT", "COVER"}:
        plataforma = "Interactive Brokers"
    elif any(s.endswith("-USD") for s in simbolos):
        plataforma = "Binance"
    else:
        plataforma = "Hapi"
    operaciones, flujos = [], []
    fuente = getattr(archivo, "name", "portfolio.csv")
    for fila in filas:
        tipo = str(fila.get("Transaction Type") or "").upper()
        simbolo = str(fila.get("Symbol") or "").strip()
        fecha_txt = fila.get("Trade Date")
        if not fecha_txt:
            continue
        fecha = _fecha_compacta(fecha_txt)
        if simbolo == "$$CASH_TX":
            if tipo in {"DEPOSIT", "WITHDRAWAL", "WITHDRAW"}:
                es_retiro = tipo != "DEPOSIT"
                monto = abs(float(fila.get("Quantity") or 0))
                flujos.append(dict(fecha=fecha, monto_origen=-monto if es_retiro else monto,
                                   moneda_origen="USD", plataforma=plataforma,
                                   moneda_cuenta="dolares", tipo_flujo="Retiro" if es_retiro else "Depósito",
                                   fuente=fuente))
            continue
        if tipo in {"BUY", "SELL", "SHORT", "COVER"}:
            simbolo_operacion = simbolo[:-4] if plataforma == "Binance" and simbolo.endswith("-USD") else simbolo
            operaciones.append(dict(
                date=fecha, time=fila.get("Time"), symbol=simbolo_operacion, type=tipo,
                quantity=float(fila.get("Quantity") or 0),
                price=float(fila.get("Purchase Price") or 0), commission=float(fila.get("Commission") or 0),
                current_price=float(fila.get("Current Price") or 0), name=simbolo_operacion,
                source=fuente, moneda="USD"))
    posiciones, historial = resumen_operaciones_inversion(operaciones, plataforma)
    for posicion in posiciones:
        posicion["moneda"] = "dolares"
        posicion["cotizacion_reportada"] = True
        if plataforma == "Interactive Brokers":
            posicion["ticker"] = posicion["ticker"].replace("Interactive Brokers - ", "IBKR - ", 1)
        if plataforma == "Binance":
            posicion["tipo"] = "Cripto"
    return posiciones, flujos, historial


def _parse_binance_csv(archivo):
    filas, campos = _contenido_csv(archivo)
    if "ID de transacción (TXID)" in campos:
        return [], [], []

    flujos = []
    fuente = getattr(archivo, "name", "Binance.csv")
    for fila in filas:
        try:
            fecha = datetime.strptime(str(fila.get("Hora")), "%Y-%m-%d %H:%M:%S").date()
            monto = float(str(fila.get("Cambiar") or 0).replace(",", ""))
        except (TypeError, ValueError):
            continue
        operacion = str(fila.get("Operación") or "").strip()
        moneda = str(fila.get("Moneda") or "").strip().upper()
        if operacion == "Deposit" and moneda == "COP" and monto > 0:
            flujos.append(dict(fecha=fecha, monto_origen=monto, moneda_origen="COP",
                               plataforma="Binance", moneda_cuenta="dolares", tipo_flujo="Depósito",
                               fuente=fuente))
        elif operacion == "P2P Trading" and monto != 0:
            flujos.append(dict(fecha=fecha, monto_origen=monto, moneda_origen=moneda,
                               plataforma="Binance", moneda_cuenta="dolares",
                               tipo_flujo="Depósito" if monto > 0 else "Retiro", fuente=fuente))
        elif operacion in {"Withdraw", "Withdrawal"} and monto < 0:
            flujos.append(dict(fecha=fecha, monto_origen=monto, moneda_origen=moneda,
                               plataforma="Binance", moneda_cuenta="dolares", tipo_flujo="Retiro",
                               fuente=fuente))
        elif operacion == "Buy Crypto With Fiat" and monto > 0:
            # Este export de Binance omite el importe fiat. El activo y la fecha
            # permiten conciliarlo solo si hay una única salida bancaria ese día.
            flujos.append(dict(fecha=fecha, monto_origen=monto, moneda_origen=moneda,
                               plataforma="Binance", moneda_cuenta="dolares", tipo_flujo="Depósito",
                               fuente=fuente, importe_fiat_desconocido=True))

    return posiciones_binance_desde_movimientos(filas), flujos, []


def _parse_acciones_valores_csv(archivo):
    filas, campos = _contenido_csv(archivo)
    fuente = getattr(archivo, "name", "Información financiera.csv")
    if "Tipo de movimiento" in campos:
        flujos = []
        for fila in filas:
            tipo_movimiento = str(fila.get("Tipo de movimiento") or "").strip()
            if fila.get("Estado") != "Aprobado" or tipo_movimiento not in {"Depósito", "Retiro"}:
                continue
            es_retiro = tipo_movimiento == "Retiro"
            monto = abs(float(fila.get("Valor total") or 0))
            flujos.append(dict(fecha=_fecha_espanol(fila.get("Fecha y hora")),
                               monto_origen=-monto if es_retiro else monto, moneda_origen="COP",
                               plataforma="Acciones y Valores", moneda_cuenta="pesos",
                               tipo_flujo="Retiro" if es_retiro else "Depósito", fuente=fuente))
        return [], flujos, []

    operaciones = []
    for fila in filas:
        if fila.get("Estado") != "Aprobado":
            continue
        tipo = {"Compra": "BUY", "Venta": "SELL"}.get(fila.get("Tipo de orden"))
        if not tipo:
            continue
        cantidad_txt = str(fila.get("Acciones completadas") or "0").split("/", 1)[0]
        operaciones.append(dict(date=_fecha_espanol(fila.get("Fecha y hora")),
                                symbol=fila.get("Símbolo de la acción"), type=tipo,
                                quantity=float(cantidad_txt), price=float(fila.get("Precio por acción") or 0),
                                commission=float(fila.get("Valor comisión") or 0), moneda="COP",
                                source=fuente))
    posiciones, historial = resumen_operaciones_inversion(operaciones, "Acciones y Valores")
    for posicion in posiciones:
        posicion["moneda"] = "pesos"
    return posiciones, [], historial


def _parse_inversion_csv(archivo):
    _, campos = _contenido_csv(archivo)
    if {"Transaction Type", "Symbol", "Trade Date"}.issubset(campos):
        return _parse_portfolio_csv(archivo)
    if {"Operación", "Moneda", "Cambiar"}.issubset(campos) or "ID de transacción (TXID)" in campos:
        return _parse_binance_csv(archivo)
    if "Fecha y hora" in campos and ({"Tipo de orden", "Símbolo de la acción"}.issubset(campos)
                                      or "Tipo de movimiento" in campos):
        return _parse_acciones_valores_csv(archivo)
    raise ValueError(f"Formato CSV no reconocido: {getattr(archivo, 'name', 'archivo.csv')}")


def _parse_ibkr_pdf(archivo):
    operaciones, depositos, efectivo_final = [], [], None
    with pdfplumber.open(archivo) as pdf:
        for pagina in pdf.pages:
            for tabla in pagina.extract_tables():
                for fila in tabla:
                    if fila and str(fila[0] or "").strip() == "Efectivo final":
                        try:
                            ultimo_valor = next(v for v in reversed(fila) if v not in (None, ""))
                            efectivo_final = float(str(ultimo_valor).replace(",", ""))
                        except (StopIteration, TypeError, ValueError):
                            pass
                    if len(fila) < 11:
                        continue
                    try:
                        fecha = datetime.strptime(str(fila[0]), "%Y-%m-%d").date()
                    except (TypeError, ValueError):
                        continue
                    tipo = str(fila[3] or "").replace("\n", " ").strip()
                    if tipo in {"Buy", "Sell"}:
                        try:
                            operaciones.append(dict(
                                date=fecha, symbol=fila[4], type=tipo.upper(), quantity=float(fila[5]),
                                price=float(fila[6]), commission=float(fila[9] or 0),
                                name=str(fila[2] or "").replace("\n", " ")))
                        except (TypeError, ValueError):
                            continue
                    elif tipo == "Deposit":
                        try:
                            monto = float(fila[10])
                        except (TypeError, ValueError):
                            continue
                        depositos.append(dict(fecha=fecha, monto_origen=monto, moneda_origen="USD",
                                               plataforma="Interactive Brokers", moneda_cuenta="dolares",
                                               fuente=getattr(archivo, "name", "IBKR.pdf")))
    posiciones, historial = resumen_operaciones_inversion(operaciones, "Interactive Brokers")
    for posicion in posiciones:
        posicion["moneda"] = "dolares"
        posicion["ticker"] = posicion["ticker"].replace("Interactive Brokers - ", "IBKR - ", 1)
    if efectivo_final:
        posiciones.append(dict(ticker="IBKR - Efectivo/Margen", tipo="Otro", cantidad=1.0,
                                precio_compra=0.0, precio_actual=efectivo_final, moneda="dolares"))
    for deposito in depositos:
        deposito["tipo_flujo"] = "Depósito"
    return posiciones, depositos, historial


def _cruzar_flujos_1031(flujos, df_egresos, df_ingresos=None):
    """Empareja depósitos y retiros del broker con la cuenta 1031."""
    palabras = {"Hapi": ("MONO COLOMBIA",), "Interactive Brokers": ("SOLUCIONES DE PAGOS",),
                "Acciones y Valores": ("ACCIONES Y VAL",),
                "Binance": ("BINANCE", "MOVII", "COLOCA GRO")}
    movimientos = []
    for idx, fila in df_egresos.iterrows():
        fecha_raw = str(fila.get("Fecha Compra") or "").strip()
        try:
            fecha_mov = (date.fromisoformat(fecha_raw) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", fecha_raw)
                         else pd.to_datetime(fecha_raw, dayfirst=True, errors="raise").date())
        except (TypeError, ValueError):
            continue
        movimientos.append(dict(idx=("egreso", idx), fecha=fecha_mov,
                                concepto=str(fila.get("Comercio / Concepto") or ""),
                                cop=abs(float(fila.get("Valor Total Compra") or 0)), direccion="salida"))
    if df_ingresos is not None:
        for idx, fila in df_ingresos.iterrows():
            try:
                fecha_mov = pd.to_datetime(fila.get("Fecha"), dayfirst=True, errors="raise").date()
            except (TypeError, ValueError):
                continue
            movimientos.append(dict(idx=("ingreso", idx), fecha=fecha_mov,
                                    concepto=str(fila.get("Concepto") or ""),
                                    cop=abs(float(fila.get("Valor") or 0)), direccion="entrada"))
    usados, cruces = set(), []
    for flujo in sorted(flujos, key=lambda d: (d["fecha"], -abs(d["monto_origen"]))):
        es_retiro = flujo.get("tipo_flujo") == "Retiro" or flujo["monto_origen"] < 0
        candidatos = []
        for mov in movimientos:
            if mov["direccion"] != ("entrada" if es_retiro else "salida"):
                continue
            retraso = ((mov["fecha"] - flujo["fecha"]).days if es_retiro
                       else (flujo["fecha"] - mov["fecha"]).days)
            if mov["idx"] in usados or not 0 <= retraso <= 7:
                continue
            palabras_plataforma = palabras.get(flujo["plataforma"], ())
            concepto_upper = mov["concepto"].upper()
            if not (any(p in concepto_upper for p in palabras_plataforma)
                    or (es_retiro and "FONDO DE INVERS" in concepto_upper)):
                continue
            monto, moneda = abs(flujo["monto_origen"]), flujo["moneda_origen"]
            if moneda == "COP":
                diferencia = abs(mov["cop"] - monto)
                if diferencia <= max(1, monto * 0.001):
                    candidatos.append((retraso + diferencia / max(monto, 1), mov, 1.0))
            elif flujo.get("importe_fiat_desconocido"):
                if retraso == 0:
                    candidatos.append((0, mov, 0.0))
            elif monto >= 10:
                tasa = mov["cop"] / monto
                if 2500 <= tasa <= 5500:
                    candidatos.append((retraso + abs(tasa - 3800) / 1000, mov, tasa))
        if flujo.get("importe_fiat_desconocido") and len(candidatos) != 1:
            candidatos = []
        if candidatos:
            _, mov, tasa = min(candidatos, key=lambda x: x[0])
            usados.add(mov["idx"])
            cruces.append({**flujo, "fecha_1031": mov["fecha"], "concepto_1031": mov["concepto"],
                           "cop": -mov["cop"] if es_retiro else mov["cop"],
                           "tasa_implicita": tasa, "estado": "Conciliado"})
        else:
            lado = "entrada" if es_retiro else "salida"
            cruces.append({**flujo, "fecha_1031": None, "concepto_1031": "", "cop": 0,
                           "tasa_implicita": 0, "estado": f"Sin {lado} 1031 identificada"})
    return cruces


def _cruzar_depositos_1031(depositos, df_egresos):
    """Compatibilidad con llamadas anteriores: cruza solo depósitos."""
    return _cruzar_flujos_1031(depositos, df_egresos)


def _render_importar_portafolios():
    with st.expander("Importar portafolios y conciliar flujos con la cuenta 1031"):
        st.caption("Admite CSV de Hapi, Binance, Interactive Brokers y Acciones y Valores, además del PDF "
                   "Transaction History de Interactive Brokers. Reconoce compras, ventas, posiciones en corto, "
                   "coberturas, depósitos y retiros.")
        c1, c2 = st.columns(2)
        csv_files = c1.file_uploader("Archivos CSV de inversión", type=["csv"], key="inversiones_csv",
                                     accept_multiple_files=True)
        pdf_file = c2.file_uploader("Interactive Brokers PDF", type=["pdf"], key="inversiones_ibkr_pdf")
        if not csv_files and not pdf_file:
            return
        posiciones, flujos, historial = [], [], []
        try:
            for csv_file in csv_files:
                pos, mov_efectivo, hist = _parse_inversion_csv(csv_file)
                posiciones += pos
                flujos += mov_efectivo
                historial += hist
            if pdf_file:
                pos, mov_efectivo, hist = _parse_ibkr_pdf(pdf_file)
                posiciones += pos
                flujos += mov_efectivo
                historial += hist
        except Exception as e:
            st.error(f"No pude interpretar los archivos: {e}")
            return

        posiciones = _deduplicar_posiciones_importadas(posiciones)
        df_pos = pd.DataFrame(posiciones)
        if not df_pos.empty:
            st.markdown(f"**Posiciones abiertas detectadas: {len(df_pos)}**")
            st.dataframe(df_pos.rename(columns={"ticker": "Posición", "tipo": "Tipo", "cantidad": "Cantidad",
                                                "precio_compra": "Precio compra promedio",
                                                "precio_actual": "Precio actual/provisional"}),
                         hide_index=True, use_container_width=True)

        df_historial = pd.DataFrame(historial)
        if not df_historial.empty:
            resumen = (df_historial.groupby(["plataforma", "activo", "estado"], as_index=False)
                       .agg(Operaciones=("operacion", "size"),
                            **{"Resultado realizado": ("resultado_realizado", "sum")}))
            st.markdown("**Posiciones abiertas, cortas y cerradas detectadas**")
            st.dataframe(resumen.rename(columns={"plataforma": "Plataforma", "activo": "Activo",
                                                  "estado": "Estado"}),
                         hide_index=True, use_container_width=True)

        df_1031 = db.read_egreso_detalle("efectivo_detalle")
        df_ingresos_1031 = db.read_otros_ingresos_tabla()
        cruces = _cruzar_flujos_1031(flujos, df_1031, df_ingresos_1031)
        flujos_guardables = []
        for cruce in cruces:
            if cruce["estado"] == "Conciliado":
                flujos_guardables.append(cruce)
            elif cruce["moneda_origen"] == "COP":
                # Si el reporte de la plataforma trae el importe exacto en COP,
                # sigue siendo una fuente válida aunque el texto bancario no haya
                # permitido emparejarlo automáticamente.
                es_retiro = cruce["tipo_flujo"] == "Retiro" or cruce["monto_origen"] < 0
                flujos_guardables.append({**cruce, "fecha_1031": cruce["fecha"],
                                          "cop": -abs(cruce["monto_origen"]) if es_retiro
                                          else abs(cruce["monto_origen"])})
        df_cruces = pd.DataFrame(cruces)
        if not df_cruces.empty:
            df_cruces["COP a registrar"] = df_cruces.apply(
                lambda fila: (fila["cop"] if fila["estado"] == "Conciliado" else
                              (-abs(fila["monto_origen"]) if fila["tipo_flujo"] == "Retiro"
                               else abs(fila["monto_origen"]))
                              if fila["moneda_origen"] == "COP" else 0), axis=1)
            vista = df_cruces[["tipo_flujo", "plataforma", "fecha", "monto_origen", "moneda_origen", "fecha_1031", "cop",
                               "COP a registrar", "tasa_implicita", "estado"]].copy()
            vista.columns = ["Flujo", "Plataforma", "Fecha broker", "Monto origen", "Moneda",
                             "Fecha cuenta 1031", "COP conciliado", "COP a registrar",
                             "Tasa implícita", "Estado"]
            st.markdown("**Cruce de depósitos y retiros con la cuenta 1031**")
            st.dataframe(vista, hide_index=True, use_container_width=True)

        if st.button("Guardar posiciones, historial y flujos", type="primary",
                     key="guardar_importacion_inversiones"):
            prefijos_importados = {
                "IBKR" if h["plataforma"] == "Interactive Brokers" else h["plataforma"]
                for h in historial
            }
            for moneda in ("pesos", "dolares"):
                nuevas = [p for p in posiciones if p.get("moneda", "dolares") == moneda]
                prefijos_moneda = {p["ticker"].split(" - ", 1)[0] for p in nuevas}
                if moneda == "dolares":
                    prefijos_moneda |= prefijos_importados & {"IBKR", "Binance", "Hapi"}
                else:
                    prefijos_moneda |= prefijos_importados & {"Acciones y Valores", "Trii"}
                if not nuevas and not prefijos_moneda:
                    continue
                existentes = db.read_posiciones_inversion(moneda)
                combinadas = {}
                for _, fila in existentes.iterrows():
                    ticker = str(fila.get("Ticker / Fondo") or "").strip()
                    if ticker:
                        prefijo = ticker.split(" - ", 1)[0]
                        es_efectivo_margen = ticker.endswith("Efectivo/Margen")
                        if prefijo in prefijos_moneda and not es_efectivo_margen:
                            continue
                        combinadas[ticker] = dict(ticker=ticker, tipo=fila.get("Tipo") or "",
                                                  cantidad=_num_o_cero(fila.get("Cantidad")),
                                                  precio_compra=_num_o_cero(fila.get("Precio Compra Promedio")),
                                                  precio_actual=_num_o_cero(fila.get("Precio Actual")))
                for posicion in nuevas:
                    anterior = combinadas.get(posicion["ticker"])
                    if anterior:
                        for campo in ("precio_compra", "precio_actual"):
                            if not posicion.get(campo) and anterior.get(campo):
                                posicion[campo] = anterior[campo]
                    combinadas[posicion["ticker"]] = posicion
                db.set_posiciones_inversion(moneda, list(combinadas.values()))
            agregados = 0
            for moneda in ("pesos", "dolares"):
                filas_aporte = [[c["fecha_1031"].isoformat(), c["plataforma"], c["cop"],
                                 f"{c['tipo_flujo']} {abs(c['monto_origen']):.8g} {c['moneda_origen']} el "
                                 f"{c['fecha'].isoformat()} — {c['fuente']}"]
                                for c in flujos_guardables if c["moneda_cuenta"] == moneda]
                agregados += db.agregar_aportes_inversion(moneda, filas_aporte)
            filas_historial = [[h["fecha"].isoformat(), h["plataforma"], h["moneda"], h["activo"],
                                h["operacion"], h["cantidad"], h["precio"], h["comision"],
                                h["resultado_realizado"], h["fuente"]] for h in historial]
            operaciones_agregadas = db.agregar_historial_inversion(filas_historial)
            db.recategorizar_comercios({"PAGO PSE MONO COLOMBIA SAS": "Inversiones",
                                        "PAGO PSE SOLUCIONES DE PAGOS": "Inversiones",
                                        "PAGO PSE ACCIONES Y VALORES S": "Inversiones",
                                        "PAGO PSE ACCIONES Y VALORES": "Inversiones"}, solo_otros=False)
            db.clear_read_cache()
            for moneda in ("pesos", "dolares"):
                _tomar_snapshot_cartera(moneda)
            st.success(f"Guardadas {len(posiciones)} posiciones, {operaciones_agregadas} operaciones históricas "
                       f"y {agregados} flujos de efectivo nuevos.")
            st.rerun()


def _form_posiciones_inversion(moneda, key_sufijo):
    columnas = ["Ticker / Fondo", "Tipo", "Cantidad", "Precio Compra Promedio",
                "Costo Total", "Precio Actual", "Valor Actual", "Ganancia/Pérdida"]
    df_pos = db.read_posiciones_inversion(moneda)
    if df_pos.empty:
        df_pos = pd.DataFrame(columns=columnas)
    df_editado = st.data_editor(
        df_pos, hide_index=True, use_container_width=True, num_rows="dynamic",
        height=min(400, 80 + 35 * (len(df_pos) + 1)), key=f"editor_posiciones_{key_sufijo}",
        disabled=["Costo Total", "Valor Actual", "Ganancia/Pérdida"],
        column_config={
            "Tipo": st.column_config.SelectboxColumn(
                options=["Acción", "ETF", "Fondo de Inversión", "Cripto", "Otro"]),
            "Cantidad": st.column_config.NumberColumn(format="%.4f"),
            "Precio Compra Promedio": st.column_config.NumberColumn(format="$%d"),
            "Precio Actual": st.column_config.NumberColumn(format="$%d"),
            "Costo Total": st.column_config.NumberColumn(format="$%d"),
            "Valor Actual": st.column_config.NumberColumn(format="$%d"),
            "Ganancia/Pérdida": st.column_config.NumberColumn(format="$%d"),
        })
    if st.button("💾 Guardar posiciones", key=f"btn_guardar_posiciones_{key_sufijo}"):
        filas = []
        for _, fila in df_editado.iterrows():
            ticker = str(fila.get("Ticker / Fondo") or "").strip()
            if not ticker:
                continue
            filas.append(dict(
                ticker=ticker, tipo=fila.get("Tipo") or "",
                cantidad=_num_o_cero(fila.get("Cantidad")),
                precio_compra=_num_o_cero(fila.get("Precio Compra Promedio")),
                precio_actual=_num_o_cero(fila.get("Precio Actual"))))
        try:
            db.set_posiciones_inversion(moneda, filas)
            db.clear_read_cache()
            _tomar_snapshot_cartera(moneda)
            st.success(f"{len(filas)} posición(es) guardadas.")
            st.rerun()
        except db.SinEspacioError as e:
            st.error(str(e))
        except Exception as e:
            st.error(f"No pude guardar: {e}")


def _serie_cierre(datos, simbolo, total_simbolos):
    if datos.empty:
        return pd.Series(dtype=float)
    if isinstance(datos.columns, pd.MultiIndex):
        if simbolo in datos.columns.get_level_values(0):
            serie = datos[simbolo]["Close"]
        elif "Close" in datos.columns.get_level_values(0):
            serie = datos["Close"][simbolo]
        else:
            return pd.Series(dtype=float)
    elif total_simbolos == 1 and "Close" in datos:
        serie = datos["Close"]
    else:
        return pd.Series(dtype=float)
    return pd.to_numeric(serie, errors="coerce").dropna()


@st.cache_data(ttl=900, show_spinner=False)
def _descargar_cotizaciones(simbolos):
    if yf is None:
        raise RuntimeError("Falta instalar yfinance para consultar precios.")
    simbolos = tuple(sorted(set(simbolos)))
    if not simbolos:
        return {}
    datos = yf.download(list(simbolos), period="5d", interval="1d", group_by="ticker",
                        auto_adjust=False, progress=False, threads=True)
    resultado = {}
    for simbolo in simbolos:
        serie = _serie_cierre(datos, simbolo, len(simbolos))
        if not serie.empty:
            resultado[simbolo] = {"precio": float(serie.iloc[-1]),
                                  "fecha": pd.Timestamp(serie.index[-1]).date().isoformat()}
    return resultado


@st.cache_data(ttl=900, show_spinner=False)
def _descargar_trm():
    """Tasa de cambio USD/COP (Yahoo Finance 'COP=X') para poder unificar el
    patrimonio de las hojas de pesos y dólares en un solo número — es
    informativa (cierre del día, no la tasa exacta de una transferencia
    real), cacheada 15 minutos igual que las cotizaciones de posiciones."""
    if yf is None:
        return None
    try:
        datos = yf.download("COP=X", period="5d", interval="1d", progress=False)
    except Exception:
        return None
    serie = _serie_cierre(datos, "COP=X", 1)
    return float(serie.iloc[-1]) if not serie.empty else None


@st.cache_data(ttl=3600, show_spinner=False)
def _descargar_historico_cierre(simbolo, fecha_inicio):
    """Serie de cierres diarios desde 'fecha_inicio' hasta hoy — para el
    benchmark, se cachea 1 hora (no hace falta la frescura de 15 min de un
    precio para operar). Serie vacía si el símbolo no existe/no hay datos."""
    if yf is None:
        return pd.Series(dtype=float)
    try:
        datos = yf.download(simbolo, start=fecha_inicio.isoformat(), progress=False)
    except Exception:
        return pd.Series(dtype=float)
    serie = _serie_cierre(datos, simbolo, 1)
    if not serie.empty:
        serie.index = pd.to_datetime(serie.index).tz_localize(None)
    return serie


BENCHMARKS = {"pesos": ("ICOLCAP.CL", "COLCAP"), "dolares": ("^GSPC", "S&P 500")}


def _valor_shadow_benchmark(moneda):
    """Si hubieras metido cada aporte/retiro real (misma fecha, mismo
    monto) a un fondo que replicara el benchmark en vez de a tu cartera,
    cuánto tendrías hoy — para comparar contra el valor real de la
    cartera. Para dólares, los aportes están en pesos transferidos, así
    que además convierte cada uno a USD con la TRM histórica de ESE día
    (no la de hoy) antes de comprarlo al precio del benchmark ese día."""
    simbolo, nombre = BENCHMARKS.get(moneda, (None, None))
    if not simbolo:
        return None, None
    df_ap = db.read_aportes_inversion(moneda)
    if df_ap.empty:
        return None, None
    flujos = []
    for _, fila in df_ap.iterrows():
        fecha = pd.to_datetime(fila["Fecha"], dayfirst=True, errors="coerce")
        if pd.isna(fecha):
            continue
        monto = _num_o_cero(fila["Monto Transferido (COP)"])
        if monto:
            flujos.append((fecha.date(), monto))
    if not flujos:
        return None, None
    fecha_inicio = min(f for f, _ in flujos)
    serie_precio = _descargar_historico_cierre(simbolo, fecha_inicio)
    if serie_precio.empty:
        return None, None
    serie_fx = _descargar_historico_cierre("COP=X", fecha_inicio) if moneda == "dolares" else None
    if moneda == "dolares" and serie_fx.empty:
        return None, None

    def _precio_en(serie, fecha):
        hasta = serie[serie.index <= pd.Timestamp(fecha)]
        return float(hasta.iloc[-1]) if not hasta.empty else float(serie.iloc[0])

    unidades = 0.0
    for fecha, monto in flujos:
        monto_en_moneda_benchmark = monto
        if serie_fx is not None:
            fx = _precio_en(serie_fx, fecha)
            monto_en_moneda_benchmark = monto / fx if fx else 0.0
        precio = _precio_en(serie_precio, fecha)
        if precio:
            unidades += monto_en_moneda_benchmark / precio
    return unidades * float(serie_precio.iloc[-1]), nombre


def _patrimonio_total_inversiones(moneda):
    df_pos = db.read_posiciones_inversion(moneda)
    if df_pos.empty:
        return 0.0
    cantidad = pd.to_numeric(df_pos["Cantidad"], errors="coerce").fillna(0.0)
    precio_actual = pd.to_numeric(df_pos["Precio Actual"], errors="coerce").fillna(0.0)
    return float((cantidad * precio_actual).sum())


def _render_patrimonio_unificado():
    st.markdown("#### 🌎 Patrimonio total en inversiones")
    patrimonio_pesos = _patrimonio_total_inversiones("pesos")
    patrimonio_dolares = _patrimonio_total_inversiones("dolares")
    trm = _descargar_trm()
    if trm is None:
        st.info("No pude consultar la TRM (USD/COP) desde Yahoo Finance ahora mismo — mostrando cada moneda "
                "por separado, sin unificar.")
        c1, c2 = st.columns(2)
        c1.metric("Pesos (COP)", f"${patrimonio_pesos:,.0f}")
        c2.metric("Dólares (USD)", f"${patrimonio_dolares:,.2f}")
        return
    patrimonio_dolares_cop = patrimonio_dolares * trm
    total = patrimonio_pesos + patrimonio_dolares_cop
    c1, c2, c3 = st.columns(3)
    c1.metric("Pesos (COP)", f"${patrimonio_pesos:,.0f}")
    c2.metric(f"Dólares → COP (TRM ${trm:,.0f})", f"${patrimonio_dolares_cop:,.0f}")
    c3.metric("Total en inversiones (COP)", f"${total:,.0f}")
    if total > 0:
        fig = px.pie(pd.DataFrame({"Moneda": ["Pesos", "Dólares (convertido)"],
                                   "Valor": [patrimonio_pesos, patrimonio_dolares_cop]}),
                    names="Moneda", values="Valor", title="Distribución por moneda (en COP)",
                    color_discrete_sequence=["#1d4ed8", "#0d9488"])
        fig.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=300)
        st.plotly_chart(fig, use_container_width=True, key="patrimonio_unificado_pie")
    st.caption(f"TRM ${trm:,.2f} COP/USD, cierre más reciente en Yahoo Finance (cacheada 15 min) — "
               "informativa, no la tasa exacta de una transferencia real.")


def _filas_posiciones(df_pos):
    filas = []
    for _, fila in df_pos.iterrows():
        ticker = str(fila.get("Ticker / Fondo") or "").strip()
        if ticker:
            filas.append(dict(ticker=ticker, tipo=fila.get("Tipo") or "",
                              cantidad=_num_o_cero(fila.get("Cantidad")),
                              precio_compra=_num_o_cero(fila.get("Precio Compra Promedio")),
                              precio_actual=_num_o_cero(fila.get("Precio Actual"))))
    return filas


def _render_actualizar_precios():
    posiciones = {moneda: db.read_posiciones_inversion(moneda) for moneda in ("pesos", "dolares")}
    referencias = {}
    for moneda, df_pos in posiciones.items():
        for idx, fila in df_pos.iterrows():
            simbolo = simbolo_cotizacion(fila.get("Ticker / Fondo"), moneda, fila.get("Tipo"))
            if simbolo:
                referencias[(moneda, idx)] = simbolo

    st.markdown("#### Precios y seguimiento")
    st.caption("Consulta el último cierre disponible en Yahoo Finance. Solo se envían símbolos públicos; "
               "cantidades, costos y saldos permanecen en la app. Las cotizaciones se conservan 15 minutos.")
    if mensaje := st.session_state.get("ultima_actualizacion_mercado"):
        st.success(mensaje)
    if st.button("Actualizar precios de mercado", type="primary", key="actualizar_precios_mercado",
                 disabled=not referencias):
        try:
            with st.spinner("Consultando cotizaciones..."):
                cotizaciones = _descargar_cotizaciones(tuple(referencias.values()))
            actualizadas, faltantes, fechas = 0, [], []
            for moneda, df_pos in posiciones.items():
                filas = _filas_posiciones(df_pos)
                for fila in filas:
                    simbolo = simbolo_cotizacion(fila["ticker"], moneda, fila["tipo"])
                    cotizacion = cotizaciones.get(simbolo)
                    if cotizacion:
                        fila["precio_actual"] = cotizacion["precio"]
                        actualizadas += 1
                        fechas.append(cotizacion["fecha"])
                    elif simbolo:
                        faltantes.append(simbolo)
                if filas:
                    db.set_posiciones_inversion(moneda, filas)
            db.clear_read_cache()
            for moneda in ("pesos", "dolares"):
                _tomar_snapshot_cartera(moneda)
            fechas_unicas = sorted(set(fechas))
            if len(fechas_unicas) > 1:
                fecha_texto = f"{fechas_unicas[0]} a {fechas_unicas[-1]}"
            else:
                fecha_texto = fechas_unicas[0] if fechas_unicas else "sin fecha disponible"
            mensaje = f"Actualicé {actualizadas} posiciones. Fechas de cotización: {fecha_texto}."
            if faltantes:
                mensaje += f" Sin cotización: {', '.join(sorted(set(faltantes)))}."
            st.session_state["ultima_actualizacion_mercado"] = mensaje
            st.rerun()
        except Exception as e:
            st.error(f"No pude actualizar las cotizaciones: {e}")


def _vista_consolidada_posiciones(df_pos):
    """Agrupa el mismo activo comprado en brokers distintos (p. ej. ADBE en
    Hapi e IBKR) y calcula cantidad total y costo promedio ponderado entre
    las compras — una vista de análisis; las filas de Posiciones siguen
    siendo por broker, no se tocan."""
    if df_pos.empty:
        return pd.DataFrame()
    tmp = df_pos.copy()
    for columna in ("Cantidad", "Precio Compra Promedio", "Precio Actual"):
        tmp[columna] = pd.to_numeric(tmp[columna], errors="coerce").fillna(0.0)
    tmp = tmp[~tmp["Ticker / Fondo"].astype(str).str.contains(r"Efectivo/Margen$", case=False, regex=True)]
    if tmp.empty:
        return pd.DataFrame()
    partes = tmp["Ticker / Fondo"].astype(str).str.split(" - ", n=1, expand=True)
    tmp["Plataforma"] = partes[0].map(_plataforma_visible)
    tmp["Activo"] = partes[1] if 1 in partes.columns else partes[0]
    tmp["Costo"] = tmp["Cantidad"].abs() * tmp["Precio Compra Promedio"]
    tmp["Valor"] = tmp["Cantidad"] * tmp["Precio Actual"]
    filas = []
    for activo, grupo in tmp.groupby("Activo"):
        cantidad_abs = grupo["Cantidad"].abs().sum()
        costo = grupo["Costo"].sum()
        valor = grupo["Valor"].sum()
        filas.append({
            "Activo": activo, "Plataformas": ", ".join(sorted(grupo["Plataforma"].unique())),
            "N.º brokers": grupo["Plataforma"].nunique(), "Cantidad": grupo["Cantidad"].sum(),
            "Precio Compra Promedio Ponderado": costo / cantidad_abs if cantidad_abs else 0.0,
            "Costo Total": costo, "Valor Actual": valor, "Ganancia/Pérdida": valor - costo,
        })
    return pd.DataFrame(filas).sort_values("Valor Actual", ascending=False)


def _render_vista_consolidada(df_pos, moneda, key_sufijo):
    consolidado = _vista_consolidada_posiciones(df_pos)
    if consolidado.empty:
        return
    multi_broker = consolidado[consolidado["N.º brokers"] > 1]
    with st.expander(f"🔀 Vista consolidada por activo ({len(consolidado)} activo(s) único(s))"):
        st.caption("Agrupa el mismo activo comprado en brokers distintos (p. ej. ADBE en Hapi e IBKR) y "
                   "calcula cantidad total y precio de compra promedio ponderado entre las compras. Es solo "
                   "una vista de análisis — las filas de 'Posiciones' de arriba siguen siendo por broker, así "
                   "que editar o vender ahí no se afecta por esta tabla.")
        if not multi_broker.empty:
            st.info(f"**{len(multi_broker)} activo(s) repetidos entre brokers:** "
                   + ", ".join(multi_broker["Activo"].tolist()))
        vista = consolidado.copy()
        for col in ("Precio Compra Promedio Ponderado", "Costo Total", "Valor Actual", "Ganancia/Pérdida"):
            vista[col] = vista[col].map(lambda v: f"${v:,.2f}" if moneda == "dolares" else f"${v:,.0f}")
        vista["Cantidad"] = vista["Cantidad"].map(lambda v: f"{v:,.4f}")
        st.dataframe(vista, hide_index=True, use_container_width=True,
                     height=min(450, 80 + 35 * len(vista)))


def _render_resumen_cartera(df_pos, moneda, key_sufijo):
    if df_pos.empty:
        return pd.DataFrame(), 0.0
    cartera = df_pos.copy()
    for columna in ("Cantidad", "Precio Compra Promedio", "Precio Actual"):
        cartera[columna] = pd.to_numeric(cartera[columna], errors="coerce").fillna(0.0)
    cartera["Costo calculado"] = cartera["Cantidad"].abs() * cartera["Precio Compra Promedio"]
    cartera["Valor calculado"] = cartera["Cantidad"] * cartera["Precio Actual"]
    cartera["Plataforma"] = (cartera["Ticker / Fondo"].astype(str).str.split(" - ", n=1).str[0]
                             .map(_plataforma_visible))
    cartera["Es efectivo/margen"] = cartera["Ticker / Fondo"].astype(str).str.contains(
        r"Efectivo/Margen$", case=False, regex=True)
    con_costo = cartera["Costo calculado"] > 0
    cartera["Ganancia calculada"] = ((cartera["Precio Actual"] - cartera["Precio Compra Promedio"])
                                      * cartera["Cantidad"]).where(con_costo)
    cartera["Rentabilidad"] = (cartera["Ganancia calculada"] / cartera["Costo calculado"]).where(con_costo)
    titulos = cartera[~cartera["Es efectivo/margen"]].copy()
    costo = titulos.loc[titulos["Costo calculado"] > 0, "Costo calculado"].sum()
    valor_titulos = titulos["Valor calculado"].sum()
    efectivo_margen = cartera.loc[cartera["Es efectivo/margen"], "Valor calculado"].sum()
    valor = valor_titulos + efectivo_margen
    exposicion = titulos["Valor calculado"].abs().sum()
    ganancia = cartera["Ganancia calculada"].sum()
    rentabilidad = ganancia / costo if costo else 0.0
    unidad = "COP" if moneda == "pesos" else "USD"
    c1, c2, c3 = st.columns(3)
    c1.metric(f"Valor de títulos ({unidad})", f"${valor_titulos:,.2f}")
    c2.metric(f"Efectivo / deuda de margen ({unidad})", f"${efectivo_margen:,.2f}")
    c3.metric(f"Patrimonio propio estimado ({unidad})", f"${valor:,.2f}")
    c4, c5, c6 = st.columns(3)
    c4.metric(f"Base de costo ({unidad})", f"${costo:,.2f}")
    c5.metric(f"Resultado no realizado ({unidad})", f"${ganancia:,.2f}")
    c6.metric("Rentabilidad sobre costo", f"{rentabilidad:.2%}")
    st.caption(f"Exposición bruta en títulos: ${exposicion:,.2f} {unidad}. "
               "Un saldo de margen negativo es financiación del broker: reduce el patrimonio, pero no es una pérdida.")
    ibkr = cartera[cartera["Plataforma"].eq("IBKR")]
    if not ibkr.empty:
        titulos_ibkr = ibkr.loc[~ibkr["Es efectivo/margen"], "Valor calculado"].sum()
        margen_ibkr = ibkr.loc[ibkr["Es efectivo/margen"], "Valor calculado"].sum()
        patrimonio_ibkr = titulos_ibkr + margen_ibkr
        apalancamiento = titulos_ibkr / patrimonio_ibkr if patrimonio_ibkr > 0 else 0.0
        st.info(f"IBKR: ${titulos_ibkr:,.2f} en títulos {margen_ibkr:+,.2f} de efectivo/margen = "
                f"${patrimonio_ibkr:,.2f} de patrimonio propio. Apalancamiento: {apalancamiento:.2f}x. "
                "El patrimonio no equivale a la ganancia.")
    sin_costo = cartera.loc[(cartera["Costo calculado"] <= 0) & (cartera["Valor calculado"] != 0),
                            "Ticker / Fondo"].tolist()
    if sin_costo:
        st.warning("No incluí en ganancia/rentabilidad las posiciones sin costo informado: "
                   + ", ".join(sin_costo))

    UMBRAL_CONCENTRACION = 0.20
    if exposicion > 0:
        concentracion = titulos[titulos["Valor calculado"] != 0].copy()
        concentracion["% de la exposición"] = concentracion["Valor calculado"].abs() / exposicion
        riesgosas = concentracion[concentracion["% de la exposición"] > UMBRAL_CONCENTRACION].sort_values(
            "% de la exposición", ascending=False)
        if not riesgosas.empty:
            detalle = ", ".join(f"{fila['Ticker / Fondo']} ({fila['% de la exposición']:.0%})"
                                for _, fila in riesgosas.iterrows())
            st.warning(f"⚠️ **Concentración**: {len(riesgosas)} posición(es) superan el "
                       f"{UMBRAL_CONCENTRACION:.0%} de tu exposición en títulos — {detalle}. Un solo activo "
                       "pesando tanto amplifica el riesgo de esa posición sobre toda la cartera.")
    por_plataforma = (cartera.groupby("Plataforma", as_index=False)
                      .agg(**{"Base de costo": ("Costo calculado", "sum"),
                              "Valor de títulos": ("Valor calculado", lambda s: s[~cartera.loc[s.index, "Es efectivo/margen"]].sum()),
                              "Efectivo / deuda de margen": ("Valor calculado", lambda s: s[cartera.loc[s.index, "Es efectivo/margen"]].sum()),
                              "Resultado no realizado": ("Ganancia calculada", "sum")}))
    por_plataforma["Patrimonio estimado"] = (por_plataforma["Valor de títulos"]
                                               + por_plataforma["Efectivo / deuda de margen"])
    st.markdown("**Resumen por plataforma**")
    st.dataframe(por_plataforma[["Plataforma", "Base de costo", "Valor de títulos", "Efectivo / deuda de margen",
                                 "Patrimonio estimado", "Resultado no realizado"]],
                 hide_index=True, use_container_width=True,
                 column_config={col: st.column_config.NumberColumn(col, format="$ %.2f") for col in
                                ["Base de costo", "Valor de títulos", "Efectivo / deuda de margen",
                                 "Patrimonio estimado", "Resultado no realizado"]})

    if moneda == "pesos" and cartera["Plataforma"].eq("Trii / Acciones y Valores").any():
        trii = cartera[cartera["Plataforma"].eq("Trii / Acciones y Valores")]
        resumen_trii = (trii.groupby("Tipo", as_index=False)
                        .agg(**{"Base de costo": ("Costo calculado", "sum"),
                                "Valor actual": ("Valor calculado", "sum")}))
        st.markdown("**Trii: acciones y fondos**")
        st.dataframe(resumen_trii, hide_index=True, use_container_width=True,
                     column_config={col: st.column_config.NumberColumn(col, format="$ %.0f")
                                    for col in ["Base de costo", "Valor actual"]})
        if not trii["Tipo"].astype(str).eq("Fondo de Inversión").any():
            st.warning("Los reportes de Trii cargados contienen las órdenes de acciones, pero no informan el "
                       "saldo actual de fondos. Agrégalo abajo como tipo 'Fondo de Inversión'; no lo estimo "
                       "desde transferencias porque esas cifras no incluyen rendimientos ni dividendos.")

    graficables = titulos[titulos["Valor calculado"] != 0].copy()
    if not graficables.empty:
        graficables["Lado"] = graficables["Cantidad"].apply(lambda q: "Largo" if q >= 0 else "Corto")
        graficables = graficables.sort_values("Valor calculado")
        g1, g2 = st.columns(2)
        fig_asignacion = px.bar(graficables, x="Valor calculado", y="Ticker / Fondo", color="Lado",
                                orientation="h", title=f"Valor de mercado por posición ({unidad})",
                                color_discrete_map={"Largo": "#1d4ed8", "Corto": "#dc2626"})
        fig_asignacion.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=520,
                                     xaxis_title=unidad, yaxis_title="", legend_title="Posición")
        g1.plotly_chart(fig_asignacion, use_container_width=True, key=f"asignacion_{key_sufijo}")
        resultados = cartera.dropna(subset=["Ganancia calculada"]).copy()
        if not resultados.empty:
            resultados = resultados.sort_values("Ganancia calculada")
            resultados["Resultado"] = resultados["Ganancia calculada"].apply(
                lambda valor_g: "Ganancia" if valor_g >= 0 else "Pérdida")
            fig_resultado = px.bar(resultados, x="Ganancia calculada", y="Ticker / Fondo", color="Resultado",
                                   orientation="h",
                                   color_discrete_map={"Ganancia": "#0d9488", "Pérdida": "#dc2626"},
                                   title=f"Resultado no realizado por posición ({unidad})")
            fig_resultado.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=520,
                                        xaxis_title=unidad, yaxis_title="", showlegend=False)
            g2.plotly_chart(fig_resultado, use_container_width=True, key=f"resultado_{key_sufijo}")

    return titulos, exposicion


def _render_metas_asignacion(titulos, exposicion, moneda, key_sufijo):
    """Metas de asignación por Tipo de activo (Acción/ETF/Cripto/...) vs. lo
    que de verdad tenés hoy — sobre la exposición en títulos (sin efectivo/
    margen, igual que el resto de las métricas de esta página)."""
    st.markdown("##### 🎯 Metas de asignación")
    if exposicion <= 0:
        st.caption("Necesitás posiciones con valor para comparar contra una meta.")
        return
    real = (titulos[titulos["Valor calculado"] != 0].groupby("Tipo")["Valor calculado"]
            .apply(lambda s: s.abs().sum() / exposicion).reset_index(name="Real %"))

    metas_todas = db.read_metas_asignacion()
    metas_moneda = metas_todas[metas_todas["Moneda"] == moneda] if not metas_todas.empty else metas_todas
    tipos_existentes = ["Acción", "ETF", "Fondo de Inversión", "Cripto", "Otro"]
    tipos = sorted(set(tipos_existentes) | set(real["Tipo"]) | set(metas_moneda["Tipo"] if not metas_moneda.empty else []))
    metas_por_tipo = dict(zip(metas_moneda["Tipo"], metas_moneda["Meta %"])) if not metas_moneda.empty else {}
    real_por_tipo = dict(zip(real["Tipo"], real["Real %"]))

    df_editor = pd.DataFrame({"Tipo": tipos,
                              "Meta %": [round(metas_por_tipo.get(t, 0.0) * 100, 1) for t in tipos],
                              "Real %": [round(real_por_tipo.get(t, 0.0) * 100, 1) for t in tipos]})
    with st.expander("Editar metas por tipo de activo"):
        st.caption("Poné el % objetivo para cada tipo (deben sumar 100 si querés cubrir toda la cartera; "
                   "podés dejar algunos en 0). 'Real %' es de referencia, no se guarda.")
        df_editado = st.data_editor(
            df_editor, hide_index=True, use_container_width=True, key=f"editor_metas_{key_sufijo}",
            disabled=["Tipo", "Real %"],
            column_config={"Meta %": st.column_config.NumberColumn(format="%.1f"),
                           "Real %": st.column_config.NumberColumn(format="%.1f")})
        if st.button("💾 Guardar metas", key=f"btn_guardar_metas_{key_sufijo}"):
            filas = [[moneda, fila["Tipo"], float(fila["Meta %"]) / 100]
                     for _, fila in df_editado.iterrows() if fila["Meta %"] > 0]
            otras_monedas = [[m, t, p] for m, t, p in metas_todas[["Moneda", "Tipo", "Meta %"]].itertuples(index=False)
                             if m != moneda] if not metas_todas.empty else []
            try:
                db.set_metas_asignacion(otras_monedas + filas)
                db.clear_read_cache()
                st.success("Metas guardadas.")
                st.rerun()
            except Exception as e:
                st.error(f"No pude guardar: {e}")

    if df_editor["Meta %"].sum() > 0:
        comparacion = df_editor[(df_editor["Meta %"] > 0) | (df_editor["Real %"] > 0)]
        df_larga = comparacion.melt(id_vars="Tipo", value_vars=["Meta %", "Real %"],
                                    var_name="Serie", value_name="Porcentaje")
        fig = px.bar(df_larga, x="Tipo", y="Porcentaje", color="Serie", barmode="group",
                    title="Meta vs. real por tipo de activo")
        fig.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=340, yaxis_title="%", xaxis_title="")
        st.plotly_chart(fig, use_container_width=True, key=f"metas_chart_{key_sufijo}")


def _xirr(flujos):
    """Tasa interna de retorno anualizada para flujos con fecha (XIRR),
    resuelta por bisección para no depender de scipy — el problema es
    chico (unas pocas decenas de flujos) y suele tener una sola raíz en el
    rango considerado. 'flujos': lista de (fecha, monto); salida negativa
    (aporte) y entrada positiva (retiro o valor final). None si no
    converge (p. ej. todos los flujos tienen el mismo signo)."""
    if len(flujos) < 2:
        return None
    fecha0 = min(f for f, _ in flujos)

    def van(tasa):
        return sum(monto / (1 + tasa) ** ((fecha - fecha0).days / 365) for fecha, monto in flujos)

    lo, hi = -0.99, 10.0
    van_lo, van_hi = van(lo), van(hi)
    if van_lo == 0:
        return lo
    if van_lo * van_hi > 0:
        return None
    mid = lo
    for _ in range(200):
        mid = (lo + hi) / 2
        van_mid = van(mid)
        if abs(van_mid) < 1e-6:
            return mid
        if van_lo * van_mid < 0:
            hi = mid
        else:
            lo, van_lo = mid, van_mid
    return mid


def _rentabilidad_xirr(moneda):
    """XIRR de la cartera: aportes/retiros con su fecha real, más el valor
    actual como si se liquidara hoy. Para dólares, convierte solo el valor
    final a COP con la TRM de hoy — los aportes ya están en pesos reales (lo
    que salió de la cuenta ese día), así que no hace falta convertir nada
    más para que el cálculo sea válido."""
    df_ap = db.read_aportes_inversion(moneda)
    flujos = []
    if not df_ap.empty:
        for _, fila in df_ap.iterrows():
            fecha = pd.to_datetime(fila["Fecha"], dayfirst=True, errors="coerce")
            if pd.isna(fecha):
                continue
            monto = _num_o_cero(fila["Monto Transferido (COP)"])
            if monto:
                flujos.append((fecha.date(), -monto))
    valor_actual = _patrimonio_total_inversiones(moneda)
    if moneda == "dolares":
        trm = _descargar_trm()
        if trm is None:
            return None
        valor_actual *= trm
    if valor_actual:
        flujos.append((date.today(), valor_actual))
    if len(flujos) < 2 or not any(m < 0 for _, m in flujos) or not any(m > 0 for _, m in flujos):
        return None
    return _xirr(flujos)


def _serie_acumulada(df, col_fecha, col_valor):
    """De una tabla con una columna de fecha (texto dd/mm/yyyy, como llega
    de sheets_backend) y una de monto, arma la serie acumulada (cumsum)
    ordenada por fecha — para gráficos de aportes netos o resultado
    realizado en el tiempo."""
    if df.empty:
        return pd.DataFrame(columns=["Fecha", "Acumulado"])
    tmp = df.copy()
    tmp["Fecha"] = pd.to_datetime(tmp[col_fecha], dayfirst=True, errors="coerce")
    tmp = tmp.dropna(subset=["Fecha"]).sort_values("Fecha")
    if tmp.empty:
        return pd.DataFrame(columns=["Fecha", "Acumulado"])
    tmp["Acumulado"] = pd.to_numeric(tmp[col_valor], errors="coerce").fillna(0.0).cumsum()
    return tmp[["Fecha", "Acumulado"]]


def _render_crecimiento_rentabilidad(moneda):
    unidad = "COP" if moneda == "pesos" else "USD"

    df_snap = db.read_historial_valor_cartera()
    df_snap = df_snap[df_snap["Moneda"] == moneda] if not df_snap.empty else df_snap
    df_ap = db.read_aportes_inversion(moneda)
    serie_aportes = _serie_acumulada(df_ap, "Fecha", "Monto Transferido (COP)")

    serie_valor = pd.DataFrame(columns=["Fecha", "Valor"])
    if not df_snap.empty:
        tmp = df_snap.copy()
        tmp["Fecha"] = pd.to_datetime(tmp["Fecha"], dayfirst=True, errors="coerce")
        tmp = tmp.dropna(subset=["Fecha"]).sort_values("Fecha")
        serie_valor = tmp[["Fecha", "Valor Actual"]].rename(columns={"Valor Actual": "Valor"})

    st.markdown("#### 📊 Crecimiento y Rentabilidad")
    st.caption("El resultado realizado acumulado (compras/ventas/cortos/coberturas) está más abajo, en "
               "'Historial de posiciones y cuenta de margen'.")
    if serie_valor.empty and serie_aportes.empty:
        st.info("Todavía no hay historial para graficar — a medida que actualices precios, guardes posiciones "
                "o cargues aportes, esta sección va a ir acumulando la serie en el tiempo.")
        return
    if len(serie_valor) < 2:
        st.caption(f"El valor de cartera se guarda como una foto cada vez que actualizás precios o guardás "
                   f"posiciones — llevás {len(serie_valor)} foto(s) para {unidad}. El gráfico se va a ir "
                   "llenando con el uso normal de la app.")

    if moneda == "pesos":
        partes = []
        if not serie_valor.empty:
            partes.append(serie_valor.rename(columns={"Valor": "Monto"}).assign(Serie="Valor de Cartera"))
        if not serie_aportes.empty:
            partes.append(serie_aportes.rename(columns={"Acumulado": "Monto"})
                          .assign(Serie="Aportes Netos Acumulados"))
        if partes:
            df_crec = pd.concat(partes, ignore_index=True)
            fig = px.line(df_crec, x="Fecha", y="Monto", color="Serie", markers=True,
                         title=f"Crecimiento de la cartera ({unidad})")
            fig.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=380, yaxis_title=unidad, xaxis_title="")
            st.plotly_chart(fig, use_container_width=True, key=f"crecimiento_{moneda}")
        m1, m2 = st.columns(2)
        if not serie_valor.empty and not serie_aportes.empty:
            ultimo_valor = serie_valor["Valor"].iloc[-1]
            ultimo_aporte = serie_aportes["Acumulado"].iloc[-1]
            if ultimo_aporte:
                m1.metric("Rentabilidad sobre aportes netos",
                         f"{(ultimo_valor - ultimo_aporte) / ultimo_aporte:.2%}")
        xirr = _rentabilidad_xirr(moneda)
        if xirr is not None:
            m2.metric("Rentabilidad anualizada (XIRR)", f"{xirr:.2%}")
    else:
        if not serie_valor.empty or not serie_aportes.empty:
            fig = make_subplots(specs=[[{"secondary_y": True}]])
            if not serie_valor.empty:
                fig.add_trace(go.Scatter(x=serie_valor["Fecha"], y=serie_valor["Valor"],
                                         name="Valor de Cartera (USD)", mode="lines+markers"), secondary_y=False)
            if not serie_aportes.empty:
                fig.add_trace(go.Scatter(x=serie_aportes["Fecha"], y=serie_aportes["Acumulado"],
                                         name="Aportes Netos Acumulados (COP)", mode="lines+markers"),
                              secondary_y=True)
            fig.update_yaxes(title_text="USD", secondary_y=False)
            fig.update_yaxes(title_text="COP (aportes transferidos)", secondary_y=True)
            fig.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=380,
                              title="Crecimiento de la cartera (USD) vs. aportes transferidos (COP)")
            st.plotly_chart(fig, use_container_width=True, key=f"crecimiento_{moneda}")
            st.caption("El valor de cartera está en dólares y los aportes en pesos transferidos — este gráfico "
                       "en particular no los convierte entre sí, así que compara el ritmo de cada uno, no una "
                       "rentabilidad exacta en una sola moneda. El XIRR de abajo sí puede combinarlos, porque "
                       "solo necesita el valor final convertido con la TRM de hoy.")
        xirr = _rentabilidad_xirr(moneda)
        if xirr is not None:
            st.metric("Rentabilidad anualizada (XIRR, valor final convertido con TRM de hoy)", f"{xirr:.2%}")

    simbolo_benchmark, nombre_benchmark = BENCHMARKS.get(moneda, (None, None))
    if simbolo_benchmark:
        valor_shadow, _ = _valor_shadow_benchmark(moneda)
        if valor_shadow is not None:
            valor_real = _patrimonio_total_inversiones(moneda)
            diferencia = valor_real - valor_shadow
            st.markdown(f"##### 📈 Comparación contra {nombre_benchmark}")
            st.caption(f"Si cada aporte/retiro real (misma fecha, mismo monto) se hubiera puesto en "
                      f"{nombre_benchmark} en vez de en tu cartera, hoy valdría lo de abajo — {unidad} porque "
                      "se compara en la moneda nativa de cada hoja (para dólares, convirtiendo cada aporte con "
                      "la TRM histórica del día que lo hiciste, no la de hoy).")
            b1, b2, b3 = st.columns(3)
            b1.metric(f"Tu cartera hoy ({unidad})", f"${valor_real:,.2f}" if moneda == "dolares"
                     else f"${valor_real:,.0f}")
            b2.metric(f"{nombre_benchmark} con los mismos aportes ({unidad})",
                     f"${valor_shadow:,.2f}" if moneda == "dolares" else f"${valor_shadow:,.0f}")
            b3.metric("Diferencia", f"{'+' if diferencia >= 0 else ''}${diferencia:,.2f}" if moneda == "dolares"
                     else f"{'+' if diferencia >= 0 else ''}${diferencia:,.0f}")
        else:
            st.caption(f"No pude descargar el histórico de {nombre_benchmark} ({simbolo_benchmark}) para "
                      "comparar — puede ser un corte temporal de Yahoo Finance, o que el símbolo no sea el "
                      "correcto; avisame si querés que pruebe otro.")


def render_inversiones():
    st.title("📈 Inversiones")
    st.caption("Flujos: depósitos que salieron de tu cuenta y retiros que regresaron desde una plataforma "
               "(Trii, Acciones y Valores, Plenti, Binance, Hapi) — la app las detecta solas al cargar un "
               "extracto de cuenta y los agrega acá, con fecha y monto (siempre el peso transferido, todavía "
               "sin convertir a dólares). Más abajo, Posiciones: qué compraste con esa plata y su valor de "
               "mercado real (cantidad × precio actual), que alimenta el Balance General de 🏢 Estados "
               "Financieros.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Muestra depósitos y retiros entre tu cuenta y las plataformas de inversión, separados en pesos "
            "y dólares, con su flujo neto por plataforma. El expander '➕ Agregar un aporte o retiro "
            "manualmente' permite cargar uno sin necesidad de un extracto de cuenta; el expander "
            "'✏️ Editar o eliminar aportes/retiros' deja corregir o borrar cualquiera ya cargado (a mano o "
            "importado).\n\n"
            "Más abajo, **Posiciones**: una tabla editable por moneda (Ticker/Fondo, Tipo, Cantidad, Precio "
            "Compra Promedio, Precio Actual) — Costo Total, Valor Actual y Ganancia/Pérdida se calculan solos. "
            "Podés actualizar los precios desde Yahoo Finance o corregirlos manualmente; el valor guardado es lo "
            "que usa el Balance General de 🏢 Estados Financieros como Activo de inversión. Cada vez que "
            "actualizás precios o guardás posiciones/aportes, la app guarda una foto del valor total del día — "
            "esa serie alimenta **📊 Crecimiento y Rentabilidad** (dentro de cada pestaña de Posiciones): valor "
            "de cartera vs. aportes netos en el tiempo, y la rentabilidad sobre lo aportado.\n\n"
            "Más abajo todavía, **Historial de posiciones y cuenta de margen** trae lo importado de reportes "
            "de brokers (compras, ventas, cortos, coberturas) — también editable/eliminable desde "
            "'✏️ Editar o eliminar operaciones del historial'.")

    _render_patrimonio_unificado()
    st.divider()
    _render_importar_portafolios()
    _render_actualizar_precios()

    col_p, col_d = st.columns(2)
    with col_p:
        st.markdown("**Pesos (COP)** — Trii / Acciones y Valores: acciones y fondos")
        _form_agregar_aporte(
            "pesos", sorted(set(PLATAFORMAS_INVERSION_PESOS.values()) | {"Fondo de Inversión (banco)"}), "pesos")
        df_ap_pesos = db.read_aportes_inversion("pesos")
        if not df_ap_pesos.empty:
            df_show_p = df_ap_pesos.copy()
            df_show_p.insert(2, "Flujo", df_show_p["Monto Transferido (COP)"].apply(
                lambda monto: "Depósito" if monto >= 0 else "Retiro"))
            df_show_p["Monto Transferido (COP)"] = df_show_p["Monto Transferido (COP)"].map(
                lambda v: f"{fmt_moneda(v)}")
            total_dep = df_ap_pesos.loc[df_ap_pesos["Monto Transferido (COP)"] > 0,
                                         "Monto Transferido (COP)"].sum()
            total_ret = -df_ap_pesos.loc[df_ap_pesos["Monto Transferido (COP)"] < 0,
                                          "Monto Transferido (COP)"].sum()
            k1, k2, k3 = st.columns(3)
            k1.metric("Depósitos", fmt_moneda(total_dep))
            k2.metric("Retiros", fmt_moneda(total_ret))
            k3.metric("Flujo neto", fmt_moneda(total_dep - total_ret))
            st.dataframe(df_show_p, hide_index=True, use_container_width=True, height=300)
        else:
            st.info("Todavía no hay aportes registrados en pesos.")
        _form_editar_aportes("pesos", "pesos")
    with col_d:
        st.markdown("**Dólares (USD)** — Plenti, Binance, Hapi, Interactive Brokers")
        _form_agregar_aporte("dolares", sorted(set(PLATAFORMAS_INVERSION_DOLARES.values())), "dolares")
        df_ap_dolares = db.read_aportes_inversion("dolares")
        if not df_ap_dolares.empty:
            df_show_d = df_ap_dolares.copy()
            df_show_d.insert(2, "Flujo", df_show_d["Monto Transferido (COP)"].apply(
                lambda monto: "Depósito" if monto >= 0 else "Retiro"))
            df_show_d["Monto Transferido (COP)"] = df_show_d["Monto Transferido (COP)"].map(
                lambda v: f"{fmt_moneda(v)}")
            total_dep = df_ap_dolares.loc[df_ap_dolares["Monto Transferido (COP)"] > 0,
                                           "Monto Transferido (COP)"].sum()
            total_ret = -df_ap_dolares.loc[df_ap_dolares["Monto Transferido (COP)"] < 0,
                                            "Monto Transferido (COP)"].sum()
            k1, k2, k3 = st.columns(3)
            k1.metric("Depósitos", fmt_moneda(total_dep))
            k2.metric("Retiros", fmt_moneda(total_ret))
            k3.metric("Flujo neto", fmt_moneda(total_dep - total_ret))
            st.dataframe(df_show_d, hide_index=True, use_container_width=True, height=300)
        else:
            st.info("Todavía no hay aportes registrados en dólares.")
        _form_editar_aportes("dolares", "dolares")
    # -- Depósitos y retiros por plataforma (pesos y dólares, aparte) --
    partes = []
    if not df_ap_pesos.empty:
        base = df_ap_pesos.copy()
        base["Flujo"] = base["Monto Transferido (COP)"].apply(
            lambda monto: "Depósito" if monto >= 0 else "Retiro")
        t = base.groupby(["Plataforma", "Flujo"])["Monto Transferido (COP)"].sum().reset_index()
        t["Moneda destino"] = "Pesos"
        partes.append(t)
    if not df_ap_dolares.empty:
        base = df_ap_dolares.copy()
        base["Flujo"] = base["Monto Transferido (COP)"].apply(
            lambda monto: "Depósito" if monto >= 0 else "Retiro")
        t = base.groupby(["Plataforma", "Flujo"])["Monto Transferido (COP)"].sum().reset_index()
        t["Moneda destino"] = "Dólares"
        partes.append(t)
    if partes:
        st.markdown("**Depósitos y retiros por plataforma** (siempre en pesos transferidos)")
        df_aportes_plat = pd.concat(partes, ignore_index=True)
        g1, g2 = st.columns(2)
        fig_aportes = px.bar(df_aportes_plat, x="Monto Transferido (COP)", y="Plataforma", color="Flujo",
                             orientation="h", barmode="relative", facet_col="Moneda destino",
                             color_discrete_map={"Depósito": "#1d4ed8", "Retiro": "#dc2626"},
                             title="Flujo acumulado por plataforma")
        fig_aportes.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=360,
                                  xaxis_title="COP", yaxis_title="")
        g1.plotly_chart(fig_aportes, use_container_width=True)

        cronologia = []
        for moneda_destino, df_ap in (("Pesos", df_ap_pesos), ("Dólares", df_ap_dolares)):
            if df_ap.empty:
                continue
            tmp = df_ap.copy()
            tmp["Fecha"] = pd.to_datetime(tmp["Fecha"], dayfirst=True, errors="coerce")
            tmp = tmp.dropna(subset=["Fecha"]).sort_values("Fecha")
            tmp["Flujo acumulado"] = tmp["Monto Transferido (COP)"].cumsum()
            tmp["Moneda destino"] = moneda_destino
            cronologia.append(tmp)
        if cronologia:
            df_cronologia = pd.concat(cronologia, ignore_index=True)
            fig_flujo = px.line(df_cronologia, x="Fecha", y="Flujo acumulado", color="Moneda destino",
                                markers=True, title="Capital neto transferido en el tiempo")
            fig_flujo.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=360,
                                    yaxis_title="COP", xaxis_title="")
            g2.plotly_chart(fig_flujo, use_container_width=True)

    st.divider()
    st.markdown("#### Posiciones")
    st.caption("Cantidad y Precio Compra Promedio se cargan una vez. Precio Actual puede actualizarse desde "
               "el mercado o editarse manualmente. Las cotizaciones son informativas y pueden tener retraso.")
    tab_pp, tab_pd = st.tabs(["Pesos (COP)", "Dólares (USD)"])
    with tab_pp:
        titulos_pesos, exposicion_pesos = _render_resumen_cartera(
            db.read_posiciones_inversion("pesos"), "pesos", "pesos")
        _render_vista_consolidada(db.read_posiciones_inversion("pesos"), "pesos", "pesos")
        _render_metas_asignacion(titulos_pesos, exposicion_pesos, "pesos", "pesos")
        _form_posiciones_inversion("pesos", "pos_pesos")
        st.divider()
        _render_crecimiento_rentabilidad("pesos")
    with tab_pd:
        titulos_dolares, exposicion_dolares = _render_resumen_cartera(
            db.read_posiciones_inversion("dolares"), "dolares", "dolares")
        _render_vista_consolidada(db.read_posiciones_inversion("dolares"), "dolares", "dolares")
        _render_metas_asignacion(titulos_dolares, exposicion_dolares, "dolares", "dolares")
        _form_posiciones_inversion("dolares", "pos_dolares")
        st.divider()
        _render_crecimiento_rentabilidad("dolares")

    st.markdown("#### Historial de posiciones y cuenta de margen")
    df_hist = db.read_historial_inversion()
    if df_hist.empty:
        st.info("Importa un reporte del broker para ver posiciones cerradas, ventas en corto y coberturas.")
    else:
        hist = df_hist.copy()
        hist["Fecha orden"] = pd.to_datetime(hist["Fecha"], dayfirst=True, errors="coerce")
        es_pasivo = hist["Operación"].astype(str).str.upper().isin({"DIVIDEND", "INTEREST"})
        pasivos, hist = hist[es_pasivo].copy(), hist[~es_pasivo].copy()
        signo = {"BUY": 1, "COVER": 1, "SELL": -1, "SHORT": -1}
        hist["Cantidad neta"] = hist.apply(
            lambda fila: signo.get(str(fila["Operación"]).upper(), 0) * float(fila["Cantidad"]), axis=1)
        grupos = []
        for (plataforma, activo, moneda), grupo in hist.groupby(["Plataforma", "Activo", "Moneda"]):
            cantidad_neta = grupo["Cantidad neta"].sum()
            estado = ("Abierta larga" if cantidad_neta > 1e-8 else
                      "Abierta corta" if cantidad_neta < -1e-8 else "Cerrada")
            estrategia = "Corto" if grupo["Operación"].astype(str).str.upper().eq("SHORT").any() else "Largo"
            primera = grupo["Fecha orden"].min()
            ultima = grupo["Fecha orden"].max()
            grupos.append({"Plataforma": plataforma, "Activo": activo, "Moneda": moneda,
                           "Estrategia": estrategia, "Estado": estado, "Cantidad neta": cantidad_neta,
                           "Resultado realizado": grupo["Resultado Realizado"].sum(),
                           "Primera operación": primera.strftime("%d/%m/%Y") if pd.notna(primera) else "",
                           "Última operación": ultima.strftime("%d/%m/%Y") if pd.notna(ultima) else ""})
        resumen_hist = pd.DataFrame(grupos).sort_values(["Estado", "Plataforma", "Activo"])
        realizado_por_moneda = hist.groupby("Moneda")["Resultado Realizado"].sum().to_dict()
        r1, r2, r3, r4 = st.columns(4)
        r1.metric("Posiciones cerradas", int((resumen_hist["Estado"] == "Cerrada").sum()))
        r2.metric("Estrategias en corto", int((resumen_hist["Estrategia"] == "Corto").sum()))
        r3.metric("Realizado en compraventas USD", f"${realizado_por_moneda.get('USD', 0):,.2f}")
        r4.metric("Realizado en compraventas COP", f"${realizado_por_moneda.get('COP', 0):,.0f}")
        st.caption("Este resultado realizado corresponde a compras, ventas, cortos y coberturas. Dividendos, "
                   "impuestos, intereses y cargos de margen se concilian aparte.")
        st.dataframe(resumen_hist, hide_index=True, use_container_width=True)
        cierres = hist[hist["Resultado Realizado"] != 0].copy()
        if not cierres.empty:
            por_activo = (cierres.groupby(["Plataforma", "Activo", "Moneda"], as_index=False)
                          ["Resultado Realizado"].sum())
            por_activo["Resultado"] = por_activo["Resultado Realizado"].apply(
                lambda valor_r: "Ganancia" if valor_r >= 0 else "Pérdida")
            por_activo["Posición"] = por_activo["Plataforma"] + " - " + por_activo["Activo"]
            por_activo = por_activo.sort_values("Resultado Realizado")
            cierres_fecha = (cierres.dropna(subset=["Fecha orden"])
                             .groupby(["Moneda", "Fecha orden"], as_index=False)["Resultado Realizado"].sum()
                             .sort_values(["Moneda", "Fecha orden"]))
            cierres_fecha["Resultado realizado acumulado"] = cierres_fecha.groupby("Moneda")[
                "Resultado Realizado"].cumsum()
            h1, h2 = st.columns(2)
            fig_realizado = px.bar(por_activo, x="Resultado Realizado", y="Posición", color="Resultado",
                                    orientation="h", facet_col="Moneda", title="Resultado realizado por posición",
                                    color_discrete_map={"Ganancia": "#0d9488", "Pérdida": "#dc2626"})
            fig_realizado.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=440,
                                         xaxis_title="", yaxis_title="", showlegend=False)
            fig_realizado.update_xaxes(matches=None)
            h1.plotly_chart(fig_realizado, use_container_width=True, key="resultado_realizado_activo")
            fig_acumulado = px.line(cierres_fecha, x="Fecha orden", y="Resultado realizado acumulado",
                                    facet_col="Moneda", markers=True, title="Resultado realizado acumulado")
            fig_acumulado.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=440,
                                         xaxis_title="", yaxis_title="")
            fig_acumulado.update_yaxes(matches=None)
            h2.plotly_chart(fig_acumulado, use_container_width=True, key="resultado_realizado_acumulado")
        with st.expander("Ver todas las compras, ventas, cortos y coberturas"):
            vista_hist = (hist.sort_values("Fecha orden", ascending=False)
                          .drop(columns=["Cantidad neta", "Fecha orden"]))
            st.dataframe(vista_hist, hide_index=True, use_container_width=True)

        if not pasivos.empty:
            st.markdown("##### 💵 Dividendos e intereses recibidos")
            st.caption("Aparte del resultado realizado de compraventas — ingreso pasivo real, no viene de "
                       "vender nada.")
            por_moneda_pasivo = pasivos.groupby("Moneda")["Resultado Realizado"].sum().to_dict()
            p1, p2 = st.columns(2)
            p1.metric("Recibido en USD", f"${por_moneda_pasivo.get('USD', 0):,.2f}")
            p2.metric("Recibido en COP", f"${por_moneda_pasivo.get('COP', 0):,.0f}")
            for moneda_pasivo in sorted(pasivos["Moneda"].unique()):
                serie_pasivo = _serie_acumulada(pasivos[pasivos["Moneda"] == moneda_pasivo],
                                                "Fecha", "Resultado Realizado")
                if len(serie_pasivo) >= 2:
                    fig_pasivo = px.line(serie_pasivo, x="Fecha", y="Acumulado", markers=True,
                                        title=f"Dividendos e intereses acumulados ({moneda_pasivo})")
                    fig_pasivo.update_layout(margin=dict(l=0, r=0, t=45, b=0), height=280,
                                            yaxis_title=moneda_pasivo, xaxis_title="")
                    st.plotly_chart(fig_pasivo, use_container_width=True,
                                    key=f"pasivo_acumulado_{moneda_pasivo}")
            with st.expander("Ver el detalle de dividendos e intereses"):
                st.dataframe(pasivos.sort_values("Fecha orden", ascending=False)
                            .drop(columns=["Fecha orden"]), hide_index=True, use_container_width=True)

    _form_agregar_dividendo()
    _form_editar_historial()

    st.divider()
    st.markdown("#### 💰 Cesantías")
    st.caption(
        "El capital de cesantías no entra a la cuenta de ahorros ni al presupuesto mensual — lo administra el "
        "fondo aparte. Cada enero, el Hospital emite un comprobante de liquidación con el mismo formato que una "
        "colilla de quincena; la app ya lo detecta solo (por 'CESANTIAS' en el texto) y no lo guarda en 📄 "
        "Colillas de Pago. Los intereses de esa liquidación sí son ingreso real y quedan solos en Otros Ingresos, "
        "categoría 'Cesantías'. Acá abajo: el saldo real del fondo, tal como venga en el extracto que te manda "
        "el fondo — cargalo a mano cuando lo recibas.")

    df_oi_cesantias = db.read_otros_ingresos_tabla()
    total_interes_cesantias = (df_oi_cesantias.loc[df_oi_cesantias["Categoría"] == "Cesantías", "Valor"].sum()
                                if not df_oi_cesantias.empty else 0.0)
    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(total_interes_cesantias)}</div>'
                f'<div class="stat-label">Intereses de cesantías recibidos (histórico)</div></div>',
                unsafe_allow_html=True)

    with st.expander("➕ Agregar/actualizar un corte del fondo de cesantías"):
        with st.form("form_agregar_cesantias"):
            c1, c2 = st.columns(2)
            fecha_corte = c1.date_input("Fecha de Corte", value=date.today(), key="fecha_corte_cesantias")
            fondo = c2.text_input("Fondo", key="fondo_cesantias")
            c3, c4 = st.columns(2)
            saldo_anterior_ces = c3.number_input("Saldo Anterior", step=1000.0, format="%.0f",
                                                   key="saldo_ant_cesantias")
            aportes_ces = c4.number_input("Aportes del Período", step=1000.0, format="%.0f",
                                            key="aportes_cesantias")
            c5, c6 = st.columns(2)
            rendimientos_ces = c5.number_input("Rendimientos", step=1000.0, format="%.0f",
                                                 key="rendimientos_cesantias")
            retiros_ces = c6.number_input("Retiros", step=1000.0, format="%.0f", key="retiros_cesantias")
            saldo_actual_ces = st.number_input("Saldo Actual", step=1000.0, format="%.0f",
                                                 key="saldo_act_cesantias")
            notas_ces = st.text_input("Notas (opcional)", key="notas_cesantias")
            guardar_ces = st.form_submit_button("💾 Guardar corte", type="primary")
        if guardar_ces:
            try:
                fila_ces = db.first_blank_row("cesantias")
                db.write_block_rows("cesantias", fila_ces,
                                     [[fecha_corte.isoformat(), fondo.strip(), saldo_anterior_ces, aportes_ces,
                                       rendimientos_ces, retiros_ces, saldo_actual_ces, notas_ces.strip()]])
                db.clear_read_cache()
                st.success("Corte de cesantías agregado.")
                st.rerun()
            except Exception as e:
                st.error(f"No pude guardar: {e}. Si es la primera vez, revisá que exista una hoja 'Cesantías' "
                         "en el Google Sheet con las columnas Fecha de Corte/Fondo/Saldo Anterior/Aportes del "
                         "Período/Rendimientos/Retiros/Saldo Actual/Notas.")

    try:
        df_cesantias = db.read_cesantias()
    except Exception as e:
        df_cesantias = None
        st.info(f"Todavía no pude leer la hoja 'Cesantías' del Sheet ({e}). Se crea sola al agregar el primer "
                "corte de arriba.")
    if df_cesantias is not None:
        if df_cesantias.empty:
            st.info("Todavía no hay cortes del fondo de cesantías cargados — subí el extracto del fondo cuando "
                    "lo tengas y cargá acá cada corte.")
        else:
            df_ces_show = df_cesantias.copy()
            for col in ["Saldo Anterior", "Aportes del Período", "Rendimientos", "Retiros", "Saldo Actual"]:
                df_ces_show[col] = df_ces_show[col].map(lambda v: f"{fmt_moneda(v)}")
            st.dataframe(df_ces_show, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_ces_show)))

    st.divider()
    st.markdown("#### 🏦 Pensión Obligatoria")
    st.caption(
        "El fondo de pensión obligatoria (AFP) lo administra la AFP aparte, igual que cesantías — el saldo no "
        "entra a la cuenta de ahorros ni al presupuesto mensual. Cada quincena se descuenta el aporte de ley "
        "('Aporte Pensión' en Colillas de Pago); el saldo real con sus rendimientos es el que reporta la AFP en "
        "su extracto — cargalo a mano acá cuando lo recibas.")

    df_desc_pension = db.read_colillas_descuentos()
    total_aportes_nomina_pension = (
        df_desc_pension.loc[df_desc_pension["Concepto"].str.startswith("Aporte Pensión", na=False), "Valor"].sum()
        if not df_desc_pension.empty else 0.0)
    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(total_aportes_nomina_pension)}</div>'
                f'<div class="stat-label">Aportes vía nómina descontados (histórico)</div></div>',
                unsafe_allow_html=True)

    with st.expander("➕ Agregar/actualizar un corte del fondo de pensión obligatoria"):
        with st.form("form_agregar_pension"):
            c1, c2 = st.columns(2)
            fecha_corte_pen = c1.date_input("Fecha de Corte", value=date.today(), key="fecha_corte_pension")
            fondo_pen = c2.text_input("Fondo (AFP)", key="fondo_pension")
            c3, c4 = st.columns(2)
            saldo_anterior_pen = c3.number_input("Saldo Anterior", step=1000.0, format="%.0f",
                                                   key="saldo_ant_pension")
            aportes_pen = c4.number_input("Aportes del Período", step=1000.0, format="%.0f",
                                            key="aportes_pension")
            c5, c6 = st.columns(2)
            rendimientos_pen = c5.number_input("Rendimientos", step=1000.0, format="%.0f",
                                                 key="rendimientos_pension")
            retiros_pen = c6.number_input("Retiros", step=1000.0, format="%.0f", key="retiros_pension")
            saldo_actual_pen = st.number_input("Saldo Actual", step=1000.0, format="%.0f",
                                                 key="saldo_act_pension")
            notas_pen = st.text_input("Notas (opcional)", key="notas_pension")
            guardar_pen = st.form_submit_button("💾 Guardar corte", type="primary")
        if guardar_pen:
            try:
                fila_pen = db.first_blank_row("pension_obligatoria")
                db.write_block_rows("pension_obligatoria", fila_pen,
                                     [[fecha_corte_pen.isoformat(), fondo_pen.strip(), saldo_anterior_pen,
                                       aportes_pen, rendimientos_pen, retiros_pen, saldo_actual_pen,
                                       notas_pen.strip()]])
                db.clear_read_cache()
                st.success("Corte de pensión obligatoria agregado.")
                st.rerun()
            except Exception as e:
                st.error(f"No pude guardar: {e}. Si es la primera vez, revisá que exista una hoja 'Pensión "
                         "Obligatoria' en el Google Sheet con las columnas Fecha de Corte/Fondo/Saldo "
                         "Anterior/Aportes del Período/Rendimientos/Retiros/Saldo Actual/Notas.")

    try:
        df_pension = db.read_pension_obligatoria()
    except Exception as e:
        df_pension = None
        st.info(f"Todavía no pude leer la hoja 'Pensión Obligatoria' del Sheet ({e}). Se crea sola al agregar "
                "el primer corte de arriba.")
    if df_pension is not None:
        if df_pension.empty:
            st.info("Todavía no hay cortes del fondo de pensión obligatoria cargados — subí el extracto de la "
                    "AFP cuando lo tengas y cargá acá cada corte.")
        else:
            df_pen_show = df_pension.copy()
            for col in ["Saldo Anterior", "Aportes del Período", "Rendimientos", "Retiros", "Saldo Actual"]:
                df_pen_show[col] = df_pen_show[col].map(lambda v: f"{fmt_moneda(v)}")
            st.dataframe(df_pen_show, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_pen_show)))


def render_verificar_datos():
    st.title("✅ Verificar Datos")
    st.info(
        "**¿Para qué sirve esta pestaña?** Es un chequeo de tranquilidad, no algo que tengas que usar "
        "seguido: sirve para confirmar que lo que cargaste en la app coincide con tus extractos reales del "
        "banco y las tarjetas — así detectás si falta cargar un movimiento, si algo quedó duplicado, o si "
        "hay algo mal clasificado.\n\n"
        "- **Tarjetas de crédito** (abajo): muestra el resumen de cada extracto tal como lo cargaste — "
        "compará estos números contra el PDF/Excel del banco.\n"
        "- **Efectivo (cuenta de ahorros)**: acá sí interactuás — escribís el saldo inicial y final que "
        "dice tu extracto bancario para un mes, y la app calcula solo cuánto DEBERÍA quedar en la cuenta "
        "según lo que ya cargaste (saldo inicial + ingresos − egresos). Si coincide con el saldo final que "
        "escribiste, ese mes está completo. Si no coincide, revisá si falta cargar algún movimiento de la "
        "cuenta de ese mes.")

    st.markdown("**Tarjetas de crédito — pesos (COP)** (resumen de cada extracto, como referencia)")
    st.caption("'Pago Total' es el pago por la totalidad de la deuda (la opción que ofrece la tarjeta para "
               "cancelar todo el saldo pendiente de una vez), no necesariamente lo que pagaste realmente ese "
               "corte — por eso todavía no comparo esto contra el detalle de compras: con compras a cuotas, "
               "el cupo usado incluye cuotas futuras que 'Compras del Mes' no cuenta, así que un chequeo "
               "automático de conciliación de tarjeta (como el que ya existe para Efectivo) necesitaría "
               "primero resolver eso.")
    filas_tarjetas = []
    for bk, nombre in [("visa_resumen", "Visa 7497"), ("mc_resumen", "Mastercard 5922")]:
        df_r = db.read_tarjeta_resumen(bk)
        for _, r in df_r.iterrows():
            filas_tarjetas.append(dict(
                Tarjeta=nombre, Periodo=r["Periodo Extracto"], **{
                    "Cupo Total": r["Cupo Total"], "Cupo Disponible": r["Cupo Disponible"],
                    "Saldo Anterior": r["Saldo Anterior"], "Compras del Mes": r["Compras del Mes"],
                    "Pago Total": r["Pago Total"]}))

    if filas_tarjetas:
        df_tarjetas = pd.DataFrame(filas_tarjetas).sort_values(["Tarjeta", "Periodo"], ascending=[True, False])
        df_tarjetas_show = df_tarjetas.copy()
        for col in ["Cupo Total", "Cupo Disponible", "Saldo Anterior", "Compras del Mes", "Pago Total"]:
            df_tarjetas_show[col] = df_tarjetas_show[col].map(lambda v: f"{fmt_moneda(v)}")
        st.dataframe(df_tarjetas_show, hide_index=True, use_container_width=True, height=400)
    else:
        st.info("Todavía no hay resúmenes de tarjeta cargados.")

    st.markdown("**Mastercard 5922 — dólares (USD)**, aparte — nunca se suma con los pesos de arriba")
    df_mc_usd_conc = db.read_egreso_detalle("mc_detalle_usd")
    df_mc_resumen = db.read_tarjeta_resumen("mc_resumen")
    if not df_mc_usd_conc.empty or not df_mc_resumen.empty:
        resumen_usd = (df_mc_usd_conc.groupby("Periodo Extracto")["Valor Cargado Este Periodo"]
                       .sum().reset_index().rename(columns={"Valor Cargado Este Periodo": "Compras del Mes (USD)"})
                       if not df_mc_usd_conc.empty
                       else pd.DataFrame(columns=["Periodo Extracto", "Compras del Mes (USD)"]))
        saldo_usd = df_mc_resumen[["Periodo Extracto", "Saldo a pagar USD"]]
        resumen_usd = resumen_usd.merge(saldo_usd, on="Periodo Extracto", how="outer").sort_values(
            "Periodo Extracto", ascending=False)
        resumen_usd["Compras del Mes (USD)"] = resumen_usd["Compras del Mes (USD)"].map(
            lambda v: f"US$ {v:,.2f}" if pd.notna(v) else "-")
        resumen_usd["Saldo a pagar USD"] = resumen_usd["Saldo a pagar USD"].map(
            lambda v: f"US$ {v:,.2f}" if pd.notna(v) else "-")
        st.dataframe(resumen_usd, hide_index=True, use_container_width=True, height=200)
    else:
        st.info("Todavía no hay compras en USD de Mastercard cargadas.")

    st.divider()
    st.markdown("**Efectivo (cuenta de ahorros)** — el extracto no trae el saldo de la cuenta, así que "
                "lo ingresás vos cada mes para poder conciliar. Todo lo que se carga desde un extracto "
                "de la cuenta cuenta acá (también lo '(no presupuestar)'), para que la cadena de saldos "
                "cuadre aunque esos movimientos no afecten tu presupuesto.")

    df_ef_all = db.read_egreso_detalle("efectivo_detalle")
    df_oi_all = db.read_otros_ingresos_tabla()
    if not df_ef_all.empty:
        am = [_extraer_anio_mes(p) for p in df_ef_all["Periodo Extracto"]]
        df_ef_all = df_ef_all.assign(_mes_str=[f"{a}-{m:02d}" if a and m else None for a, m in am])
    if not df_oi_all.empty:
        am = [_extraer_anio_mes(p) for p in df_oi_all["Fecha"]]
        df_oi_all = df_oi_all.assign(_mes_str=[f"{a}-{m:02d}" if a and m else None for a, m in am])

    def _ingresos_egresos_mes(mes_str):
        egresos = (df_ef_all[df_ef_all["_mes_str"] == mes_str]["Valor Cargado Este Periodo"].sum()
                   if not df_ef_all.empty else 0)
        ingresos = (df_oi_all[df_oi_all["_mes_str"] == mes_str]["Valor"].sum()
                   if not df_oi_all.empty else 0)
        return ingresos, egresos

    hoy = date.today()
    anios_conc = list(range(2023, 2033))
    c1, c2 = st.columns(2)
    anio_conc = c1.selectbox("Año", anios_conc,
                              index=anios_conc.index(hoy.year) if hoy.year in anios_conc else 0,
                              key="anio_conciliacion")
    mes_conc_num = c2.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                 format_func=lambda m: MESES_NOMBRE[m], key="mes_conciliacion")
    mes_conc = f"{anio_conc}-{mes_conc_num:02d}"
    mes_anterior_conc = f"{anio_conc}-{mes_conc_num - 1:02d}" if mes_conc_num > 1 else f"{anio_conc - 1}-12"

    df_conc = db.read_conciliacion_efectivo()
    existente = df_conc[df_conc["Mes"] == mes_conc] if not df_conc.empty else df_conc
    if not existente.empty:
        saldo_inicial_prev = float(existente.iloc[0]["Saldo Inicial"])
        saldo_final_prev = float(existente.iloc[0]["Saldo Final"])
    else:
        # Sin saldo inicial guardado para este mes todavía: lo autocompletamos
        # con el saldo final que quedó guardado el mes anterior (si lo hay),
        # para no tener que volver a escribirlo cada vez.
        fila_anterior = df_conc[df_conc["Mes"] == mes_anterior_conc] if not df_conc.empty else df_conc
        saldo_inicial_prev = float(fila_anterior.iloc[0]["Saldo Final"]) if not fila_anterior.empty else 0.0
        saldo_final_prev = 0.0

    if existente.empty and (df_conc.empty or fila_anterior.empty):
        st.warning(f"⚠️ No hay un saldo guardado para {mes_anterior_conc} — el $0 de abajo es solo un valor "
                   "de partida, no un cálculo. Ingresá el saldo real de tu cuenta al cierre de ese mes (el "
                   "saldo al 31 de diciembre si este es el primer mes que cargás) para que la cadena de "
                   "saldos arranque bien.")

    c3, c4 = st.columns(2)
    saldo_inicial = c3.number_input("Saldo inicial del mes", value=saldo_inicial_prev, step=1000.0,
                                     format="%.0f", key=f"saldo_ini_{mes_conc}")
    saldo_final = c4.number_input("Saldo final del mes", value=saldo_final_prev, step=1000.0,
                                   format="%.0f", key=f"saldo_fin_{mes_conc}")

    if st.button("💾 Guardar saldos de este mes", key="btn_guardar_conciliacion"):
        db.guardar_conciliacion_efectivo(mes_conc, saldo_inicial, saldo_final)
        db.clear_read_cache()
        st.success(f"Saldos de {mes_conc} guardados.")
        st.rerun()

    ingresos_mes, egresos_mes = _ingresos_egresos_mes(mes_conc)
    saldo_calculado_ef = saldo_inicial + ingresos_mes - egresos_mes
    diferencia_ef = saldo_calculado_ef - saldo_final

    c5, c6, c7, c8 = st.columns(4)
    c5.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(ingresos_mes)}</div>'
                f'<div class="stat-label">Ingresos del mes</div></div>', unsafe_allow_html=True)
    c6.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(egresos_mes)}</div>'
                f'<div class="stat-label">Egresos del mes</div></div>', unsafe_allow_html=True)
    c7.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(saldo_calculado_ef)}</div>'
                f'<div class="stat-label">Saldo calculado</div></div>', unsafe_allow_html=True)
    c8.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(diferencia_ef)}</div>'
                f'<div class="stat-label">Diferencia vs. saldo final</div></div>', unsafe_allow_html=True)

    if not (saldo_inicial == 0 and saldo_final == 0):
        if abs(diferencia_ef) > 100:
            st.warning(f"La conciliación de {mes_conc} no cuadra — diferencia de {fmt_moneda(diferencia_ef)}. "
                       "Revisá si falta cargar algún movimiento de ese mes o si el saldo ingresado está mal.")
        else:
            st.success(f"Conciliación de {mes_conc} cuadra ✅")

    if not df_conc.empty:
        with st.expander(f"Ver histórico de conciliación ({len(df_conc)} mes(es) con saldos guardados)"):
            filas_hist = []
            for _, fila in df_conc.sort_values("Mes").iterrows():
                ing, egr = _ingresos_egresos_mes(fila["Mes"])
                calc = fila["Saldo Inicial"] + ing - egr
                diff = calc - fila["Saldo Final"]
                if fila["Saldo Inicial"] == 0 and fila["Saldo Final"] == 0:
                    estado = "Pendiente saldos"
                elif abs(diff) <= 100:
                    estado = "✅ Conciliado"
                else:
                    estado = "⚠️ Revisar"
                filas_hist.append({"Mes": fila["Mes"], "Saldo Inicial": fila["Saldo Inicial"],
                                    "Ingresos": ing, "Egresos": egr, "Saldo Esperado": calc,
                                    "Saldo Final": fila["Saldo Final"], "Diferencia": diff, "Estado": estado})
            df_hist = pd.DataFrame(filas_hist)
            for col in ["Saldo Inicial", "Ingresos", "Egresos", "Saldo Esperado", "Saldo Final", "Diferencia"]:
                df_hist[col] = df_hist[col].map(lambda v: f"{fmt_moneda(v)}")
            st.dataframe(df_hist, hide_index=True, use_container_width=True, height=min(400, 45 + 35 * len(df_hist)))


def render_salud_datos():
    st.title("🔍 Salud de los Datos")
    st.caption("Chequeos automáticos sobre lo que ya está cargado en el Sheet, más herramientas para "
               "recategorizar en bloque — no compara contra tus extractos del banco (para eso está "
               "✅ Verificar Datos), sino que busca inconsistencias dentro de los datos mismos: quincenas "
               "duplicadas o faltantes, primas mal etiquetadas, gasto/ingreso sin categorizar y comercios o "
               "conceptos con categoría inconsistente entre un movimiento y otro.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "La mayoría son chequeos de solo lectura; dos son herramientas que sí escriben en el Sheet "
            "(recategorizar):\n\n"
            "- **Quincenas duplicadas** *(lectura)*: la misma quincena de Colillas de Pago aparece más de una "
            "vez en el resumen — puede ser una carga repetida, o dos registros con valores distintos que hay "
            "que revisar a mano.\n"
            "- **Comprobantes faltantes** *(lectura)*: de los 24 quincenas + 2 primas (junio y diciembre) "
            "esperadas por año, cuáles no están — para años ya en curso, hasta el mes anterior al actual.\n"
            "- **Primas mal etiquetadas** *(lectura)*: una 'quincena' de Colillas de Pago cuyo devengo es en "
            "realidad la Prima Legal — pasa cuando el comprobante de Prima se sube con el parser viejo (antes "
            "de detectarla por el devengo 'PRIMA LEGAL' en vez del encabezado 'Quincena De...', que el "
            "Hospital reusa para los dos tipos de comprobante) y queda pisando el lugar de la quincena real.\n"
            "- **Cesantías mal etiquetadas** *(escribe)*: mismo bug que la Prima, pero con la liquidación anual "
            "de cesantías de enero — a diferencia de la Prima, acá el botón corrige solo: relabels el período, "
            "saca el capital de Colillas de Pago (no tiene efecto en caja) y manda los intereses a Otros "
            "Ingresos.\n"
            "- **Gasto sin categorizar / Recategorizar cualquier comercio** *(escribe)*: por defecto solo "
            "muestra los comercios en 'Otros'; activá 'Mostrar todos' para poder cambiarle la categoría a "
            "cualquier comercio de gastos (Efectivo/Visa/Mastercard), tenga la categoría que tenga hoy.\n"
            "- **Recategorizar Otros Ingresos** *(escribe)*: lo mismo pero para la hoja de Otros Ingresos, por "
            "concepto — por defecto muestra los que están en la categoría genérica 'Otro'.\n"
            "- **Categoría inconsistente por comercio/concepto** *(lectura)*: el mismo comercio (en gastos) o "
            "concepto (en Otros Ingresos) aparece con más de una categoría distinta entre un movimiento y "
            "otro — puede ser un error de tipeo o un comercio que de verdad cambia de categoría según la "
            "compra. Usá las herramientas de arriba para corregirlo.")

    st.subheader("📅 Quincenas duplicadas en Colillas de Pago")
    df_cp = db.read_colillas_resumen()
    if df_cp.empty:
        st.info("Todavía no hay colillas cargadas.")
    else:
        conteo = df_cp.groupby("Periodo").size()
        duplicadas = conteo[conteo > 1].index.tolist()
        if not duplicadas:
            st.success("✅ No hay quincenas repetidas.")
        else:
            df_dup = df_cp[df_cp["Periodo"].isin(duplicadas)].sort_values(["Periodo", "Fecha de Pago"])
            df_dup_show = df_dup.copy()
            for col in ["Devengos Totales", "Descuentos Totales"]:
                df_dup_show[col] = df_dup_show[col].map(lambda v: f"{fmt_moneda(v)}")
            st.warning(f"⚠️ {len(duplicadas)} quincena(s) aparecen más de una vez: {', '.join(duplicadas)}")
            st.dataframe(df_dup_show, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_dup)))

    st.divider()

    st.subheader("🧾 Comprobantes de colillas faltantes")
    if df_cp.empty:
        st.info("Todavía no hay colillas cargadas.")
    else:
        periodos_presentes = set(df_cp["Periodo"].astype(str).str.strip())
        anios_con_datos = sorted({a for p in periodos_presentes for a, _ in [_extraer_anio_mes(p)] if a})
        hoy = date.today()
        faltantes_por_anio = {}
        for anio in anios_con_datos:
            esperados = []
            for mes_num in range(1, 13):
                if anio == hoy.year and mes_num >= hoy.month:
                    break  # el mes en curso (o futuro) todavía no tiene por qué estar completo
                mes_abr = _MESES_3LETRAS[mes_num - 1]
                esperados += [f"1a quincena {mes_abr}-{anio}", f"2a quincena {mes_abr}-{anio}"]
                if mes_num == 6:
                    esperados.append(f"Prima jun-{anio}")
                if mes_num == 12:
                    esperados.append(f"Prima dic-{anio}")
            faltantes = [p for p in esperados if p not in periodos_presentes]
            if faltantes:
                faltantes_por_anio[anio] = faltantes
        if not faltantes_por_anio:
            st.success("✅ No falta ningún comprobante esperado (24 quincenas + primas de junio/diciembre por "
                       "año, hasta el mes anterior al actual).")
        else:
            for anio, faltantes in faltantes_por_anio.items():
                st.warning(f"⚠️ {anio}: faltan {len(faltantes)} comprobante(s) — {', '.join(faltantes)}")
            st.caption("Subilos en 📄 Cargar Colillas de Pago si tenés el PDF, o revisá si el nombre del "
                       "período quedó distinto al esperado (ver el chequeo de abajo, por si una Prima quedó "
                       "etiquetada como si fuera la quincena que falta).")

    st.divider()

    st.subheader("🎁 Primas mal etiquetadas como quincena")
    df_dev = db.read_colillas_devengos()
    if df_dev.empty:
        st.info("Todavía no hay devengos de colillas cargados.")
    else:
        es_prima = df_dev["Categoría"] == "Prima de Servicios"
        es_quincena = ~df_dev["Quincena"].astype(str).str.strip().str.lower().str.startswith("prima")
        periodos_sospechosos = sorted(df_dev.loc[es_prima & es_quincena, "Quincena"].unique().tolist())
        if not periodos_sospechosos:
            st.success("✅ Ninguna 'quincena' trae un devengo de Prima de Servicios.")
        else:
            st.warning(f"⚠️ {len(periodos_sospechosos)} período(s) etiquetados como quincena traen un devengo "
                       f"de Prima de Servicios: {', '.join(periodos_sospechosos)}. Es probable que ese "
                       "comprobante sea en realidad la Prima completa (el Hospital usa el mismo encabezado "
                       "'Quincena De...' para los dos) y esté pisando el lugar de la quincena real. El botón "
                       "de abajo conserva los valores y renombra el comprobante como 'Prima {mes}-{año}'.")
            if st.button(f"🔧 Corregir {len(periodos_sospechosos)} Prima(s) automáticamente", type="primary",
                         key="btn_corregir_primas"):
                resultados, errores = [], []
                for p in periodos_sospechosos:
                    try:
                        resultados.append(db.corregir_prima_mal_etiquetada(p))
                    except Exception as e:
                        errores.append((p, str(e)))
                db.clear_read_cache()
                for r in resultados:
                    st.success(f"'{r['periodo_viejo']}' → '{r['periodo_nuevo']}' "
                               f"({r['devengos']} devengo(s), {r['descuentos']} descuento(s))")
                for p, e in errores:
                    st.error(f"No pude corregir '{p}': {e}")
                if resultados:
                    st.info("Revisá '🧾 Comprobantes de colillas faltantes' más arriba — la quincena real de "
                            "ese período probablemente ahora aparezca como faltante; subí su PDF si lo tenés.")
                st.rerun()

    st.divider()

    st.subheader("💰 Cesantías mal etiquetadas como quincena")
    if df_dev.empty:
        st.info("Todavía no hay devengos de colillas cargados.")
    else:
        es_cesantias = df_dev["Concepto"] == "Cesantías Año Anterior"
        no_etiquetado_cesantias = ~df_dev["Quincena"].astype(str).str.strip().str.lower().str.startswith("cesantías")
        periodos_cesantias_mal = sorted(df_dev.loc[es_cesantias & no_etiquetado_cesantias, "Quincena"].unique().tolist())
        if not periodos_cesantias_mal:
            st.success("✅ Ninguna 'quincena' trae un devengo de Cesantías Año Anterior.")
        else:
            st.warning(f"⚠️ {len(periodos_cesantias_mal)} período(s) etiquetados como quincena traen en realidad "
                       f"la liquidación de cesantías: {', '.join(periodos_cesantias_mal)}. Mismo tipo de bug que "
                       "la Prima (el Hospital reusa el encabezado 'Quincena De...' para los dos), ya arreglado "
                       "en el parser — el botón de abajo corrige lo que ya estaba cargado de antes: relabels el "
                       "período a 'Cesantías {mes}-{año}', saca el capital de Colillas de Pago (no tiene efecto "
                       "en caja) y manda los intereses, si los hay, a Otros Ingresos como ingreso real.")
            if st.button(f"🔧 Corregir {len(periodos_cesantias_mal)} período(s) automáticamente", type="primary",
                         key="btn_corregir_cesantias"):
                resultados, errores = [], []
                for p in periodos_cesantias_mal:
                    try:
                        resultados.append(db.corregir_cesantias_mal_etiquetada(p))
                    except Exception as e:
                        errores.append((p, str(e)))
                db.clear_read_cache()
                for r in resultados:
                    detalle_interes = f" (intereses {fmt_moneda(r['interes'])} → Otros Ingresos)" if r["interes"] else ""
                    st.success(f"'{r['periodo_viejo']}' → '{r['periodo_nuevo']}'{detalle_interes}")
                for p, e in errores:
                    st.error(f"No pude corregir '{p}': {e}")
                if resultados:
                    st.info("Revisá '🧾 Comprobantes de colillas faltantes' más arriba — la quincena real de "
                            "ese período probablemente ahora aparezca como faltante; subí su PDF si lo tenés.")
                st.rerun()

    st.divider()

    frames_gastos = []
    for bk in ("efectivo_detalle", "visa_detalle", "mc_detalle"):
        df_b = db.read_egreso_detalle(bk)
        if not df_b.empty:
            frames_gastos.append(df_b[df_b["Comercio / Concepto"].astype(str).str.strip() != ""])
    df_gastos_all = pd.concat(frames_gastos, ignore_index=True) if frames_gastos else pd.DataFrame()

    st.subheader("🗂️ Gasto sin categorizar ('Otros')")
    if df_gastos_all.empty:
        st.info("Todavía no hay gastos cargados.")
    else:
        df_gastos_pres = df_gastos_all[(df_gastos_all["Moneda"] == "COP") & (df_gastos_all["Presupuestar"] == "Sí")]
        total_gasto = df_gastos_pres["Valor Cargado Este Periodo"].sum()
        df_otros = df_gastos_pres[df_gastos_pres["Categoría"] == "Otros"]
        total_otros = df_otros["Valor Cargado Este Periodo"].sum()
        pct = (total_otros / total_gasto * 100) if total_gasto else 0
        c1, c2 = st.columns(2)
        c1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(total_otros)}</div>'
                    f'<div class="stat-label">Total en "Otros"</div></div>', unsafe_allow_html=True)
        c2.markdown(f'<div class="stat-card"><div class="stat-num">{pct:.0f}%</div>'
                    f'<div class="stat-label">% del gasto real total (efectivo+tarjetas, COP)</div></div>',
                    unsafe_allow_html=True)

        mostrar_todos = st.checkbox("Mostrar todos los comercios (no solo 'Otros') para recategorizar cualquiera",
                                     key="chk_mostrar_todos_comercios")
        base = df_gastos_pres if mostrar_todos else df_otros

        if base.empty:
            if mostrar_todos:
                st.info("Todavía no hay gastos cargados en pesos presupuestables.")
        else:
            # Filtro año/mes — solo para encontrar más rápido qué comercio
            # tocar; la recategorización en sí sigue siendo por comercio en
            # TODO el historial (ver aviso más abajo), no solo el mes
            # filtrado acá.
            anios_meses_g = [_extraer_anio_mes(f) for f in base["Fecha Compra"]]
            base = base.assign(_anio=[a for a, _ in anios_meses_g], _mes=[m for _, m in anios_meses_g])
            anios_g = ["(todos)"] + sorted({str(a) for a in base["_anio"] if a}, reverse=True)
            meses_g = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]
            cfg1, cfg2 = st.columns(2)
            anio_sel_g = cfg1.selectbox("Año", anios_g, key="anio_recat_gastos")
            mes_sel_g = cfg2.selectbox("Mes", meses_g, key="mes_recat_gastos")
            if anio_sel_g != "(todos)":
                base = base[base["_anio"] == int(anio_sel_g)]
            if mes_sel_g != "(todos)":
                base = base[base["_mes"] == int(mes_sel_g[:2])]
            base = base.drop(columns=["_anio", "_mes"])

            if base.empty:
                st.info("No hay movimientos con ese filtro de año/mes.")
            else:
                top = (base.groupby("Comercio / Concepto")
                       .agg(Total=("Valor Cargado Este Periodo", "sum"), **{"N° compras": ("Valor Cargado Este Periodo", "count")},
                            **({"Categoría Actual": ("Categoría", "first")} if mostrar_todos else {}))
                       .reset_index().sort_values("Total", ascending=False))
                sin_cambiar = "(sin cambiar)"
                top.insert(len(top.columns), "Nueva Categoría", sin_cambiar)
                aviso_alcance = (" La categoría nueva se aplica a TODAS las compras de ese comercio en "
                                 "cualquier mes, no solo las del filtro de año/mes de arriba." if
                                 (anio_sel_g != "(todos)" or mes_sel_g != "(todos)") else "")
                if mostrar_todos:
                    st.caption(f"{len(top)} comercio(s) en total — elegí una categoría nueva en la última "
                               "columna para los que quieras corregir (sin importar cuál tengan hoy), y guardá "
                               f"con el botón de abajo. Los que dejes en '(sin cambiar)' quedan sin "
                               f"tocar.{aviso_alcance}")
                else:
                    st.caption(f"{len(top)} comercio(s) distintos en 'Otros' — elegí una categoría nueva en la "
                               "última columna para los que quieras corregir, y guardá con el botón de abajo. "
                               f"Los que dejes en '(sin cambiar)' quedan sin tocar.{aviso_alcance}")
                categorias_gasto = db.read_categorias_gasto()
                disabled_cols = (["Comercio / Concepto", "Total", "N° compras"]
                                 + (["Categoría Actual"] if mostrar_todos else []))
                top_editado = st.data_editor(
                    top, hide_index=True, use_container_width=True, height=min(500, 45 + 35 * len(top)),
                    disabled=disabled_cols, key="editor_recategorizar_otros",
                    column_config={
                        "Total": st.column_config.NumberColumn(format="$%d"),
                        "Nueva Categoría": st.column_config.SelectboxColumn(options=[sin_cambiar] + categorias_gasto),
                    })
                mapeo = {fila["Comercio / Concepto"]: fila["Nueva Categoría"]
                         for _, fila in top_editado.iterrows() if fila["Nueva Categoría"] != sin_cambiar}
                if mapeo and st.button(f"💾 Aplicar {len(mapeo)} categoría(s) nueva(s)", type="primary",
                                        key="btn_recategorizar_otros"):
                    try:
                        resultado = db.recategorizar_comercios(mapeo, solo_otros=not mostrar_todos)
                        db.clear_read_cache()
                        total_filas = sum(resultado.values())
                        st.success(f"{total_filas} fila(s) recategorizadas — "
                                   f"{resultado['efectivo_detalle']} en Efectivo, {resultado['visa_detalle']} en "
                                   f"Visa, {resultado['mc_detalle']} en Mastercard.")
                        st.rerun()
                    except Exception as e:
                        st.error(f"No pude guardar: {e}")

    st.divider()

    st.subheader("⚠️ Categoría inconsistente por comercio (gastos)")
    if df_gastos_all.empty:
        st.info("Todavía no hay gastos cargados.")
    else:
        variab = df_gastos_all.groupby("Comercio / Concepto")["Categoría"].nunique()
        comercios_incons = variab[variab > 1].index.tolist()
        if not comercios_incons:
            st.success("✅ Cada comercio de gastos usa siempre la misma categoría.")
        else:
            filas = []
            for com in comercios_incons:
                sub = df_gastos_all[df_gastos_all["Comercio / Concepto"] == com]
                cats = sub.groupby("Categoría").size().to_dict()
                filas.append({"Comercio / Concepto": com,
                               "Categorías usadas": ", ".join(f"{c} ({n})" for c, n in cats.items())})
            df_incons = pd.DataFrame(filas).sort_values("Comercio / Concepto")
            st.warning(f"⚠️ {len(comercios_incons)} comercio(s) de gastos tienen más de una categoría asignada.")
            st.dataframe(df_incons, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_incons)))

    st.divider()

    st.subheader("🔄 Recategorizar Otros Ingresos")
    df_oi_todo = db.read_otros_ingresos_tabla()
    df_oi_todo = (df_oi_todo[df_oi_todo["Concepto"].astype(str).str.strip() != ""]
                  if not df_oi_todo.empty else df_oi_todo)
    if df_oi_todo.empty:
        st.info("Todavía no hay Otros Ingresos cargados.")
    else:
        solo_otro = st.checkbox("Mostrar solo los conceptos en la categoría genérica 'Otro'",
                                 value=True, key="chk_solo_otro_ingresos")
        base_oi = df_oi_todo[df_oi_todo["Categoría"] == "Otro"] if solo_otro else df_oi_todo
        if base_oi.empty:
            st.success("✅ Ningún concepto está en la categoría genérica 'Otro'.")
        else:
            # Filtro año/mes — solo para encontrar más rápido qué concepto
            # tocar; la recategorización en sí sigue siendo por concepto en
            # TODO el historial (ver aviso más abajo), no solo el mes filtrado acá.
            anios_meses_oi = [_extraer_anio_mes(f) for f in base_oi["Fecha"]]
            base_oi = base_oi.assign(_anio=[a for a, _ in anios_meses_oi], _mes=[m for _, m in anios_meses_oi])
            anios_oi = ["(todos)"] + sorted({str(a) for a in base_oi["_anio"] if a}, reverse=True)
            meses_oi = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]
            cfg1_oi, cfg2_oi = st.columns(2)
            anio_sel_oi = cfg1_oi.selectbox("Año", anios_oi, key="anio_recat_oi")
            mes_sel_oi = cfg2_oi.selectbox("Mes", meses_oi, key="mes_recat_oi")
            if anio_sel_oi != "(todos)":
                base_oi = base_oi[base_oi["_anio"] == int(anio_sel_oi)]
            if mes_sel_oi != "(todos)":
                base_oi = base_oi[base_oi["_mes"] == int(mes_sel_oi[:2])]
            base_oi = base_oi.drop(columns=["_anio", "_mes"])

            if base_oi.empty:
                st.info("No hay movimientos con ese filtro de año/mes.")
            else:
                top_oi = (base_oi.groupby("Concepto")
                          .agg(Total=("Valor", "sum"), **{"N° movimientos": ("Valor", "count")},
                               **{"Categoría Actual": ("Categoría", "first")})
                          .reset_index().sort_values("Total", ascending=False))
                sin_cambiar_oi = "(sin cambiar)"
                top_oi.insert(len(top_oi.columns), "Nueva Categoría", sin_cambiar_oi)
                aviso_alcance_oi = (" La categoría nueva se aplica a TODOS los movimientos de ese concepto en "
                                     "cualquier mes, no solo las del filtro de año/mes de arriba." if
                                     (anio_sel_oi != "(todos)" or mes_sel_oi != "(todos)") else "")
                st.caption(f"{len(top_oi)} concepto(s) — elegí una categoría nueva en la última columna para los "
                           "que quieras corregir (sin importar cuál tengan hoy), y guardá con el botón de abajo. "
                           f"Los que dejes en '(sin cambiar)' quedan sin tocar.{aviso_alcance_oi}")
                top_oi_editado = st.data_editor(
                    top_oi, hide_index=True, use_container_width=True, height=min(500, 45 + 35 * len(top_oi)),
                    disabled=["Concepto", "Total", "N° movimientos", "Categoría Actual"],
                    key="editor_recategorizar_oi",
                    column_config={
                        "Total": st.column_config.NumberColumn(format="$%d"),
                        "Nueva Categoría": st.column_config.SelectboxColumn(
                            options=[sin_cambiar_oi] + CATEGORIAS_OTROS_INGRESOS),
                    })
                mapeo_oi = {fila["Concepto"]: fila["Nueva Categoría"]
                            for _, fila in top_oi_editado.iterrows() if fila["Nueva Categoría"] != sin_cambiar_oi}
                if mapeo_oi and st.button(f"💾 Aplicar {len(mapeo_oi)} categoría(s) nueva(s)", type="primary",
                                           key="btn_recategorizar_oi"):
                    try:
                        total_filas_oi = db.recategorizar_conceptos_ingreso(mapeo_oi)
                        db.clear_read_cache()
                        st.success(f"{total_filas_oi} fila(s) recategorizadas en Otros Ingresos.")
                        st.rerun()
                    except Exception as e:
                        st.error(f"No pude guardar: {e}")

    st.divider()

    st.subheader("⚠️ Categoría inconsistente por concepto (Otros Ingresos)")
    df_oi = db.read_otros_ingresos_tabla()
    df_oi = df_oi[df_oi["Concepto"].astype(str).str.strip() != ""] if not df_oi.empty else df_oi
    if df_oi.empty:
        st.info("Todavía no hay Otros Ingresos cargados.")
    else:
        variab_oi = df_oi.groupby("Concepto")["Categoría"].nunique()
        conceptos_incons = variab_oi[variab_oi > 1].index.tolist()
        if not conceptos_incons:
            st.success("✅ Cada concepto de Otros Ingresos usa siempre la misma categoría.")
        else:
            filas = []
            for con in conceptos_incons:
                sub = df_oi[df_oi["Concepto"] == con]
                cats = sub.groupby("Categoría").size().to_dict()
                filas.append({"Concepto": con,
                               "Categorías usadas": ", ".join(f"{c} ({n})" for c, n in cats.items())})
            df_incons_oi = pd.DataFrame(filas).sort_values("Concepto")
            st.warning(f"⚠️ {len(conceptos_incons)} concepto(s) de Otros Ingresos tienen más de una categoría "
                       "asignada.")
            st.dataframe(df_incons_oi, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_incons_oi)))


def render_presupuesto():
    st.title("📋 Presupuesto")
    st.caption("Definí una meta mensual por categoría y comparala contra el gasto real de un mes puntual — "
               "igual que Balance Mensual, el efectivo cuenta en su propio mes y las tarjetas en el mes "
               "siguiente al extracto, cuando se pagan. Los descuentos de nómina (Ahorro, Seguros, Impuestos, "
               "cuotas de deuda...) se comparan aparte, abajo, porque no son gasto discrecional.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Elegís un mes, escribís una meta de gasto para cada categoría (y para cada descuento de nómina) "
            "en el formulario, y las guardás con el botón de abajo — quedan escritas directo en la hoja "
            "'Presupuesto' del Sheet. El botón '💡 Sugerir metas' completa el formulario automáticamente con "
            "el promedio de gasto real de los últimos meses con datos, para que solo tengas que ajustar. "
            "Debajo del formulario aparece un gráfico comparando presupuesto vs. "
            "gasto real para las categorías que ya tengan una meta definida.")

    hoy = date.today()
    anios_pres = list(range(2023, 2033))
    c1, c2 = st.columns(2)
    anio_pres = c1.selectbox("Año", anios_pres,
                              index=anios_pres.index(hoy.year) if hoy.year in anios_pres else 0,
                              key="anio_presupuesto")
    mes_pres_num = c2.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                 format_func=lambda m: MESES_NOMBRE[m], key="mes_presupuesto")
    mes_pres = f"{anio_pres}-{mes_pres_num:02d}"

    if st.session_state.get("_presupuesto_mes_aplicado") != mes_pres:
        try:
            try:
                db.set_presupuesto_mes(mes_pres)
            except Exception:
                time.sleep(2)  # reintento único — la API de Sheets a veces falla de forma transitoria
                db.set_presupuesto_mes(mes_pres)
            db.clear_read_cache()
            st.session_state["_presupuesto_mes_aplicado"] = mes_pres
        except Exception as e:
            st.error(f"No pude cambiar el mes en la hoja 'Presupuesto': {e}. Probá refrescar la página.")
            st.stop()

    p = db.read_presupuesto()

    if st.button("💡 Sugerir metas según el promedio de los últimos meses con datos", key="btn_sugerir_metas"):
        hoy_sug = date.today()
        mes_actual_sug = f"{hoy_sug.year}-{hoy_sug.month:02d}"
        meses_prev = [_shift_mes(mes_actual_sug, -i) for i in range(1, 4)]
        cat_por_mes = [_gasto_real_categoria_mes(m) for m in meses_prev]
        desc_por_mes = [_descuentos_categoria_mes(m) for m in meses_prev]

        for cat in p["categorias"]:
            valores_no_cero = [d.get(cat["categoria"], 0.0) for d in cat_por_mes if d.get(cat["categoria"], 0.0) > 0]
            if valores_no_cero:
                st.session_state[f"meta_cat_{cat['fila']}"] = round(sum(valores_no_cero) / len(valores_no_cero), -3)
        for desc in p["descuentos"]:
            valores_no_cero = [d.get(desc["categoria"], 0.0) for d in desc_por_mes
                                if d.get(desc["categoria"], 0.0) > 0]
            if valores_no_cero:
                st.session_state[f"meta_desc_{desc['fila']}"] = round(sum(valores_no_cero) / len(valores_no_cero), -3)
        st.success(f"Metas sugeridas con el promedio de {', '.join(meses_prev)} (donde hubo datos) — "
                   "revisalas y ajustalas antes de guardar.")
        st.rerun()

    with st.form("form_presupuesto"):
        st.markdown("#### Metas por categoría de gasto")
        entradas = {}
        for cat in p["categorias"]:
            cc1, cc2, cc3, cc4 = st.columns([3, 2, 2, 2])
            cc1.markdown(cat["categoria"])
            cc2.caption(f"Gasto real: {fmt_moneda(cat['gasto_real'])}")
            entradas[cat["fila"]] = cc3.number_input(
                "Meta", value=float(cat["presupuesto"]), step=10000.0, format="%.0f",
                key=f"meta_cat_{cat['fila']}", label_visibility="collapsed")
            meta_val = entradas[cat["fila"]]
            cc4.caption(f"{fmt_moneda(meta_val - cat['gasto_real'])} disponible" if meta_val > 0 else "—")

        st.divider()
        st.markdown("#### Metas de descuentos de nómina")
        st.caption("Ahorro, Seguros, Impuestos, Aportes de Ley y las cuotas de deuda que se descuentan directo "
                   "de la colilla — no son gasto discrecional, pero sí reducen lo que te queda disponible.")
        for desc in p["descuentos"]:
            cc1, cc2, cc3, cc4 = st.columns([3, 2, 2, 2])
            cc1.markdown(desc["categoria"])
            cc2.caption(f"Real: {fmt_moneda(desc['real'])}")
            entradas[desc["fila"]] = cc3.number_input(
                "Meta", value=float(desc["presupuesto"]), step=10000.0, format="%.0f",
                key=f"meta_desc_{desc['fila']}", label_visibility="collapsed")
            meta_val = entradas[desc["fila"]]
            cc4.caption(f"{fmt_moneda(meta_val - desc['real'])} disponible" if meta_val > 0 else "—")

        guardar = st.form_submit_button("💾 Guardar todas las metas", type="primary")

    if guardar:
        try:
            db.set_presupuesto_metas(entradas)
            db.clear_read_cache()
            st.success("Metas guardadas.")
            st.rerun()
        except Exception as e:
            st.error(f"No pude guardar las metas: {e}")

    con_meta = [c for c in p["categorias"] if c["presupuesto"] > 0]
    if con_meta:
        st.divider()
        st.markdown("#### Presupuesto vs. gasto real (categorías con meta definida)")
        df_comp = pd.DataFrame(con_meta)
        df_long = df_comp.melt(id_vars=["categoria"], value_vars=["presupuesto", "gasto_real"],
                                var_name="Concepto", value_name="Valor")
        df_long["Concepto"] = df_long["Concepto"].map({"presupuesto": "Presupuesto", "gasto_real": "Gasto Real"})
        fig_pres = px.bar(df_long, x="categoria", y="Valor", color="Concepto", barmode="group")
        fig_pres.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=340, xaxis_title="")
        st.plotly_chart(fig_pres, use_container_width=True)
    else:
        st.info("Definí al menos una meta arriba (y guardala) para ver el gráfico de comparación.")


def _form_agregar_colilla():
    with st.expander("➕ Agregar una quincena manualmente"):
        st.caption("Alternativa a subir el PDF: escribí la fecha de pago, la quincena, y cada devengo y "
                   "descuento con su categoría y valor — agregá o quitá filas con los botones de la tabla.")
        hoy = date.today()
        anios_q = list(range(2023, 2033))
        c1, c2, c3 = st.columns(3)
        anio_q = c1.selectbox("Año", anios_q, index=anios_q.index(hoy.year) if hoy.year in anios_q else 0,
                               key="anio_colilla_manual")
        mes_q = c2.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                              format_func=lambda m: MESES_NOMBRE[m], key="mes_colilla_manual")
        quincena_q = c3.selectbox("Quincena", ["1a", "2a"], key="quincena_colilla_manual")
        fecha_pago = st.date_input("Fecha de pago", value=hoy, key="fecha_pago_colilla_manual")
        periodo = f"{quincena_q} quincena {_MESES_3LETRAS[mes_q - 1]}-{anio_q}"
        st.caption(f"Periodo: **{periodo}**")

        st.markdown("**Devengos**")
        df_dev_edit = st.data_editor(
            pd.DataFrame([{"Concepto": "", "Categoría": "Salario Base", "Valor": 0.0}]),
            num_rows="dynamic", hide_index=True, key="devengos_colilla_manual", use_container_width=True,
            column_config={
                "Categoría": st.column_config.SelectboxColumn(options=DEVENGOS_CATEGORIAS_MANUAL, required=True),
                "Valor": st.column_config.NumberColumn(format="%.0f")})

        st.markdown("**Descuentos**")
        df_desc_edit = st.data_editor(
            pd.DataFrame([{"Concepto": "", "Categoría": "Ahorro", "Valor": 0.0}]),
            num_rows="dynamic", hide_index=True, key="descuentos_colilla_manual", use_container_width=True,
            column_config={
                "Categoría": st.column_config.SelectboxColumn(options=DESCUENTOS_CATEGORIAS_MANUAL, required=True),
                "Valor": st.column_config.NumberColumn(format="%.0f")})

        if st.button("💾 Guardar quincena", type="primary", key="btn_guardar_colilla_manual"):
            devengos_validos = [r for r in df_dev_edit.to_dict("records")
                                 if str(r.get("Concepto") or "").strip() and (r.get("Valor") or 0) > 0]
            descuentos_validos = [r for r in df_desc_edit.to_dict("records")
                                   if str(r.get("Concepto") or "").strip() and (r.get("Valor") or 0) > 0]
            df_cp_existente = db.read_colillas_resumen()
            if not df_cp_existente.empty and periodo in df_cp_existente["Periodo"].values:
                st.error(f"Ya existe una quincena para '{periodo}' — editala directo en el Sheet si necesitás "
                         "corregirla, no puedo agregar otra para el mismo período.")
            elif not devengos_validos and not descuentos_validos:
                st.error("Agregá al menos un devengo o un descuento con concepto y valor.")
            else:
                total_devengos = sum(r["Valor"] for r in devengos_validos)
                total_descuentos = sum(r["Valor"] for r in descuentos_validos)
                try:
                    rr = db.first_blank_row("colillas_resumen")
                    db.write_block_rows("colillas_resumen", rr,
                                         [[fecha_pago.isoformat(), periodo, total_devengos, total_descuentos]])
                    if devengos_validos:
                        dev_row = db.first_blank_row("colillas_devengos")
                        rows = [[periodo, str(r["Concepto"]).strip(), r["Categoría"], r["Valor"]]
                                for r in devengos_validos]
                        db.write_block_rows("colillas_devengos", dev_row, rows)
                    if descuentos_validos:
                        ded_row = db.first_blank_row("colillas_descuentos")
                        rows = [[periodo, str(r["Concepto"]).strip(), r["Categoría"], r["Valor"]]
                                for r in descuentos_validos]
                        db.write_block_rows("colillas_descuentos", ded_row, rows)
                    db.clear_read_cache()
                    st.success(f"Quincena '{periodo}' agregada.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"No pude guardar: {e}")


def _form_eliminar_colilla():
    with st.expander("🗑️ Eliminar una quincena"):
        df_cp = db.read_colillas_resumen()
        if df_cp.empty:
            st.info("Todavía no hay quincenas cargadas.")
            return
        periodos = df_cp["Periodo"].tolist()
        periodo_sel = st.selectbox("Quincena a eliminar", periodos, key="periodo_eliminar_colilla")
        fila_sel = df_cp[df_cp["Periodo"] == periodo_sel].iloc[0]
        n_dev = (db.read_colillas_devengos()["Quincena"] == periodo_sel).sum()
        n_desc = (db.read_colillas_descuentos()["Quincena"] == periodo_sel).sum()
        st.caption(f"Devengos: {fmt_moneda(fila_sel['Devengos Totales'])} ({n_dev} línea(s)) — "
                   f"Descuentos: {fmt_moneda(fila_sel['Descuentos Totales'])} ({n_desc} línea(s))")
        confirmar = st.checkbox(f"Confirmo que quiero borrar '{periodo_sel}' — no se puede deshacer",
                                 key="confirmar_eliminar_colilla")
        if st.button("🗑️ Eliminar quincena", type="primary", disabled=not confirmar, key="btn_eliminar_colilla"):
            try:
                resultado = db.eliminar_colilla(periodo_sel)
                db.clear_read_cache()
                st.success(f"'{periodo_sel}' eliminada — {resultado['resumen']} fila(s) de resumen, "
                           f"{resultado['devengos']} de devengos, {resultado['descuentos']} de descuentos.")
                st.rerun()
            except Exception as e:
                st.error(f"No pude borrar: {e}")


def render_ingresos_colillas():
    st.markdown("#### Colillas de Pago")
    _form_agregar_colilla()
    _form_eliminar_colilla()
    df_cp = db.read_colillas_resumen()
    if not df_cp.empty:
        anios_meses_cp = [_extraer_anio_mes(p) for p in df_cp["Periodo"]]
        df_cp = df_cp.assign(_anio=[a for a, _ in anios_meses_cp], _mes=[m for _, m in anios_meses_cp])

        anios_cp = ["(todos)"] + sorted({str(a) for a in df_cp["_anio"] if a}, reverse=True)
        meses_cp = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]
        c1, c2 = st.columns(2)
        anio_sel_cp = c1.selectbox("Año", anios_cp, key="anio_colillas")
        mes_sel_cp = c2.selectbox("Mes", meses_cp, key="mes_colillas")

        c3, c4 = st.columns(2)
        c3.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(df_cp["Devengos Totales"].sum())}</div>'
                    f'<div class="stat-label">Devengos totales (histórico)</div></div>', unsafe_allow_html=True)
        c4.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(df_cp["Descuentos Totales"].sum())}</div>'
                    f'<div class="stat-label">Descuentos totales (histórico)</div></div>', unsafe_allow_html=True)

        st.markdown("**Tendencia por quincena** (últimas 24)")
        df_tend_cp = df_cp.tail(24)
        df_long_cp = df_tend_cp.melt(id_vars=["Periodo"], value_vars=["Devengos Totales", "Descuentos Totales"],
                                      var_name="Concepto", value_name="Valor")
        fig_cp = px.bar(df_long_cp, x="Periodo", y="Valor", color="Concepto", barmode="group")
        fig_cp.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=360, xaxis_title="")
        st.plotly_chart(fig_cp, use_container_width=True)

        if anio_sel_cp == "(todos)" or mes_sel_cp == "(todos)":
            # Sin un mes puntual elegido: la vista general, tal como antes.
            st.markdown("**Resumen por quincena**")
            st.dataframe(df_cp.drop(columns=["_anio", "_mes"]).iloc[::-1], hide_index=True,
                         use_container_width=True, height=350)

            df_dev = db.read_colillas_devengos()
            df_desc = db.read_colillas_descuentos()
            c5, c6 = st.columns(2)
            with c5:
                st.markdown("**Devengos**")
                st.dataframe(df_dev, hide_index=True, use_container_width=True, height=400)
            with c6:
                st.markdown("**Descuentos**")
                st.dataframe(df_desc, hide_index=True, use_container_width=True, height=400)
        else:
            # Un mes puntual: el detalle "tal como aparece en la colilla",
            # quincena por quincena — igual que en la hoja de cálculo.
            mes_num_cp = int(mes_sel_cp[:2])
            df_mes_cp = df_cp[(df_cp["_anio"] == int(anio_sel_cp)) & (df_cp["_mes"] == mes_num_cp)]
            df_mes_cp = df_mes_cp.sort_values("Periodo")
            if df_mes_cp.empty:
                st.info(f"No hay colillas cargadas para {mes_sel_cp[5:]} de {anio_sel_cp}.")
            else:
                df_dev_all = db.read_colillas_devengos()
                df_desc_all = db.read_colillas_descuentos()
                for _, fila in df_mes_cp.iterrows():
                    quincena = fila["Periodo"]
                    st.markdown(f"**{quincena}** — fecha de pago: {fila['Fecha de Pago']}")
                    c5, c6 = st.columns(2)
                    c5.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(fila["Devengos Totales"])}</div>'
                                f'<div class="stat-label">Devengos</div></div>', unsafe_allow_html=True)
                    c6.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(fila["Descuentos Totales"])}</div>'
                                f'<div class="stat-label">Descuentos</div></div>', unsafe_allow_html=True)
                    c7, c8 = st.columns(2)
                    with c7:
                        st.markdown("Devengos")
                        st.dataframe(df_dev_all[df_dev_all["Quincena"] == quincena], hide_index=True,
                                     use_container_width=True, height=250)
                    with c8:
                        st.markdown("Descuentos")
                        st.dataframe(df_desc_all[df_desc_all["Quincena"] == quincena], hide_index=True,
                                     use_container_width=True, height=250)
                    st.divider()
    else:
        st.info("Todavía no hay colillas cargadas.")


def render_ingresos_otros():
    st.markdown("#### Otros Ingresos")

    with st.expander("➕ Agregar un ingreso manualmente"):
        with st.form("form_agregar_ingreso"):
            c1, c2 = st.columns(2)
            fecha_ing = c1.date_input("Fecha", value=date.today(), key="fecha_nuevo_ingreso")
            categoria_ing = c2.selectbox("Categoría", CATEGORIAS_OTROS_INGRESOS, key="categoria_nuevo_ingreso")
            concepto_ing = st.text_input("Concepto", key="concepto_nuevo_ingreso")
            valor_ing = st.number_input("Valor", min_value=0.0, step=1000.0, format="%.0f", key="valor_nuevo_ingreso")
            notas_ing = st.text_input("Notas (opcional)", key="notas_nuevo_ingreso")
            guardar_ing = st.form_submit_button("💾 Guardar ingreso", type="primary")
        if guardar_ing:
            if not concepto_ing.strip():
                st.error("Escribí un concepto.")
            elif valor_ing <= 0:
                st.error("El valor tiene que ser mayor que cero.")
            else:
                try:
                    fila = db.first_blank_row("otros_ingresos")
                    db.write_block_rows("otros_ingresos", fila,
                                         [[fecha_ing.isoformat(), concepto_ing.strip(), categoria_ing,
                                           valor_ing, notas_ing.strip()]])
                    db.clear_read_cache()
                    st.success(f"Ingreso de {fmt_moneda(valor_ing)} agregado.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"No pude guardar: {e}")

    df_oi = db.read_otros_ingresos_tabla()
    if not df_oi.empty:
        anios_meses = [_extraer_anio_mes(p) for p in df_oi["Fecha"]]
        df_oi = df_oi.assign(_anio=[a for a, _ in anios_meses], _mes=[m for _, m in anios_meses],
                              _fecha_dt=pd.to_datetime(df_oi["Fecha"], format="%d/%m/%Y", errors="coerce"))

        ocultar_np_oi = st.checkbox(
            "Ocultar movimientos de conciliación (no presupuestar)", value=True, key="ocultar_np_otros_ing",
            help="Movimientos que se cargan igual para poder conciliar el saldo real de la cuenta, pero no "
                 "cuentan en tus totales de ingreso. Desmarcá esto para verlos también.")
        agrupar_rend = st.checkbox(
            "Agrupar 'Rendimientos Financieros' por mes", value=True, key="agrupar_rendimientos_otros_ing",
            help="Los intereses de ahorros se abonan casi a diario, en montos muy chicos — agruparlos por mes "
                 "hace la tabla mucho más legible sin perder el total. Desmarcá esto para ver cada abono.")

        busqueda_oi = st.text_input("🔍 Buscar en concepto", "", key="busqueda_otros_ing")
        categorias_oi = ["(todas)"] + sorted([c for c in df_oi["Categoría"].unique().tolist() if c])
        anios_oi = ["(todos)"] + sorted({str(a) for a in df_oi["_anio"] if a}, reverse=True)
        meses_oi = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]

        c1, c2, c3 = st.columns(3)
        cat_sel_oi = c1.selectbox("Categoría", categorias_oi, key="cat_otros_ing")
        anio_sel_oi = c2.selectbox("Año", anios_oi, key="anio_otros_ing")
        mes_sel_oi = c3.selectbox("Mes", meses_oi, key="mes_otros_ing")

        df_f = df_oi
        if ocultar_np_oi:
            df_f = df_f[df_f["Presupuestar"] == "Sí"]
        if busqueda_oi:
            df_f = df_f[df_f["Concepto"].str.contains(busqueda_oi, case=False, na=False)]
        if cat_sel_oi != "(todas)":
            df_f = df_f[df_f["Categoría"] == cat_sel_oi]
        if anio_sel_oi != "(todos)":
            df_f = df_f[df_f["_anio"] == int(anio_sel_oi)]
        if mes_sel_oi != "(todos)":
            df_f = df_f[df_f["_mes"] == int(mes_sel_oi[:2])]

        st.caption(f"{len(df_f):,} de {len(df_oi):,} movimientos — suma: {fmt_moneda(df_f['Valor'].sum())}")

        if not df_f.empty:
            st.markdown("**Total por categoría** (según el filtro de arriba)")
            resumen_cat = df_f.groupby("Categoría")["Valor"].sum().sort_values(ascending=False)
            cols_cat = st.columns(len(resumen_cat))
            for col, (categoria, valor) in zip(cols_cat, resumen_cat.items()):
                col.metric(categoria, f"{fmt_moneda(valor)}")

        df_mostrar = df_f
        if agrupar_rend:
            es_rend = df_mostrar["Categoría"] == "Rendimientos Financieros"
            df_rend, df_resto = df_mostrar[es_rend], df_mostrar[~es_rend]
            if not df_rend.empty:
                resumen_rend = (df_rend.groupby(["_anio", "_mes"])
                                 .agg(Valor=("Valor", "sum"), _n=("Valor", "size"),
                                      _fecha_dt=("_fecha_dt", "max"))
                                 .reset_index())
                resumen_rend["Fecha"] = resumen_rend["_fecha_dt"].dt.strftime("%d/%m/%Y")
                resumen_rend["Concepto"] = "Rendimientos Financieros (resumen del mes)"
                resumen_rend["Categoría"] = "Rendimientos Financieros"
                resumen_rend["Notas"] = resumen_rend["_n"].map(lambda n: f"{n} abono(s) de interés este mes")
                resumen_rend["Presupuestar"] = "Sí"
                df_mostrar = pd.concat([df_resto, resumen_rend], ignore_index=True)

        df_mostrar = df_mostrar.sort_values("_fecha_dt", ascending=False)
        columnas_tabla = ["Fecha", "Concepto", "Categoría", "Valor", "Notas", "Presupuestar"]
        st.dataframe(df_mostrar[columnas_tabla], hide_index=True, use_container_width=True, height=600)

        if not df_f.empty:
            with st.expander("🗑️ Eliminar un ingreso"):
                st.caption("Para limpiar un duplicado puntual (p. ej. una corrección que corrió dos veces) — "
                           "elegí la fila exacta de la lista filtrada de arriba y borrala. Si hay dos filas "
                           "idénticas, esto borra una sola por vez; repetí para borrar la otra.")
                vista_borrar = df_f.sort_values("_fecha_dt", ascending=False).reset_index(drop=True)
                etiquetas_oi = [f"{i}: {fila['Fecha']} — {fila['Concepto']} — {fmt_moneda(fila['Valor'])}"
                                for i, fila in vista_borrar.iterrows()]
                seleccion_oi = st.selectbox("Ingreso a eliminar", etiquetas_oi, key="seleccion_eliminar_oi")
                if seleccion_oi:
                    fila_sel = vista_borrar.iloc[int(seleccion_oi.split(":", 1)[0])]
                    confirmar_oi = st.checkbox("Confirmo que quiero borrar esta fila — no se puede deshacer",
                                               key="confirmar_eliminar_oi")
                    if st.button("🗑️ Eliminar ingreso", type="primary", disabled=not confirmar_oi,
                                 key="btn_eliminar_oi"):
                        try:
                            borrado = db.eliminar_otro_ingreso(fila_sel["Fecha"], fila_sel["Concepto"],
                                                               fila_sel["Categoría"], float(fila_sel["Valor"]))
                            db.clear_read_cache()
                            if borrado:
                                st.success("Ingreso eliminado.")
                            else:
                                st.warning("No encontré esa fila exacta en el Sheet — puede que ya se haya borrado.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"No pude borrar: {e}")

        if not df_f.empty:
            st.markdown("**Ingresos por categoría** (según el filtro de arriba)")
            df_cat_oi = (df_f.groupby("Categoría")["Valor"].sum().reset_index()
                         .sort_values("Valor", ascending=False))
            fig_oi = px.bar(df_cat_oi, x="Valor", y="Categoría", orientation="h")
            fig_oi.update_layout(yaxis=dict(autorange="reversed"), margin=dict(l=0, r=0, t=10, b=0),
                                 height=max(280, 32 * len(df_cat_oi)))
            st.plotly_chart(fig_oi, use_container_width=True)
    else:
        st.info("Todavía no hay otros ingresos cargados.")


def render_egresos_efectivo():
    with st.expander("➕ Agregar un gasto en efectivo manualmente"):
        with st.form("form_agregar_gasto_efectivo"):
            c1, c2 = st.columns(2)
            fecha_gto = c1.date_input("Fecha de compra", value=date.today(), key="fecha_nuevo_gasto_efectivo")
            categoria_gto = c2.selectbox("Categoría", db.read_categorias_gasto(), key="categoria_nuevo_gasto_efectivo")
            comercio_gto = st.text_input("Comercio / Concepto", key="comercio_nuevo_gasto_efectivo")
            c3, c4 = st.columns(2)
            valor_gto = c3.number_input("Valor", min_value=0.0, step=1000.0, format="%.0f",
                                         key="valor_nuevo_gasto_efectivo")
            reembolsable_gto = c4.checkbox("Reembolsable (es en realidad gasto de tu esposa)",
                                            key="reembolsable_nuevo_gasto_efectivo")
            notas_gto = st.text_input("Notas (opcional)", key="notas_nuevo_gasto_efectivo")
            guardar_gto = st.form_submit_button("💾 Guardar gasto", type="primary")
        if guardar_gto:
            if not comercio_gto.strip():
                st.error("Escribí un comercio o concepto.")
            elif valor_gto <= 0:
                st.error("El valor tiene que ser mayor que cero.")
            else:
                try:
                    fila = db.first_blank_row("efectivo_detalle")
                    db.write_block_rows("efectivo_detalle", fila,
                                         [[db.as_text(fecha_gto.strftime("%Y-%m")), fecha_gto.isoformat(),
                                           comercio_gto.strip(), "COP", db.as_text("1/1"), valor_gto, valor_gto, 0,
                                           categoria_gto, "Sí" if reembolsable_gto else "No", notas_gto.strip()]])
                    db.clear_read_cache()
                    st.success(f"Gasto de {fmt_moneda(valor_gto)} agregado.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"No pude guardar: {e}")

    _render_egreso_tab("efectivo_detalle", "Egresos - Efectivo")


def _form_agregar_extracto_resumen(tarjeta_label, key_sufijo):
    blocks = db.TARJETA_BLOCKS[tarjeta_label]
    resumen_block = blocks["resumen"]
    tiene_usd = "detalle_usd" in blocks
    st.caption("Un renglón por período de corte — si ese período ya tiene un resumen cargado, no se puede "
               "agregar otro (editalo directo en el Sheet si necesitás corregirlo).")
    anios_ext = list(range(2023, 2033))
    with st.form(f"form_agregar_extracto_{key_sufijo}"):
        hoy = date.today()
        c1, c2 = st.columns(2)
        anio_ext = c1.selectbox("Año del corte", anios_ext,
                                 index=anios_ext.index(hoy.year) if hoy.year in anios_ext else 0,
                                 key=f"anio_ext_{key_sufijo}")
        mes_ext = c2.selectbox("Mes del corte", list(range(1, 13)), index=hoy.month - 1,
                                format_func=lambda m: MESES_NOMBRE[m], key=f"mes_ext_{key_sufijo}")
        c3, c4 = st.columns(2)
        fecha_corte = c3.date_input("Fecha de Corte", value=hoy, key=f"fcorte_{key_sufijo}")
        fecha_limite = c4.date_input("Fecha Límite de Pago", value=hoy, key=f"flimite_{key_sufijo}")
        c5, c6 = st.columns(2)
        cupo_total = c5.number_input("Cupo Total", min_value=0.0, step=100000.0, format="%.0f",
                                      key=f"cupototal_{key_sufijo}")
        cupo_disponible = c6.number_input("Cupo Disponible", min_value=0.0, step=100000.0, format="%.0f",
                                           key=f"cupodisp_{key_sufijo}")
        saldo_anterior = st.number_input("Saldo Anterior", min_value=0.0, step=10000.0, format="%.0f",
                                          key=f"saldoant_{key_sufijo}")
        c7, c8 = st.columns(2)
        pago_minimo = c7.number_input("Pago Mínimo", min_value=0.0, step=10000.0, format="%.0f",
                                       key=f"pagomin_{key_sufijo}")
        pago_total = c8.number_input("Pago Total", min_value=0.0, step=10000.0, format="%.0f",
                                      key=f"pagototal_{key_sufijo}")
        saldo_usd = None
        if tiene_usd:
            saldo_usd = st.number_input("Saldo a pagar USD", min_value=0.0, step=10.0, format="%.2f",
                                         key=f"saldousd_{key_sufijo}")
        guardar_ext = st.form_submit_button("💾 Guardar extracto", type="primary")
    if guardar_ext:
        periodo = f"{anio_ext}-{mes_ext:02d}"
        df_existente = db.read_tarjeta_resumen(resumen_block)
        if not df_existente.empty and periodo in df_existente["Periodo Extracto"].values:
            st.error(f"Ya existe un resumen para el período {periodo} — editalo directo en el Sheet si "
                     "necesitás corregirlo, no puedo agregar otro para el mismo período.")
        else:
            try:
                fila = db.first_blank_row(resumen_block)
                segments = [
                    (1, [db.as_text(periodo), fecha_corte.isoformat(), fecha_limite.isoformat(),
                         cupo_total, cupo_disponible]),
                    (7, [saldo_anterior]),
                    (9, [pago_minimo, pago_total]),
                ]
                if tiene_usd and saldo_usd:
                    segments.append((11, [saldo_usd]))
                db.write_row_segments(resumen_block, fila, segments)
                db.clear_read_cache()
                st.success(f"Extracto de {periodo} agregado.")
                st.rerun()
            except db.SinEspacioError as e:
                st.error(str(e))
            except Exception as e:
                st.error(f"No pude guardar: {e}")


def _form_agregar_compra_tarjeta(block_key_cop, block_key_usd, key_sufijo):
    anios_per = list(range(2023, 2033))
    with st.form(f"form_agregar_compra_{key_sufijo}"):
        hoy = date.today()
        c1, c2 = st.columns(2)
        anio_per = c1.selectbox("Año del período (corte)", anios_per,
                                 index=anios_per.index(hoy.year) if hoy.year in anios_per else 0,
                                 key=f"anioper_{key_sufijo}")
        mes_per = c2.selectbox("Mes del período (corte)", list(range(1, 13)), index=hoy.month - 1,
                                format_func=lambda m: MESES_NOMBRE[m], key=f"mesper_{key_sufijo}")
        fecha_compra = st.date_input("Fecha de compra", value=hoy, key=f"fcompra_{key_sufijo}")
        comercio = st.text_input("Comercio / Concepto", key=f"comercio_{key_sufijo}")
        c3, c4 = st.columns(2)
        moneda = c3.selectbox("Moneda", ["COP", "USD"], key=f"moneda_{key_sufijo}") if block_key_usd else "COP"
        cuotas = c4.text_input("Cuotas", value="1/1", key=f"cuotas_{key_sufijo}")
        c5, c6 = st.columns(2)
        valor_total = c5.number_input("Valor Total Compra", min_value=0.0, step=1000.0, format="%.0f",
                                       key=f"vtotal_{key_sufijo}")
        valor_periodo = c6.number_input("Valor Cargado Este Periodo", min_value=0.0, step=1000.0,
                                         format="%.0f", key=f"vperiodo_{key_sufijo}")
        saldo_pendiente = st.number_input("Saldo Pendiente (cuotas)", min_value=0.0, step=1000.0,
                                           format="%.0f", key=f"spend_{key_sufijo}")
        c7, c8 = st.columns(2)
        categoria = c7.selectbox("Categoría", db.read_categorias_gasto(), key=f"catcompra_{key_sufijo}")
        reembolsable = c8.checkbox("Reembolsable", key=f"reembcompra_{key_sufijo}")
        notas = st.text_input("Notas (opcional)", key=f"notascompra_{key_sufijo}")
        guardar = st.form_submit_button("💾 Guardar compra", type="primary")
    if guardar:
        if not comercio.strip():
            st.error("Escribí un comercio o concepto.")
        elif valor_total <= 0:
            st.error("El valor total tiene que ser mayor que cero.")
        else:
            periodo = f"{anio_per}-{mes_per:02d}"
            block_key = block_key_usd if moneda == "USD" else block_key_cop
            fila_val = valor_periodo if valor_periodo > 0 else valor_total
            try:
                fila = db.first_blank_row(block_key)
                db.write_block_rows(block_key, fila,
                                     [[db.as_text(periodo), fecha_compra.isoformat(), comercio.strip(),
                                       moneda, db.as_text(cuotas), valor_total, fila_val,
                                       saldo_pendiente or 0, categoria,
                                       "Sí" if reembolsable else "No", notas.strip()]])
                db.clear_read_cache()
                st.success(f"Compra de {fmt_moneda(valor_total)} agregada.")
                st.rerun()
            except db.SinEspacioError as e:
                st.error(str(e))
            except Exception as e:
                st.error(f"No pude guardar: {e}")


def _form_eliminar_extracto(tarjeta_label, key_sufijo):
    blocks = db.TARJETA_BLOCKS[tarjeta_label]
    df_resumen = db.read_tarjeta_resumen(blocks["resumen"])
    if df_resumen.empty:
        st.info("Todavía no hay extractos cargados.")
        return
    periodos = sorted(df_resumen["Periodo Extracto"].tolist(), reverse=True)
    periodo_sel = st.selectbox("Período a eliminar", periodos, key=f"periodo_eliminar_extracto_{key_sufijo}")
    fila_sel = df_resumen[df_resumen["Periodo Extracto"] == periodo_sel].iloc[0]
    df_detalle = db.read_egreso_detalle(blocks["detalle"])
    n_detalle = (df_detalle["Periodo Extracto"] == periodo_sel).sum() if not df_detalle.empty else 0
    st.caption(f"Pago total: {fmt_moneda(fila_sel['Pago Total'])} — {n_detalle} compra(s) en el detalle de este período.")
    if "detalle_usd" in blocks:
        df_usd = db.read_egreso_detalle(blocks["detalle_usd"])
        n_usd = (df_usd["Periodo Extracto"] == periodo_sel).sum() if not df_usd.empty else 0
        if n_usd:
            st.caption(f"También tiene {n_usd} compra(s) en USD para este período — se borran igual.")
    confirmar = st.checkbox(f"Confirmo que quiero borrar el extracto '{periodo_sel}' — no se puede deshacer",
                             key=f"confirmar_eliminar_extracto_{key_sufijo}")
    if st.button("🗑️ Eliminar extracto", type="primary", disabled=not confirmar,
                 key=f"btn_eliminar_extracto_{key_sufijo}"):
        try:
            resultado = db.eliminar_extracto(tarjeta_label, periodo_sel)
            db.clear_read_cache()
            detalle_txt = f"{resultado['detalle']} compra(s)"
            if "detalle_usd" in resultado:
                detalle_txt += f" + {resultado['detalle_usd']} en USD"
            st.success(f"Extracto '{periodo_sel}' eliminado — {resultado['resumen']} fila(s) de resumen, "
                       f"{detalle_txt}.")
            st.rerun()
        except Exception as e:
            st.error(f"No pude borrar: {e}")


def _ver_extracto_puntual(tarjeta_label, key_sufijo):
    blocks = db.TARJETA_BLOCKS[tarjeta_label]
    tiene_usd = "detalle_usd" in blocks
    st.markdown("#### 🔍 Ver un extracto puntual")
    st.caption("Elegí un período de corte para ver su resumen (cupo, saldo, pago) junto con el detalle de "
               "compras de ese extracto — igual que elegir un mes en Colillas de Pago.")
    df_resumen = db.read_tarjeta_resumen(blocks["resumen"])
    if df_resumen.empty:
        st.info("Todavía no hay extractos cargados.")
        return
    periodos = sorted(df_resumen["Periodo Extracto"].tolist(), reverse=True)
    periodo_sel = st.selectbox("Período", periodos, key=f"periodo_ver_extracto_{key_sufijo}")
    fila = df_resumen[df_resumen["Periodo Extracto"] == periodo_sel].iloc[0]

    c1, c2, c3 = st.columns(3)
    c1.metric("Cupo Total", f"{fmt_moneda(fila['Cupo Total'])}")
    c2.metric("Cupo Disponible", f"{fmt_moneda(fila['Cupo Disponible'])}")
    c3.metric("% Cupo Utilizado", f"{fila['% Cupo Utilizado'] * 100:,.1f}%")
    c4, c5, c6 = st.columns(3)
    c4.metric("Saldo Anterior", f"{fmt_moneda(fila['Saldo Anterior'])}")
    c5.metric("Pago Mínimo", f"{fmt_moneda(fila['Pago Mínimo'])}")
    c6.metric("Pago Total", f"{fmt_moneda(fila['Pago Total'])}")
    st.caption(f"Fecha de Corte: {fila['Fecha de Corte']} — Fecha Límite de Pago: {fila['Fecha Límite de Pago']}")
    if tiene_usd and fila["Saldo a pagar USD"]:
        st.caption(f"Saldo a pagar en USD (aparte, no se suma con pesos): US$ {fila['Saldo a pagar USD']:,.2f}")

    df_detalle = db.read_egreso_detalle(blocks["detalle"])
    df_periodo = df_detalle[df_detalle["Periodo Extracto"] == periodo_sel] if not df_detalle.empty else df_detalle
    st.markdown(f"**Compras de este período** ({len(df_periodo)})")
    if not df_periodo.empty:
        st.dataframe(df_periodo, hide_index=True, use_container_width=True, height=350)
    else:
        st.info("Sin compras en COP para este período.")

    if tiene_usd:
        df_usd = db.read_egreso_detalle(blocks["detalle_usd"])
        df_usd_periodo = df_usd[df_usd["Periodo Extracto"] == periodo_sel] if not df_usd.empty else df_usd
        if not df_usd_periodo.empty:
            st.markdown(f"**Compras en USD de este período** ({len(df_usd_periodo)})")
            st.dataframe(df_usd_periodo, hide_index=True, use_container_width=True, height=250)


def _gestionar_extractos_tarjeta(tarjeta_label, block_key_cop, block_key_usd, key_sufijo):
    st.markdown("#### 🛠️ Gestionar extractos y compras")
    tab_agregar_ext, tab_agregar_compra, tab_eliminar = st.tabs(
        ["➕ Agregar extracto", "➕ Agregar compra", "🗑️ Eliminar extracto"])
    with tab_agregar_ext:
        _form_agregar_extracto_resumen(tarjeta_label, key_sufijo)
    with tab_agregar_compra:
        _form_agregar_compra_tarjeta(block_key_cop, block_key_usd, key_sufijo)
    with tab_eliminar:
        _form_eliminar_extracto(tarjeta_label, key_sufijo)


def render_egresos_visa():
    _gestionar_extractos_tarjeta("Visa ****7497", "visa_detalle", None, "visa")
    _ver_extracto_puntual("Visa ****7497", "visa")
    st.divider()
    _render_egreso_tab("visa_detalle", "Egresos - Tarjeta Visa 7497")


def render_egresos_mc():
    _gestionar_extractos_tarjeta("Mastercard ****5922", "mc_detalle", "mc_detalle_usd", "mc")
    _ver_extracto_puntual("Mastercard ****5922", "mc")
    st.divider()
    _render_egreso_tab("mc_detalle", "Egresos - Mastercard 5922")
    st.divider()
    st.markdown("#### Compras en USD (separado de los pesos de arriba — no se suman entre sí)")
    df_mc_usd = db.read_egreso_detalle("mc_detalle_usd")
    if not df_mc_usd.empty:
        st.caption(f"{len(df_mc_usd):,} movimientos — suma: US$ {df_mc_usd['Valor Cargado Este Periodo'].sum():,.2f}")
        st.dataframe(df_mc_usd, hide_index=True, use_container_width=True, height=300)
    else:
        st.info("Todavía no hay compras en USD cargadas.")


def render_egresos_tendencia():
    st.markdown("#### Tendencia de Tarjetas de Crédito")
    st.caption("Visa y Mastercard en pesos, comparadas período a período — no incluye Efectivo (no es tarjeta "
               "de crédito) ni las compras en USD de Mastercard (moneda distinta, no se puede sumar con pesos).")
    partes = []
    for block_key, nombre in [("visa_detalle", "Visa 7497"), ("mc_detalle", "Mastercard 5922")]:
        df = db.read_egreso_detalle(block_key)
        if not df.empty:
            df = df[(df["Moneda"] == "COP") & (df["Presupuestar"] == "Sí")]
        if not df.empty:
            t = df.groupby("Periodo Extracto")["Valor Cargado Este Periodo"].sum().reset_index()
            t["Tarjeta"] = nombre
            partes.append(t)

    if partes:
        df_tend = pd.concat(partes, ignore_index=True).sort_values("Periodo Extracto")
        fig = px.bar(df_tend, x="Periodo Extracto", y="Valor Cargado Este Periodo", color="Tarjeta",
                     barmode="group")
        fig.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=380, xaxis_title="Período")
        _fijar_eje_periodos(fig, df_tend["Periodo Extracto"])
        st.plotly_chart(fig, use_container_width=True)

        st.markdown("**Total por tarjeta (histórico, según lo cargado)**")
        totales = df_tend.groupby("Tarjeta")["Valor Cargado Este Periodo"].sum().reset_index()
        cols = st.columns(len(totales))
        for col, (_, fila) in zip(cols, totales.iterrows()):
            col.metric(fila["Tarjeta"], f"{fmt_moneda(fila['Valor Cargado Este Periodo'])}")
    else:
        st.info("Todavía no hay compras de tarjeta cargadas.")


def render_cargar_colillas():
    st.title("📄 Cargar Colillas de Pago")
    st.write("Sube una o varias colillas del Hospital Pablo Tobón Uribe, en PDF.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Leé cada colilla en PDF (devengos y descuentos, categorizados solos) y te muestra una vista "
            "previa antes de aplicar nada. Las que ya estaban cargadas, o repetidas entre los archivos de esta "
            "subida, se detectan y se omiten. Al aplicar, "
            "escribe directo en la hoja 'Colillas de Pago' del Sheet.")
    archivos = st.file_uploader("Colillas PDF", type=["pdf"], accept_multiple_files=True, key="colillas")

    if archivos:
        try:
            ya_cargadas = db.column_values("colillas_resumen", col=2)
        except Exception as e:
            st.error(f"No pude leer la hoja 'Colillas de Pago' del Sheet: {e}")
            st.stop()

        parsed = []
        vistos_en_subida = set()
        for f in archivos:
            try:
                r = parse_colilla(f, f.name)
            except Exception as e:
                st.error(f"No pude leer {f.name}: {e}")
                continue
            r["_archivo"] = f.name
            r["_duplicado_subida"] = r["periodo"] in vistos_en_subida
            r["_nueva"] = r["periodo"] not in ya_cargadas and not r["_duplicado_subida"]
            vistos_en_subida.add(r["periodo"])
            parsed.append(r)

        st.markdown("#### Vista previa")
        for r in parsed:
            if r["_nueva"]:
                estado = '<span class="pill-ok">nueva</span>'
            elif r["_duplicado_subida"]:
                estado = '<span class="pill-skip">repetida en esta subida — se omitirá</span>'
            else:
                estado = '<span class="pill-skip">ya cargada — se omitirá</span>'
            with st.container(border=True):
                st.markdown(f"**{r['periodo']}** ({r['_archivo']}) &nbsp; {estado}", unsafe_allow_html=True)
                c1, c2, c3, c4 = st.columns(4)
                c1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(r["total_devengos"])}</div>'
                            f'<div class="stat-label">Devengos</div></div>', unsafe_allow_html=True)
                c2.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(r["total_descuentos"])}</div>'
                            f'<div class="stat-label">Descuentos</div></div>', unsafe_allow_html=True)
                c3.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(r["neto"])}</div>'
                            f'<div class="stat-label">Neto</div></div>', unsafe_allow_html=True)
                c4.markdown(f'<div class="stat-card"><div class="stat-num">{r["fecha"]}</div>'
                            f'<div class="stat-label">Fecha de pago</div></div>', unsafe_allow_html=True)
                with st.expander("Ver devengos y descuentos detallados"):
                    st.write("**Devengos**")
                    st.dataframe([{"Concepto": i["concepto"], "Categoría": i["categoria"], "Valor": i["valor"]}
                                  for i in r["devengos"]], hide_index=True, use_container_width=True)
                    st.write("**Descuentos**")
                    st.dataframe([{"Concepto": i["concepto"], "Categoría": i["categoria"], "Valor": i["valor"]}
                                  for i in r["descuentos"]], hide_index=True, use_container_width=True)
                if r["revisar"]:
                    st.warning("Conceptos no reconocidos (quedarán sin categoría, para que la pongas tú): " +
                               ", ".join(f'"{c[1]}"' for c in r["revisar"]))

        nuevas = [r for r in parsed if r["_nueva"]]
        if nuevas:
            if st.button(f"✅ Aplicar {len(nuevas)} colilla(s) nueva(s)", type="primary", key="btn_colillas"):
                try:
                    aplicar_colillas(nuevas)
                    st.success(f"Listo — {len(nuevas)} colilla(s) agregada(s) directamente al Google Sheet. "
                               "Mirá el Dashboard para ver los totales actualizados.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"Algo falló al aplicar los cambios: {e}")
        elif parsed:
            st.info("Todas las colillas subidas ya estaban cargadas — no hay nada que aplicar.")


def render_cargar_extractos():
    st.title("💳 Cargar Extractos de Tarjeta")
    st.write("Sube uno o varios extractos detallados de la Visa o la Mastercard (.xlsx, tal como los descargas de Bancolombia).")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Leé cada extracto de tarjeta (compras categorizadas solas, resumen del corte) y te muestra una "
            "vista previa antes de aplicar nada. Los que ya estaban cargados, o repetidos entre los archivos "
            "de esta subida, se detectan y se omiten. Al aplicar, escribe directo en la hoja de Egresos de "
            "esa tarjeta en el Sheet.")
    archivos_ext = st.file_uploader("Extractos .xlsx", type=["xlsx"], accept_multiple_files=True, key="extractos")

    if archivos_ext:
        # Se lee UNA sola vez qué periodos ya están cargados por tarjeta (no por
        # archivo/hoja) — si no, subir muchos meses de una vuelta agota la cuota
        # de lectura de la API de Sheets (error 429) repitiendo la misma consulta.
        ya_por_tarjeta = {}

        def _ya_cargados(tarjeta_label):
            if tarjeta_label not in ya_por_tarjeta:
                blocks = db.TARJETA_BLOCKS[tarjeta_label]
                ya_por_tarjeta[tarjeta_label] = db.column_values(blocks["resumen"], col=1)
            return ya_por_tarjeta[tarjeta_label]

        parsed_ext = []
        vistos_en_subida = set()
        for f in archivos_ext:
            try:
                wbx = load_workbook(f, data_only=True)
            except Exception as e:
                st.error(f"No pude abrir {f.name}: {e}")
                continue
            for sn in wbx.sheetnames:
                try:
                    r = parse_extracto_sheet(wbx[sn], nombre_archivo=f.name)
                except Exception as e:
                    st.error(f"No pude leer la hoja '{sn}' de {f.name}: {e}")
                    continue
                r["_archivo"] = f.name
                r["_sheet"] = sn
                try:
                    ya = _ya_cargados(r["tarjeta"]["label"])
                    clave_subida = (r["tarjeta"]["label"], r["statement"]["periodo"], r["moneda"])
                    r["_duplicado_subida"] = clave_subida in vistos_en_subida
                    r["_nueva"] = (r["statement"]["periodo"] not in ya
                                   and not r["_duplicado_subida"])
                    vistos_en_subida.add(clave_subida)
                except Exception as e:
                    st.error(f"No pude revisar la hoja '{r['tarjeta']['sheet']}' del Sheet: {e}")
                    r["_nueva"] = False
                parsed_ext.append(r)

        st.markdown("#### Vista previa")
        for r in parsed_ext:
            if r["_nueva"]:
                estado = '<span class="pill-ok">nuevo</span>'
            elif r.get("_duplicado_subida"):
                estado = '<span class="pill-skip">repetido en esta subida — se omitirá</span>'
            else:
                estado = '<span class="pill-skip">ya cargado — se omitirá</span>'
            with st.container(border=True):
                st.markdown(f"**{r['tarjeta']['label']}** — {r['statement']['periodo']} ({r['moneda']}) "
                            f"&nbsp; <span style='color:#667085'>{r['_archivo']} [{r['_sheet']}]</span> &nbsp; {estado}",
                            unsafe_allow_html=True)
                s = r["statement"]
                c1, c2, c3, c4 = st.columns(4)
                c1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(s["pago_total"])}</div>'
                            f'<div class="stat-label">Pago total</div></div>', unsafe_allow_html=True)
                c2.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(s["cupo_disponible"])}</div>'
                            f'<div class="stat-label">Cupo disponible</div></div>', unsafe_allow_html=True)
                c3.markdown(f'<div class="stat-card"><div class="stat-num">{len(r["txns"])}</div>'
                            f'<div class="stat-label">Movimientos</div></div>', unsafe_allow_html=True)
                c4.markdown(f'<div class="stat-card"><div class="stat-num">{s["fecha_limite"]}</div>'
                            f'<div class="stat-label">Pagar antes de</div></div>', unsafe_allow_html=True)
                with st.expander("Ver movimientos"):
                    st.dataframe([{"Fecha": t["fecha_compra"], "Comercio": t["comercio"], "Cuotas": t["cuotas"],
                                   "Valor este período": t["valor_periodo"], "Categoría": t["categoria"]}
                                  for t in r["txns"]], hide_index=True, use_container_width=True)
                nuevos_comercios = sorted({t["comercio"] for t in r["txns"]
                                            if t["categoria"] == "Otros" and t["nota"] == "Comercio no reconocido, clasificar manualmente"})
                if nuevos_comercios:
                    st.warning("Comercios no reconocidos (quedarán en 'Otros', clasifícalos tú): " + ", ".join(nuevos_comercios))

        nuevos_ext = [r for r in parsed_ext if r["_nueva"]]
        if nuevos_ext:
            if st.button(f"✅ Aplicar {len(nuevos_ext)} extracto(s) nuevo(s)", type="primary", key="btn_extractos"):
                try:
                    aplicar_extractos(nuevos_ext)
                    st.success(f"Listo — {len(nuevos_ext)} extracto(s) agregado(s) directamente al Google Sheet. "
                               "Mirá el Dashboard para ver los totales actualizados.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"Algo falló al aplicar los cambios: {e}")
        elif parsed_ext:
            st.info("Todos los extractos subidos ya estaban cargados — no hay nada que aplicar.")


def render_cargar_cuenta():
    st.title("🏦 Cargar Cuenta de Ahorros")
    st.write("Sube uno o varios extractos o detalles de transacciones de tu cuenta de ahorros "
             "(.xlsx, tal como los descargas de Bancolombia). "
             "Los débitos de la cuenta terminada en 1031 se registran como Egresos - Efectivo. "
             "Los movimientos que ya se contabilizan por otro lado — pago automático de tarjetas, intereses "
             "diarios de ahorros, tu sueldo del hospital (ya está en Colillas de Pago) — se cargan igual, para "
             "que la conciliación de la cuenta cuadre, pero con una categoría '(no presupuestar)' para que no "
             "dupliquen tus totales de ingreso/gasto.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Leé el extracto de cuenta y categoriza cada movimiento solo (ingreso, gasto, aporte a "
            "inversión, o '(no presupuestar)' si ya se cuenta por otro lado). Te muestra un resumen antes de "
            "aplicar nada — nuevos, ya cargados, repetidos entre archivos. Al aplicar, escribe en Egresos - "
            "Efectivo, Otros Ingresos, y en Inversiones si detecta un aporte a una plataforma.")
    archivos_cta = st.file_uploader("Extracto de cuenta .xlsx", type=["xlsx"], accept_multiple_files=True, key="cuenta")

    if archivos_cta:
        try:
            ya_ingresos_raw = db.existing_keys("otros_ingresos", cols=[1, 2, 4], date_col=1)
            ya_egresos_raw = db.existing_keys("efectivo_detalle", cols=[2, 3, 7], date_col=2)
            ya_ingresos = [(fecha, concepto_cuenta(descripcion), valor)
                           for fecha, descripcion, valor in ya_ingresos_raw]
            ya_egresos = [(fecha, concepto_cuenta(descripcion), valor)
                          for fecha, descripcion, valor in ya_egresos_raw]
        except Exception as e:
            st.error(f"No pude revisar lo ya cargado en el Sheet: {e}")
            st.stop()

        todos = []
        for f in archivos_cta:
            try:
                wbx = load_workbook(f, data_only=True)
            except Exception as e:
                st.error(f"No pude abrir {f.name}: {e}")
                continue
            for sn in wbx.sheetnames:
                try:
                    txns = parse_extracto_cuenta_sheet(wbx[sn])
                except Exception as e:
                    st.error(f"No pude leer la hoja '{sn}' de {f.name}: {e}")
                    continue
                for txn in txns:
                    txn["_archivo"] = f.name
                    txn["_sheet"] = sn
                todos.extend(txns)

        marcar_novedad_movimientos(todos, ya_ingresos, ya_egresos)

        nuevos_ingresos = [t for t in todos if t["valor"] > 0 and t["_nuevo"]]
        nuevos_egresos = [t for t in todos if t["valor"] < 0 and t["_nuevo"]]
        duplicados_subida = [t for t in todos if t["_duplicado_subida"]]
        ya_cargados = [t for t in todos if not t["_nuevo"] and not t["_duplicado_subida"]]
        no_presupuestar_nuevos = [t for t in (nuevos_ingresos + nuevos_egresos) if t["no_presupuestar"]]

        st.markdown("#### Resumen")
        c1, c2, c3, c4, c5 = st.columns(5)
        c1.markdown(f'<div class="stat-card"><div class="stat-num">{len(nuevos_ingresos)}</div>'
                    f'<div class="stat-label">Ingresos nuevos</div></div>', unsafe_allow_html=True)
        c2.markdown(f'<div class="stat-card"><div class="stat-num">{len(nuevos_egresos)}</div>'
                    f'<div class="stat-label">Gastos nuevos</div></div>', unsafe_allow_html=True)
        c3.markdown(f'<div class="stat-card"><div class="stat-num">{len(ya_cargados)}</div>'
                    f'<div class="stat-label">Ya cargados</div></div>', unsafe_allow_html=True)
        c4.markdown(f'<div class="stat-card"><div class="stat-num">{len(duplicados_subida)}</div>'
                    f'<div class="stat-label">Repetidos en subida</div></div>', unsafe_allow_html=True)
        c5.markdown(f'<div class="stat-card"><div class="stat-num">{len(no_presupuestar_nuevos)}</div>'
                    f'<div class="stat-label">No presupuestar (se cargan igual)</div></div>', unsafe_allow_html=True)

        if nuevos_ingresos:
            st.markdown("#### Ingresos nuevos")
            st.dataframe([{"Fecha": t["fecha"], "Concepto": t["descripcion"], "Categoría": t["categoria"],
                           "Valor": t["valor"], "Notas": t["nota"]} for t in nuevos_ingresos],
                         hide_index=True, use_container_width=True)
            no_reconocidos_ing = sorted({t["descripcion"] for t in nuevos_ingresos if t["categoria"] == "Otro"})
            if no_reconocidos_ing:
                st.warning("Ingresos no reconocidos (quedarán en 'Otro', clasifícalos tú en el Sheet y avisame "
                           "para agregarlos a la lista): " + ", ".join(no_reconocidos_ing))

        if nuevos_egresos:
            st.markdown("#### Gastos nuevos")
            st.dataframe([{"Fecha": t["fecha"], "Concepto": t["descripcion"], "Categoría": t["categoria"],
                           "Valor": abs(t["valor"]), "Notas": t["nota"]} for t in nuevos_egresos],
                         hide_index=True, use_container_width=True)
            no_reconocidos = sorted({t["descripcion"] for t in nuevos_egresos if t["categoria"] == "Otros"})
            if no_reconocidos:
                st.warning("Movimientos no reconocidos (quedarán en 'Otros', clasifícalos tú en el Sheet y avisame "
                           "para agregarlos a la lista): " + ", ".join(no_reconocidos))

        if no_presupuestar_nuevos:
            with st.expander(f"Ver los {len(no_presupuestar_nuevos)} movimiento(s) '(no presupuestar)' — se cargan "
                              "igual, para conciliar la cuenta, pero no cuentan en tus totales de ingreso/gasto"):
                st.dataframe([{"Fecha": t["fecha"], "Concepto": t["descripcion"], "Categoría": t["categoria"],
                               "Valor": t["valor"], "Notas": t["nota"]} for t in no_presupuestar_nuevos],
                             hide_index=True, use_container_width=True)

        if ya_cargados:
            st.info(f"{len(ya_cargados)} movimiento(s) ya estaban cargados — se omiten.")

        if duplicados_subida:
            st.warning(f"{len(duplicados_subida)} movimiento(s) están repetidos entre los archivos de esta subida — "
                       "se aplicará únicamente la primera aparición.")
            with st.expander("Ver movimientos repetidos en esta subida"):
                st.dataframe([{"Fecha": t["fecha"], "Concepto": t["descripcion"],
                               "Valor": abs(t["valor"]), "Archivo": t["_archivo"]}
                              for t in duplicados_subida], hide_index=True, use_container_width=True)

        if nuevos_ingresos or nuevos_egresos:
            total_nuevos = len(nuevos_ingresos) + len(nuevos_egresos)
            if st.button(f"✅ Aplicar {total_nuevos} movimiento(s) nuevo(s)", type="primary", key="btn_cuenta"):
                try:
                    aplicar_movimientos_cuenta(nuevos_ingresos, nuevos_egresos)
                    st.success(f"Listo — {len(nuevos_ingresos)} ingreso(s) y {len(nuevos_egresos)} gasto(s) "
                               "agregados directamente al Google Sheet. Mirá el Dashboard para ver los totales actualizados.")
                    st.rerun()
                except db.SinEspacioError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"Algo falló al aplicar los cambios: {e}")
        elif todos:
            st.info("Todos los movimientos ya estaban cargados — no hay nada que aplicar.")


def _ingresos_gastos_periodo(coincide):
    """Ingresos, gasto real (con desglose por categoría) y descuentos de
    nómina para el período que define coincide(anio, mes) -> bool — vista
    **devengado** (cada colilla/compra/ingreso en su propia fecha). Reusado
    por Resumen y por Estado de Resultados/Flujo de Efectivo (🏢 Estados
    Financieros) para no repetir esta misma agregación tres veces."""
    df_cp = db.read_colillas_resumen()
    if not df_cp.empty:
        am = [_extraer_anio_mes(p) for p in df_cp["Periodo"]]
        mask = [coincide(a, m) for a, m in am]
        ingresos_colillas = df_cp.loc[mask, "Devengos Totales"].sum()
        # Descuentos Totales viene impreso en cada comprobante — es el total
        # confiable, a diferencia del detalle categorizado de abajo
        # (colillas_descuentos), que puede estar incompleto si alguna
        # quincena no se desglosó ítem por ítem. Usar el detalle acá
        # subestimaría los descuentos (e infla la Utilidad Neta y el Flujo
        # de Efectivo) cuando falte categorizar algo.
        descuentos_nomina = df_cp.loc[mask, "Descuentos Totales"].sum()
    else:
        ingresos_colillas = 0.0
        descuentos_nomina = 0.0

    df_oi = db.read_otros_ingresos_tabla()
    if not df_oi.empty:
        am = [_extraer_anio_mes(f) for f in df_oi["Fecha"]]
        mask = [coincide(a, m) and pres == "Sí" for (a, m), pres in zip(am, df_oi["Presupuestar"])]
        otros_ingresos = df_oi.loc[mask, "Valor"].sum()
    else:
        otros_ingresos = 0.0

    total_ingresos = ingresos_colillas + otros_ingresos

    # Ingresos por categoría — para discriminar el Estado de Resultados igual
    # que ya se discrimina el gasto (ver gasto_por_categoria más abajo).
    # Colillas: se usa el detalle de devengos (no el resumen por quincena),
    # excluyendo las categorías "(no presupuestar)" (cesantías) igual que se
    # excluyen del total. Otros ingresos: se reusa la misma máscara de arriba
    # (ya filtra Presupuestar=="Sí"), agrupando por su propia categoría.
    ingresos_por_categoria = {}
    df_dev = db.read_colillas_devengos()
    if not df_dev.empty:
        am_dev = [_extraer_anio_mes(q) for q in df_dev["Quincena"]]
        mask_dev = [coincide(a, m) and not es_no_presupuestar(cat)
                    for (a, m), cat in zip(am_dev, df_dev["Categoría"])]
        if any(mask_dev):
            sub = df_dev.loc[mask_dev].groupby("Categoría")["Valor"].sum()
            for cat, val in sub.items():
                ingresos_por_categoria[cat] = ingresos_por_categoria.get(cat, 0.0) + val
    if not df_oi.empty and any(mask):
        sub = df_oi.loc[mask].groupby("Categoría")["Valor"].sum()
        for cat, val in sub.items():
            ingresos_por_categoria[cat] = ingresos_por_categoria.get(cat, 0.0) + val

    # Gasto real: también devengado — efectivo + Visa + Mastercard en pesos,
    # cada compra en su propia fecha de compra (no en el mes del corte de
    # la tarjeta, que puede incluir hasta dos meses calendario distintos,
    # ni en el mes en que se paga la tarjeta).
    gasto_real = 0.0
    gasto_sin_categorizar = 0.0
    gasto_por_categoria = {}
    filas_inversiones = []
    for block_key in ("efectivo_detalle", "visa_detalle", "mc_detalle"):
        df_b = db.read_egreso_detalle(block_key)
        if df_b.empty:
            continue
        am = [_extraer_anio_mes(f) for f in df_b["Fecha Compra"]]
        mask = [coincide(a, m) and moneda == "COP" and pres == "Sí"
                for (a, m), moneda, pres in zip(am, df_b["Moneda"], df_b["Presupuestar"])]
        gasto_real += df_b.loc[mask, "Valor Cargado Este Periodo"].sum()
        mask_otros = [c and cat == "Otros" for c, cat in zip(mask, df_b["Categoría"])]
        gasto_sin_categorizar += df_b.loc[mask_otros, "Valor Cargado Este Periodo"].sum()
        if any(mask):
            sub = df_b.loc[mask].groupby("Categoría")["Valor Cargado Este Periodo"].sum()
            for cat, val in sub.items():
                gasto_por_categoria[cat] = gasto_por_categoria.get(cat, 0.0) + val
        # Detalle de cada movimiento categorizado "Inversiones" — el total
        # de gasto_inversiones (más abajo) es un solo bulto que puede
        # esconder que en realidad fue a varios destinos distintos (otro
        # broker, un fondo del banco, o plata que ni el importador supo
        # identificar). Se guarda el detalle acá para poder discriminarlo.
        mask_inv = [c and cat == "Inversiones" for c, cat in zip(mask, df_b["Categoría"])]
        if any(mask_inv):
            cols_detalle = ["Fecha Compra", "Comercio / Concepto", "Valor Cargado Este Periodo", "Notas"]
            sub_inv = df_b.loc[mask_inv, cols_detalle].copy()
            sub_inv["Fuente"] = block_key
            filas_inversiones.append(sub_inv)

    # El detalle de descuentos (categorizado ítem por ítem) solo se usa para
    # el desglose por categoría y para detectar si falta categorizar algo —
    # el total en sí ya se tomó arriba de Descuentos Totales (resumen).
    df_desc = db.read_colillas_descuentos()
    descuentos_detalle_sum = 0.0
    seguros_nomina = 0.0
    descuento_ahorro = 0.0
    descuento_fondo_empleados = 0.0
    descuento_deuda_nomina = 0.0
    if not df_desc.empty:
        am = [_extraer_anio_mes(q) for q in df_desc["Quincena"]]
        mask_desc = [coincide(a, m) for a, m in am]
        df_desc_periodo = df_desc.loc[mask_desc]
        descuentos_detalle_sum = df_desc_periodo["Valor"].sum()
        # Categorías reales de descuentos de nómina (ver 'Ahorro',
        # 'Aportes de Ley', 'Deuda (Leasing Habitacional)', 'Deuda
        # (Préstamo Fondo Empleados)', 'Fondo de Empleados', 'Impuestos',
        # 'Seguros', 'Transporte'). 'Fondo de Empleados' a secas es el
        # aporte/ahorro cooperativo (no la cuota del préstamo, que es la
        # categoría separada 'Deuda (Préstamo Fondo Empleados)').
        por_categoria_desc = df_desc_periodo.groupby("Categoría")["Valor"].sum()
        seguros_nomina = por_categoria_desc.get("Seguros", 0.0)
        descuento_ahorro = por_categoria_desc.get("Ahorro", 0.0)
        descuento_fondo_empleados = por_categoria_desc.get("Fondo de Empleados", 0.0)
        descuento_deuda_nomina = (por_categoria_desc.get("Deuda (Leasing Habitacional)", 0.0)
                                  + por_categoria_desc.get("Deuda (Préstamo Fondo Empleados)", 0.0))
    descuentos_sin_categorizar = max(0.0, descuentos_nomina - descuentos_detalle_sum)

    gasto_ahorro = gasto_por_categoria.get("Ahorro", 0.0)
    ingreso_cesantias = ingresos_por_categoria.get("Cesantías", 0.0)

    # El total de "Inversiones" es un solo bulto que puede esconder que en
    # realidad fue a varios destinos: un bróker con posiciones trackeadas
    # en 📈 Inversiones (Acciones y Valores), un fondo de inversión del
    # banco que la app no trackea como portafolio, o plata que el
    # importador nunca pudo identificar en el extracto y quedó en
    # Inversiones sin más rastro. Antes de rendirse a "sin identificar" se
    # cruza cada movimiento (misma fecha, mismo monto) contra los aportes
    # ya registrados en 📈 Inversiones (Pesos y Dólares) — el extracto de
    # cuenta suele nombrar la pasarela de pago (p. ej. "PSE Soluciones de
    # Pagos", "PSE Mono Colombia") en vez de la plataforma real, así que
    # ese cruce es la única forma confiable de saber a dónde fue.
    inversiones_por_destino = {}
    inversiones_por_destino_bruto = {}
    inversiones_sin_identificar = 0.0
    if filas_inversiones:
        df_inv_detalle = pd.concat(filas_inversiones, ignore_index=True)

        # aportes_index sirve para identificar la plataforma de un cargo
        # puntual (fecha+monto exactos); retiros_por_plataforma acumula por
        # separado los montos negativos (retiros) de ese mismo período — es
        # plata que salió de la plataforma y volvió a la cuenta corriente,
        # así que no debe seguir contando como "sigue invertida" en el
        # total de Inversiones. Un aporte o retiro nuevo cargado en 📈
        # Inversiones se refleja solo con esto, sin tocar código.
        aportes_index = {}
        retiros_por_plataforma = {}
        for moneda_ap in ("pesos", "dolares"):
            df_ap = db.read_aportes_inversion(moneda_ap)
            for _, fila_ap in df_ap.iterrows():
                fecha_ap = pd.to_datetime(fila_ap["Fecha"], dayfirst=True, errors="coerce")
                monto_ap = _num_o_cero(fila_ap["Monto Transferido (COP)"])
                if pd.isna(fecha_ap) or not monto_ap:
                    continue
                aportes_index[(fecha_ap.date(), round(monto_ap))] = fila_ap["Plataforma"]
                if monto_ap < 0 and coincide(fecha_ap.year, fecha_ap.month):
                    plataforma_ret = fila_ap["Plataforma"]
                    retiros_por_plataforma[plataforma_ret] = (
                        retiros_por_plataforma.get(plataforma_ret, 0.0) - monto_ap)

        def _destino_inversion(comercio, notas, fecha, valor):
            comercio = str(comercio).strip()
            if "acciones y val" in comercio.lower():
                return "Acciones y Valores"
            if "fondo de inversion" in comercio.lower():
                return "Fondo de Inversión (banco)"
            fecha_dt = pd.to_datetime(fecha, dayfirst=True, errors="coerce")
            if not pd.isna(fecha_dt):
                plataforma = aportes_index.get((fecha_dt.date(), round(valor)))
                if plataforma:
                    return plataforma
            if "no reconocido" in str(notas).lower():
                return f"Sin identificar ({comercio})" if comercio else "Sin identificar"
            return comercio or "Otro"

        df_inv_detalle["Destino"] = [
            _destino_inversion(c, n, f, v) for c, n, f, v in zip(
                df_inv_detalle["Comercio / Concepto"], df_inv_detalle["Notas"],
                df_inv_detalle["Fecha Compra"], df_inv_detalle["Valor Cargado Este Periodo"])
        ]
        sub = df_inv_detalle.groupby("Destino")["Valor Cargado Este Periodo"].sum().sort_values(ascending=False)
        inversiones_por_destino_bruto = sub.to_dict()
        inversiones_por_destino = dict(inversiones_por_destino_bruto)
        for plataforma_ret, retiro in retiros_por_plataforma.items():
            if plataforma_ret in inversiones_por_destino:
                inversiones_por_destino[plataforma_ret] = max(
                    0.0, inversiones_por_destino[plataforma_ret] - retiro)
        inversiones_sin_identificar = sum(v for k, v in inversiones_por_destino.items()
                                          if k.startswith("Sin identificar"))

    # "Inversiones" y "Ahorro" son categorías presupuestables normales
    # (Presupuestar=Sí, ya sumadas en gasto_real) pero no son gasto real —
    # es plata que sigue siendo tuya, solo que ahora en otra forma. Se
    # separan acá para que Estado de Resultados/Flujo de Efectivo/Resumen
    # puedan excluirlas del gasto operativo. gasto_inversiones ya descuenta
    # los retiros (dinero que volvió a la cuenta corriente y se gastó bajo
    # otra categoría) — sin esto, esa plata contaría dos veces: una como
    # inversión y otra como el gasto real en el que terminó.
    gasto_inversiones_bruto = gasto_por_categoria.get("Inversiones", 0.0)
    gasto_inversiones = (sum(inversiones_por_destino.values()) if inversiones_por_destino
                         else gasto_inversiones_bruto)
    gasto_operativo = gasto_real - gasto_inversiones - gasto_ahorro

    return dict(
        ingresos_colillas=ingresos_colillas, otros_ingresos=otros_ingresos, total_ingresos=total_ingresos,
        ingresos_por_categoria=ingresos_por_categoria, ingreso_cesantias=ingreso_cesantias,
        gasto_real=gasto_real, gasto_sin_categorizar=gasto_sin_categorizar,
        gasto_por_categoria=gasto_por_categoria, gasto_inversiones=gasto_inversiones,
        gasto_inversiones_bruto=gasto_inversiones_bruto,
        inversiones_por_destino=inversiones_por_destino,
        inversiones_por_destino_bruto=inversiones_por_destino_bruto,
        inversiones_sin_identificar=inversiones_sin_identificar,
        gasto_ahorro=gasto_ahorro, gasto_operativo=gasto_operativo,
        descuentos_nomina=descuentos_nomina, seguros_nomina=seguros_nomina,
        descuento_ahorro=descuento_ahorro, descuento_fondo_empleados=descuento_fondo_empleados,
        descuento_deuda_nomina=descuento_deuda_nomina,
        descuentos_sin_categorizar=descuentos_sin_categorizar,
    )


def _render_desglose_inversiones(datos):
    """Desglose de la categoría de gasto 'Inversiones' por destino real
    (bróker, fondo del banco, o sin identificar) — reusado por Resumen e
    Informe de Presupuesto para no repetir el mismo bloque dos veces."""
    por_destino = datos.get("inversiones_por_destino") or {}
    if not por_destino:
        return
    por_destino_bruto = datos.get("inversiones_por_destino_bruto") or {}
    retiros_total = datos.get("gasto_inversiones_bruto", 0.0) - datos.get("gasto_inversiones", 0.0)
    with st.expander(f"Ver los {len(por_destino)} destino(s) de 'Inversiones'"):
        filas_destino = []
        for destino, neto in sorted(por_destino.items(), key=lambda x: -x[1]):
            bruto = por_destino_bruto.get(destino, neto)
            filas_destino.append({"Destino": destino, "Aportado (bruto)": fmt_moneda(bruto),
                                   "Retirado": fmt_moneda(bruto - neto) if bruto > neto else "—",
                                   "Sigue invertido (neto)": fmt_moneda(neto)})
        st.dataframe(pd.DataFrame(filas_destino), hide_index=True, use_container_width=True)
        if retiros_total > 0:
            st.caption(f"'Sigue invertido (neto)' ya descuenta {fmt_moneda(retiros_total)} en retiros que "
                       "volvieron a la cuenta corriente y se gastaron bajo otra categoría — es la cifra que "
                       "usan Resumen, Estado de Resultados, Flujo de Efectivo e Informe de Presupuesto para "
                       "calcular cuánto se ahorró/invirtió de verdad. Cargá o corregí retiros en 📈 Inversiones.")
    sin_identificar = datos.get("inversiones_sin_identificar", 0.0)
    if sin_identificar > 0:
        pct = (sin_identificar / datos["gasto_inversiones"] * 100) if datos["gasto_inversiones"] else 0
        st.warning(f"⚠️ {fmt_moneda(sin_identificar)} ({pct:.0f}% de Inversiones) nunca se conectó con una "
                  "plataforma concreta — el importador no reconoció ese movimiento del extracto y quedó "
                  "categorizado como Inversiones sin más rastro de a dónde fue. Revisalo en 📊 Análisis → "
                  "Movimientos, o directo en la hoja 'Egresos - Efectivo'/'Tarjeta' del Sheet (columna Notas).")


def render_resumen():
    st.title("🏠 Resumen")
    st.caption("Vista general: cuánto entra, cuánto sale, qué balance queda, cuánta deuda tenés pendiente y "
               "cuánto llevás aportado a inversiones. Los números vienen directo de las fórmulas del Sheet — "
               "este es el punto de partida para entender tu mes; para el detalle día a día usá Análisis, "
               "Ingresos o Egresos.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "El punto de partida: tarjetas de 💰 Ingresos, 💳 Gastos, ⚖️ Balance, 🏦 Deudas y 📈 Inversiones, "
            "más un gráfico de tendencia de los últimos meses.\n\n"
            "El selector **'Ver resumen de:'** cambia Ingresos, Gastos y Balance entre Total histórico, un "
            "año puntual o un mes puntual (siempre vista **devengado** — cada compra o ingreso en su propia "
            "fecha). Deudas e Inversiones no cambian con el selector: son saldos de hoy, no algo que se "
            "acumule mes a mes.")

    alcance = st.radio("Ver resumen de:", ["Total histórico", "Un año", "Un mes"], horizontal=True,
                       key="alcance_resumen")
    anio_sel = mes_sel = None
    hoy = date.today()
    anios_disp = list(range(2023, 2033))
    if alcance == "Un año":
        anio_sel = st.selectbox("Año", anios_disp,
                                 index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0,
                                 key="anio_resumen")
    elif alcance == "Un mes":
        c_a, c_m = st.columns(2)
        anio_sel = c_a.selectbox("Año", anios_disp,
                                  index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0,
                                  key="anio_resumen_mes")
        mes_sel = c_m.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                 format_func=lambda m: MESES_NOMBRE[m], key="mes_resumen_mes")

    def coincide(anio, mes):
        if alcance == "Total histórico":
            return True
        if anio is None:
            return False
        if alcance == "Un año":
            return anio == anio_sel
        return anio == anio_sel and mes == mes_sel

    datos = _ingresos_gastos_periodo(coincide)
    ingresos_colillas = datos["ingresos_colillas"]
    otros_ingresos = datos["otros_ingresos"]
    total_ingresos = datos["total_ingresos"]
    gasto_real = datos["gasto_real"]
    gasto_sin_categorizar = datos["gasto_sin_categorizar"]
    descuentos_nomina = datos["descuentos_nomina"]
    seguros_nomina = datos["seguros_nomina"]

    st.markdown("#### 💰 Ingresos")
    c1, c2, c3 = st.columns(3)
    c1.metric("Colillas de pago", f"{fmt_moneda(ingresos_colillas)}")
    c2.metric("Otros ingresos", f"{fmt_moneda(otros_ingresos)}")
    c3.metric("Total ingresos brutos", f"{fmt_moneda(total_ingresos)}")

    st.markdown("#### 💳 Gastos")
    c4, c5, c6 = st.columns(3)
    c4.metric("Gasto real (efectivo+tarjetas)", f"{fmt_moneda(gasto_real)}")
    c5.metric("Descuentos de nómina", f"{fmt_moneda(descuentos_nomina)}")
    c6.metric("Seguros vía nómina", f"{fmt_moneda(seguros_nomina)}")
    if gasto_sin_categorizar > 0:
        pct_otros = (gasto_sin_categorizar / gasto_real * 100) if gasto_real else 0
        st.warning(f"⚠️ {fmt_moneda(gasto_sin_categorizar)} ({pct_otros:.0f}% del gasto real de este alcance) está en la "
                   "categoría genérica 'Otros' — revisalo en Análisis → Movimientos para categorizarlo mejor.")

    total_egresos = gasto_real + descuentos_nomina
    balance = total_ingresos - total_egresos
    # El "ahorro" no es un solo bulto — puede venir de la categoría de
    # gasto Inversiones, la categoría de gasto Ahorro, el descuento de
    # nómina "Ahorro" o el aporte al Fondo de Empleados (no la cuota del
    # préstamo del fondo, esa sí es deuda). Se discriminan por separado en
    # vez de sumarlos en una sola cifra ciega. Sumarlos de vuelta al
    # balance es lo mismo que ya hace Estado de Resultados con su
    # "Utilidad Neta"; sin esto, la tasa de ahorro subestima cuánto
    # ahorraste de verdad.
    ahorro_gasto = datos["gasto_ahorro"]
    ahorro_nomina = datos["descuento_ahorro"] + datos["descuento_fondo_empleados"]
    ahorro_total = datos["gasto_inversiones"] + ahorro_gasto + ahorro_nomina
    total_ahorrado = balance + ahorro_total
    tasa_ahorro = (total_ahorrado / total_ingresos * 100) if total_ingresos else 0
    st.markdown("#### ⚖️ Balance")
    st.caption("Ingresos brutos menos gasto real y descuentos de nómina, para el mismo alcance elegido arriba "
               "(vista devengado). No es lo mismo que 'Disponible del mes' de Balance Mensual/Resumen Mensual, "
               "que cuenta la tarjeta en el mes en que se paga el extracto, no en el mes de la compra.")
    c_bal1, c_bal2 = st.columns(2)
    c_bal1.metric("Total ingresos", f"{fmt_moneda(total_ingresos)}")
    c_bal2.metric("Total egresos", f"{fmt_moneda(total_egresos)}")
    c_bal3, c_bal4 = st.columns(2)
    c_bal3.metric("Balance (efectivo)", f"{fmt_moneda(balance)}")
    c_bal4.metric("Tasa de ahorro (efectivo + ahorro + inversión)", f"{tasa_ahorro:.0f}%")

    if ahorro_total > 0:
        st.markdown("###### No es gasto de consumo — discriminado por origen")
        ca1, ca2, ca3 = st.columns(3)
        ca1.metric("Inversiones", fmt_moneda(datos["gasto_inversiones"]))
        ca2.metric("Ahorro (efectivo/tarjeta)", fmt_moneda(ahorro_gasto))
        ca3.metric("Ahorro vía nómina", fmt_moneda(ahorro_nomina))
        st.caption("'Total egresos' y 'Balance' arriba sí cuentan estos montos como salida de efectivo, porque "
                  "literalmente dejaron la cuenta corriente — pero no dejaron de ser tuyos, así que se suman "
                  "de vuelta en la tasa de ahorro. 'Ahorro vía nómina' junta el descuento 'Ahorro' y el aporte "
                  "al Fondo de Empleados.")
        _render_desglose_inversiones(datos)

    if datos["descuento_deuda_nomina"] > 0:
        st.caption(f"Aparte, {fmt_moneda(datos['descuento_deuda_nomina'])} de este alcance se fue a cuotas de "
                  "deuda descontadas directo de la nómina (leasing habitacional / préstamo del Fondo de "
                  "Empleados) — tampoco es consumo, pero no se suma a la tasa de ahorro porque cada cuota "
                  "mezcla capital e interés y no hay forma de separarlos acá (el detalle de cada crédito está "
                  "en 🏦 Deudas).")

    if datos["ingreso_cesantias"] > 0:
        st.caption(f"De los ingresos, {fmt_moneda(datos['ingreso_cesantias'])} son de la categoría Cesantías "
                  "(intereses o retiro del fondo) — ya están incluidos en 'Total ingresos', discriminados acá "
                  "para que sepas cuánto de lo que entró no fue tu sueldo habitual.")

    if total_ahorrado >= 0:
        if tasa_ahorro >= 20:
            st.success(f"✅ Con este alcance, no gastaste {fmt_moneda(total_ahorrado)} (efectivo + ahorro + "
                       f"inversión) — una tasa de ahorro de {tasa_ahorro:.0f}% (saludable: 20% o más).")
        else:
            st.success(f"✅ Con este alcance, no gastaste {fmt_moneda(total_ahorrado)} (efectivo + ahorro + "
                       f"inversión) — una tasa de ahorro de {tasa_ahorro:.0f}% (referencia: 20% o más se "
                       "considera saludable).")
    else:
        st.warning(f"⚠️ Con este alcance, el gasto superó el ingreso por {fmt_moneda(abs(total_ahorrado))} "
                   f"incluso contando lo aportado a ahorro e inversión (tasa de ahorro de {tasa_ahorro:.0f}%).")

    st.markdown("#### 🏦 Deudas")
    st.caption("Estado actual — no cambia según el selector de arriba (es el saldo de hoy, no un monto que se "
               "acumule mes a mes).")
    df_deudas_r = db.read_deudas_resumen()
    c7, c8 = st.columns(2)
    c7.metric("Saldo total pendiente", f"{fmt_moneda(df_deudas_r['Saldo Actual'].sum())}")
    c8.metric("Cuota mensual total", f"{fmt_moneda(df_deudas_r['Cuota Mensual'].sum())}")

    st.markdown("#### 📈 Inversiones")
    st.caption("Estado actual — tampoco cambia según el selector de arriba.")
    kpis = dict(db.read_resumen_kpis())
    c9, c10 = st.columns(2)
    c9.metric("Portafolio en pesos", f"{fmt_moneda(kpis.get('Valor actual del portafolio de inversiones en pesos', 0))}")
    c10.metric("Portafolio en dólares",
               f"US$ {kpis.get('Valor actual del portafolio de inversiones en dólares', 0):,.2f}")

    st.divider()
    st.markdown("#### 📊 Tendencia de los últimos meses")
    st.caption("Vista **efectivo real** (igual que Balance Mensual/Evolución): ingresos, gastos y lo que "
               "quedó disponible cada mes, para los meses que ya tienen datos cargados.")
    df_mes_r = db.read_resumen_mensual()
    df_mes_r = df_mes_r[df_mes_r["Ingresos ganados"] != 0].tail(12)
    if not df_mes_r.empty:
        df_long_r = df_mes_r.melt(id_vars=["Mes"],
                                   value_vars=["Ingresos ganados", "Gastos personales", "Disponible del mes"],
                                   var_name="Concepto", value_name="Valor")
        fig_r = px.bar(df_long_r, x="Mes", y="Valor", color="Concepto", barmode="group")
        fig_r.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=360)
        _fijar_eje_periodos(fig_r, df_mes_r["Mes"])
        st.plotly_chart(fig_r, use_container_width=True)
    else:
        st.info("Todavía no hay suficientes meses cargados para ver la tendencia.")


def _render_estado_resultados():
    st.caption("Ingresos menos gastos operativos del período — la 'Utilidad Neta', igual que el estado de "
               "resultados de una empresa. No incluye aportes a Inversiones (no es un gasto, es una compra "
               "de activo — ver Balance General) ni pagos de deuda (eso va en Flujo de Efectivo).")

    alcance = st.radio("Alcance:", ["Total histórico", "Un año", "Un mes"], horizontal=True, key="alcance_er")
    anio_sel = mes_sel = None
    hoy = date.today()
    anios_disp = list(range(2023, 2033))
    if alcance == "Un año":
        anio_sel = st.selectbox("Año", anios_disp,
                                 index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0, key="anio_er")
    elif alcance == "Un mes":
        c_a, c_m = st.columns(2)
        anio_sel = c_a.selectbox("Año", anios_disp,
                                  index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0, key="anio_er_mes")
        mes_sel = c_m.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                 format_func=lambda m: MESES_NOMBRE[m], key="mes_er_mes")

    def coincide(anio, mes):
        if alcance == "Total histórico":
            return True
        if anio is None:
            return False
        if alcance == "Un año":
            return anio == anio_sel
        return anio == anio_sel and mes == mes_sel

    datos = _ingresos_gastos_periodo(coincide)

    st.markdown("##### Ingresos")
    c1, c2, c3 = st.columns(3)
    c1.metric("Colillas de pago", fmt_moneda(datos["ingresos_colillas"]))
    c2.metric("Otros ingresos", fmt_moneda(datos["otros_ingresos"]))
    c3.metric("Total Ingresos", fmt_moneda(datos["total_ingresos"]))

    if datos["ingresos_por_categoria"]:
        with st.expander("Ver ingresos por categoría"):
            df_ing_cat = pd.DataFrame(sorted(datos["ingresos_por_categoria"].items(), key=lambda x: -x[1]),
                                       columns=["Categoría", "Valor"])
            df_ing_cat["Valor"] = df_ing_cat["Valor"].map(fmt_moneda)
            st.dataframe(df_ing_cat, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_ing_cat)))

    st.markdown("##### Gastos operativos")
    total_gastos = datos["gasto_operativo"] + datos["descuentos_nomina"]
    c4, c5, c6 = st.columns(3)
    c4.metric("Gasto de consumo (sin inversiones ni ahorro)", fmt_moneda(datos["gasto_operativo"]))
    c5.metric("Descuentos de nómina", fmt_moneda(datos["descuentos_nomina"]))
    c6.metric("Total Gastos", fmt_moneda(total_gastos))

    cat_sin_inversiones = {k: v for k, v in datos["gasto_por_categoria"].items() if k != "Inversiones"}
    if cat_sin_inversiones:
        with st.expander("Ver gasto por categoría"):
            df_cat = pd.DataFrame(sorted(cat_sin_inversiones.items(), key=lambda x: -x[1]),
                                   columns=["Categoría", "Valor"])
            df_cat["Valor"] = df_cat["Valor"].map(fmt_moneda)
            st.dataframe(df_cat, hide_index=True, use_container_width=True,
                         height=min(400, 45 + 35 * len(df_cat)))

    if datos["descuentos_sin_categorizar"] > 100:
        st.warning(f"⚠️ {fmt_moneda(datos['descuentos_sin_categorizar'])} de los descuentos de nómina de este "
                   "período no tiene detalle categorizado (el total viene del resumen de cada comprobante, "
                   "pero falta desglosarlo ítem por ítem en 'Colillas de Pago') — completalo para poder ver "
                   "en qué categoría cayó.")

    utilidad_neta = datos["total_ingresos"] - total_gastos
    tasa_ahorro_er = (utilidad_neta / datos["total_ingresos"] * 100) if datos["total_ingresos"] else 0
    st.markdown("##### Utilidad Neta del Período")
    c_un1, c_un2 = st.columns(2)
    c_un1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(utilidad_neta)}</div>'
                    f'<div class="stat-label">Utilidad Neta</div></div>', unsafe_allow_html=True)
    c_un2.markdown(f'<div class="stat-card"><div class="stat-num">{tasa_ahorro_er:.0f}%</div>'
                    f'<div class="stat-label">Tasa de Ahorro</div></div>', unsafe_allow_html=True)
    if utilidad_neta >= 0:
        st.success(f"✅ Resultado positivo: sobraron {fmt_moneda(utilidad_neta)} en este período "
                   f"({tasa_ahorro_er:.0f}% de los ingresos).")
    else:
        st.warning(f"⚠️ Resultado negativo: el gasto superó el ingreso por {fmt_moneda(abs(utilidad_neta))}.")


def _render_balance_general():
    st.caption("Foto de hoy: qué tenés (Activos) menos qué debés (Pasivos) = Patrimonio Neto. No cambia con "
               "ningún selector de período — es un saldo, no algo que se acumule mes a mes.")

    df_conc = db.read_conciliacion_efectivo()
    if not df_conc.empty:
        fila_ef = df_conc.sort_values("Mes").iloc[-1]
        saldo_efectivo = float(fila_ef["Saldo Final"])
        mes_efectivo = fila_ef["Mes"]
    else:
        saldo_efectivo = 0.0
        mes_efectivo = None

    kpis = dict(db.read_resumen_kpis())
    valor_portafolio_pesos = kpis.get("Valor actual del portafolio de inversiones en pesos", 0) or 0
    valor_portafolio_dolares = kpis.get("Valor actual del portafolio de inversiones en dólares", 0) or 0
    total_activos = saldo_efectivo + valor_portafolio_pesos

    st.markdown("##### Activos")
    if mes_efectivo:
        st.caption(f"Efectivo: saldo final de {mes_efectivo} (cargalo/actualizalo en Flujo de Efectivo).")
    else:
        st.caption("Todavía no cargaste ningún saldo de cuenta en Flujo de Efectivo.")
    c1, c2, c3 = st.columns(3)
    c1.metric("Efectivo (cuenta de ahorros)", fmt_moneda(saldo_efectivo))
    c2.metric("Portafolio de inversiones (pesos)", fmt_moneda(valor_portafolio_pesos))
    c3.metric("Total Activos", fmt_moneda(total_activos))
    if valor_portafolio_dolares:
        st.caption(f"Portafolio de inversiones en dólares (aparte, no sumado): US$ {valor_portafolio_dolares:,.2f}")
    if valor_portafolio_pesos == 0 and valor_portafolio_dolares == 0:
        st.info("El portafolio de inversiones da $0 porque todavía no cargaste posiciones (Cantidad/Precio "
                "Actual) en 📈 Inversiones → Posiciones.")

    st.markdown("##### Pasivos")
    df_deudas = db.read_deudas_resumen()
    total_pasivos = df_deudas["Saldo Actual"].sum() if not df_deudas.empty else 0.0
    cuota_mensual_total = df_deudas["Cuota Mensual"].sum() if not df_deudas.empty else 0.0
    deuda_usd = db.read_deuda_tarjeta_usd()
    c4, c5 = st.columns(2)
    c4.metric("Saldo total de deudas (pesos)", fmt_moneda(total_pasivos))
    c5.metric("Cuota mensual total", fmt_moneda(cuota_mensual_total))
    if deuda_usd:
        entidad_usd, nota_usd, valor_usd = deuda_usd
        st.caption(f"{entidad_usd} — deuda en dólares (aparte, no sumada): US$ {valor_usd:,.2f}. {nota_usd}")

    patrimonio_neto = total_activos - total_pasivos
    st.markdown("##### Patrimonio Neto")
    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(patrimonio_neto)}</div>'
                f'<div class="stat-label">Activos − Pasivos (pesos)</div></div>', unsafe_allow_html=True)


def _flujo_efectivo_periodo(coincide):
    """Financiación (pago de deuda) y Conciliación (el resto de movimientos
    '(no presupuestar)' de la cuenta: pago automático de tarjeta, intereses/
    4x1000, transferencias entre cuentas propias) para el período que define
    coincide(anio, mes) -> bool. Se separa de _ingresos_gastos_periodo
    porque esas categorías quedan afuera del presupuesto (Presupuestar=No)
    pero sí mueven plata real de la cuenta — hace falta sumarlas para que el
    Saldo Final de Flujo de Efectivo/Auditoría Anual cuadre contra el
    extracto real (el pago automático de tarjeta suele ser el monto más
    grande de esta categoría). Reusado por ambas pantallas para no repetir
    esta agregación con un período distinto cada vez (mes vs. año)."""
    df_ef = db.read_egreso_detalle("efectivo_detalle")
    pago_deuda = 0.0
    otros_conciliacion_egreso = 0.0
    if not df_ef.empty:
        am = [_extraer_anio_mes(f) for f in df_ef["Fecha Compra"]]
        mask = [coincide(a, m) for a, m in am]
        mask_deuda = [m and cat == "Pago de deuda (no presupuestar)"
                      for m, cat in zip(mask, df_ef["Categoría"])]
        pago_deuda = df_ef.loc[mask_deuda, "Valor Cargado Este Periodo"].sum()
        # OJO: "PAGO AUTOM TC VISA/MASTERCARD" (categoría "Pago Tarjeta de
        # Crédito (no presupuestar)") es la transferencia real que sale de
        # la cuenta de ahorros para pagar el extracto — pero esa plata YA
        # está en gasto_operativo vía el detalle devengado de cada compra
        # de la tarjeta (Egresos - Tarjeta Visa/Mastercard, en su propia
        # fecha de compra). Sumar también esta transferencia acá la cuenta
        # dos veces y descuadra el saldo calculado por el valor íntegro del
        # pago de tarjeta cada mes. Se detecta por la nota ("ya
        # contabilizado en el detalle de la tarjeta"), mismo criterio que
        # ya se usa abajo para no duplicar la nómina.
        notas_ef_lower = df_ef["Notas"].fillna("").str.lower()
        es_dup_ef = notas_ef_lower.str.contains("ya contabilizad")
        mask_otros_np = [m and cat != "Pago de deuda (no presupuestar)" and es_no_presupuestar(cat) and not dup
                         for m, cat, dup in zip(mask, df_ef["Categoría"], es_dup_ef)]
        otros_conciliacion_egreso = df_ef.loc[mask_otros_np, "Valor Cargado Este Periodo"].sum()

    # OJO: "PAGO DE NOMI HOSPITAL HPTU" (nómina del hospital vista desde el
    # extracto de cuenta) es "(no presupuestar)" por una razón distinta a
    # las demás — no es "no es ingreso nuevo", es "ya contabilizada en
    # Colillas de Pago, no duplicar" (ver su nota y cuenta_formatos.py). Ese
    # dinero YA está en flujo_operacion vía el neto de colillas
    # (Devengos Totales - descuentos_nomina); sumarlo también acá lo cuenta
    # dos veces y descuadra el saldo calculado por el valor íntegro de la
    # nómina. Se detecta por la nota (no por texto de categoría/concepto a
    # mano) para que cualquier "ya contabilizada.../no duplicar" futuro
    # quede afuera igual, sin tener que acordarse de actualizar esta lista.
    df_oi = db.read_otros_ingresos_tabla()
    otros_conciliacion_ingreso = 0.0
    if not df_oi.empty:
        am = [_extraer_anio_mes(f) for f in df_oi["Fecha"]]
        notas_lower = df_oi["Notas"].fillna("").str.lower()
        es_duplicado = notas_lower.str.contains("no duplicar") | notas_lower.str.contains("ya contabilizad")
        mask_np = [coincide(a, m) and es_no_presupuestar(cat) and not dup
                   for (a, m), cat, dup in zip(am, df_oi["Categoría"], es_duplicado)]
        otros_conciliacion_ingreso = df_oi.loc[mask_np, "Valor"].sum()

    return dict(
        pago_deuda=pago_deuda, otros_conciliacion_ingreso=otros_conciliacion_ingreso,
        otros_conciliacion_egreso=otros_conciliacion_egreso,
        flujo_conciliacion=otros_conciliacion_ingreso - otros_conciliacion_egreso,
    )


def _saldo_inicial_encadenado(anio_obj, mes_obj):
    """Saldo inicial de (anio_obj, mes_obj) partiendo del saldo real guardado
    más reciente ANTES de ese mes (no necesariamente el mes inmediatamente
    anterior — puede ser dic-2024 si no se guardó nada más cerca) y sumando
    el flujo calculado (Operación+Inversión+Financiación+Conciliación) de
    cada mes intermedio. Antes, si el mes anterior no tenía un saldo
    guardado a mano, se asumía $0 — lo cual era casi siempre falso, porque
    el saldo real de partida (31/12/2024) sí se conoce y de ahí en adelante
    todo lo demás ya está cargado (colillas, extractos, movimientos de
    cuenta). Devuelve (saldo, mes_ancla) o (None, None) si no hay ningún
    saldo guardado antes de ese mes."""
    mes_objetivo = f"{anio_obj}-{mes_obj:02d}"
    df_conc = db.read_conciliacion_efectivo()
    if df_conc.empty:
        return None, None
    anteriores = df_conc[df_conc["Mes"] < mes_objetivo]
    if anteriores.empty:
        return None, None
    fila_ancla = anteriores.sort_values("Mes").iloc[-1]
    mes_ancla = fila_ancla["Mes"]
    anio_a, mes_a = (int(x) for x in mes_ancla.split("-"))
    idx_ancla = anio_a * 12 + (mes_a - 1)
    idx_objetivo = anio_obj * 12 + (mes_obj - 1)
    saldo = float(fila_ancla["Saldo Final"])
    for idx in range(idx_ancla + 1, idx_objetivo):
        a, m = idx // 12, idx % 12 + 1

        def coincide(anio, mes, _a=a, _m=m):
            return anio == _a and mes == _m

        d = _ingresos_gastos_periodo(coincide)
        f = _flujo_efectivo_periodo(coincide)
        fo = d["total_ingresos"] - d["gasto_operativo"] - d["descuentos_nomina"]
        fi = -d["gasto_inversiones"]
        ff = -f["pago_deuda"]
        fc = f["flujo_conciliacion"]
        saldo += fo + fi + ff + fc
    return saldo, mes_ancla


def _render_flujo_efectivo():
    st.caption("Cuánta plata entra y sale realmente de tu cuenta de ahorros en un mes — empezando por cuánto "
               "tenías al arrancar. Se separa en Operación (tu día a día), Inversión (aportes a plataformas "
               "de inversión), Financiación (cuotas de deuda pagadas automáticamente desde la cuenta) y "
               "Conciliación (pago automático de tarjeta, intereses/4x1000, transferencias entre tus propias "
               "cuentas — no son gasto ni ingreso real, pero sí mueven la plata de la cuenta).")

    hoy = date.today()
    anios_fe = list(range(2023, 2033))
    c1, c2 = st.columns(2)
    anio_fe = c1.selectbox("Año", anios_fe, index=anios_fe.index(hoy.year) if hoy.year in anios_fe else 0,
                            key="anio_fe")
    mes_fe_num = c2.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                               format_func=lambda m: MESES_NOMBRE[m], key="mes_fe")
    mes_fe = f"{anio_fe}-{mes_fe_num:02d}"

    df_conc = db.read_conciliacion_efectivo()
    existente = df_conc[df_conc["Mes"] == mes_fe] if not df_conc.empty else df_conc
    if not existente.empty:
        saldo_inicial_prev = float(existente.iloc[0]["Saldo Inicial"])
        saldo_final_prev = float(existente.iloc[0]["Saldo Final"])
    else:
        # Sin saldo guardado para este mes puntual: en vez de asumir $0,
        # encadenamos desde el saldo real más reciente guardado antes de
        # este mes (aunque sea de varios meses atrás, p. ej. dic-2024) más
        # el flujo calculado de cada mes intermedio — todo eso ya está
        # cargado (colillas, extractos, movimientos de cuenta).
        saldo_inicial_prev, mes_ancla = _saldo_inicial_encadenado(anio_fe, mes_fe_num)
        if saldo_inicial_prev is None:
            st.warning(f"⚠️ No hay ningún saldo real guardado antes de {mes_fe} — el $0 de abajo es solo un "
                       "valor de partida, no un cálculo. Ingresá el saldo real de tu cuenta al cierre de "
                       "diciembre de 2024 (el primer mes con datos) en Verificar Datos para que la cadena de "
                       "saldos arranque bien y esto deje de mostrar $0.")
            saldo_inicial_prev = 0.0
        else:
            st.caption(f"Saldo inicial calculado encadenando desde el último saldo real guardado "
                       f"({mes_ancla}) más el flujo de los meses intermedios.")
        saldo_final_prev = 0.0

    c3, c4 = st.columns(2)
    saldo_inicial = c3.number_input("Saldo inicial del mes", value=saldo_inicial_prev, step=1000.0,
                                     format="%.0f", key=f"fe_saldo_ini_{mes_fe}")
    saldo_final_manual = c4.number_input("Saldo final del mes (según tu extracto)", value=saldo_final_prev,
                                          step=1000.0, format="%.0f", key=f"fe_saldo_fin_{mes_fe}")
    if st.button("💾 Guardar saldos de este mes", key="btn_guardar_fe"):
        db.guardar_conciliacion_efectivo(mes_fe, saldo_inicial, saldo_final_manual)
        db.clear_read_cache()
        st.success(f"Saldos de {mes_fe} guardados.")
        st.rerun()

    def coincide(anio, mes):
        return anio == anio_fe and mes == mes_fe_num

    datos = _ingresos_gastos_periodo(coincide)
    fe = _flujo_efectivo_periodo(coincide)

    flujo_operacion = datos["total_ingresos"] - datos["gasto_operativo"] - datos["descuentos_nomina"]
    flujo_inversion = -datos["gasto_inversiones"]
    flujo_financiacion = -fe["pago_deuda"]
    flujo_conciliacion = fe["flujo_conciliacion"]
    saldo_final_calculado = (saldo_inicial + flujo_operacion + flujo_inversion + flujo_financiacion
                              + flujo_conciliacion)

    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(saldo_inicial)}</div>'
                f'<div class="stat-label">Saldo Inicial del Mes</div></div>', unsafe_allow_html=True)

    st.markdown("##### Flujo del mes")
    c5, c6, c7, c8 = st.columns(4)
    c5.metric("Operación", fmt_moneda(flujo_operacion))
    c6.metric("Inversión", fmt_moneda(flujo_inversion))
    c7.metric("Financiación", fmt_moneda(flujo_financiacion))
    c8.metric("Conciliación", fmt_moneda(flujo_conciliacion))

    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(saldo_final_calculado)}</div>'
                f'<div class="stat-label">Saldo Final Calculado</div></div>', unsafe_allow_html=True)

    if not (saldo_inicial == 0 and saldo_final_manual == 0):
        diferencia = saldo_final_calculado - saldo_final_manual
        if abs(diferencia) > 100:
            st.warning(f"El saldo calculado no cuadra con el saldo final que ingresaste — diferencia de "
                       f"{fmt_moneda(diferencia)}. Revisá si falta cargar algún movimiento de este mes.")
        else:
            st.success(f"✅ Cuadra con el saldo final ingresado ({fmt_moneda(saldo_final_manual)}).")

    _render_desglose_flujo_efectivo(coincide, datos)


def _tabla_categorias(diccionario, nombre_valor="Valor"):
    if not diccionario:
        st.caption("Nada este mes.")
        return
    df = pd.DataFrame(sorted(diccionario.items(), key=lambda x: -abs(x[1])), columns=["Categoría", nombre_valor])
    df[nombre_valor] = df[nombre_valor].map(fmt_moneda)
    st.dataframe(df, hide_index=True, use_container_width=True, height=min(320, 45 + 35 * len(df)))


def _tabla_movimientos(df, columnas=("Fecha Compra", "Comercio / Concepto", "Categoría", "Valor Cargado Este Periodo")):
    if df.empty:
        st.caption("Nada este mes.")
        return
    vista = df[list(columnas)].rename(columns={"Valor Cargado Este Periodo": "Valor", "Fecha Compra": "Fecha"})
    vista["Valor"] = vista["Valor"].map(fmt_moneda)
    st.dataframe(vista.sort_values("Fecha"), hide_index=True, use_container_width=True,
                 height=min(360, 45 + 35 * len(vista)))


def _render_desglose_flujo_efectivo(coincide, datos):
    """Detalle línea por línea de las 4 categorías del flujo del mes — para
    entender qué mueve cada número, en especial Conciliación (donde suelen
    aparecer los traslados grandes de inversión que descuadran el saldo si
    no se revisan)."""
    with st.expander("🔎 Ver desglose del mes"):
        st.markdown("**Operación**")
        oc1, oc2, oc3 = st.columns(3)
        with oc1:
            st.caption("Ingresos por categoría")
            _tabla_categorias(datos["ingresos_por_categoria"])
        with oc2:
            st.caption("Gasto operativo por categoría")
            _tabla_categorias({k: v for k, v in datos["gasto_por_categoria"].items() if k != "Inversiones"})
        with oc3:
            st.caption("Descuentos de nómina por categoría")
            df_desc = db.read_colillas_descuentos()
            if not df_desc.empty:
                am = [_extraer_anio_mes(q) for q in df_desc["Quincena"]]
                mask = [coincide(a, m) for a, m in am]
                sub = df_desc.loc[mask].groupby("Categoría")["Valor"].sum().to_dict()
            else:
                sub = {}
            _tabla_categorias(sub)

        st.markdown("**Inversión** (aportes a plataformas de inversión)")
        lineas_inv = []
        for block_key in ("efectivo_detalle", "visa_detalle", "mc_detalle"):
            df_b = db.read_egreso_detalle(block_key)
            if df_b.empty:
                continue
            am = [_extraer_anio_mes(f) for f in df_b["Fecha Compra"]]
            mask = [coincide(a, m) and cat == "Inversiones" for (a, m), cat in zip(am, df_b["Categoría"])]
            if any(mask):
                lineas_inv.append(df_b.loc[mask])
        _tabla_movimientos(pd.concat(lineas_inv) if lineas_inv else pd.DataFrame(
            columns=["Fecha Compra", "Comercio / Concepto", "Categoría", "Valor Cargado Este Periodo"]))

        st.markdown("**Financiación** (cuotas de deuda pagadas automáticamente)")
        df_ef = db.read_egreso_detalle("efectivo_detalle")
        if not df_ef.empty:
            am = [_extraer_anio_mes(f) for f in df_ef["Fecha Compra"]]
            mask = [coincide(a, m) and cat == "Pago de deuda (no presupuestar)"
                    for (a, m), cat in zip(am, df_ef["Categoría"])]
            _tabla_movimientos(df_ef.loc[mask])
        else:
            st.caption("Nada este mes.")

        st.markdown("**Conciliación** (pago automático de tarjeta ya contado en Operación, intereses/4x1000, "
                     "transferencias entre tus propias cuentas, traslados de/hacia fondos de inversión)")
        cc1, cc2 = st.columns(2)
        with cc1:
            st.caption("Entradas de conciliación (Otros Ingresos)")
            df_oi = db.read_otros_ingresos_tabla()
            if not df_oi.empty:
                am_oi = [_extraer_anio_mes(f) for f in df_oi["Fecha"]]
                notas_lower = df_oi["Notas"].fillna("").str.lower()
                dup = notas_lower.str.contains("no duplicar") | notas_lower.str.contains("ya contabilizad")
                mask = [coincide(a, m) and es_no_presupuestar(cat) and not d
                        for (a, m), cat, d in zip(am_oi, df_oi["Categoría"], dup)]
                vista = df_oi.loc[mask, ["Fecha", "Concepto", "Categoría", "Valor"]]
                if not vista.empty:
                    vista = vista.copy()
                    vista["Valor"] = vista["Valor"].map(fmt_moneda)
                    st.dataframe(vista, hide_index=True, use_container_width=True,
                                 height=min(360, 45 + 35 * len(vista)))
                else:
                    st.caption("Nada este mes.")
            else:
                st.caption("Nada este mes.")
        with cc2:
            st.caption("Salidas de conciliación (Egresos - Efectivo)")
            if not df_ef.empty:
                notas_ef_lower = df_ef["Notas"].fillna("").str.lower()
                dup_ef = notas_ef_lower.str.contains("ya contabilizad")
                mask = [coincide(a, m) and cat != "Pago de deuda (no presupuestar)" and es_no_presupuestar(cat)
                        and not d
                        for (a, m), cat, d in zip(am, df_ef["Categoría"], dup_ef)]
                _tabla_movimientos(df_ef.loc[mask], columnas=("Fecha Compra", "Comercio / Concepto", "Categoría",
                                                               "Valor Cargado Este Periodo"))
            else:
                st.caption("Nada este mes.")


def _render_auditoria_anual():
    st.caption("Para un año completo: ingresos y gastos discriminados por categoría, y si el saldo calculado "
               "a fin de año cuadra contra el saldo real que tenías el 31 de diciembre — para auditar un año "
               "(ej. 2024 o 2025) contra tus extractos antes de confiar en el presupuesto hacia adelante.")

    hoy = date.today()
    anios_aud = list(range(2023, 2033))
    anio_aud = st.selectbox("Año a auditar", anios_aud,
                             index=anios_aud.index(hoy.year) if hoy.year in anios_aud else 0, key="anio_aud")

    def coincide_anio(a, m):
        return a == anio_aud

    df_conc = db.read_conciliacion_efectivo()
    mes_dic_anterior = f"{anio_aud - 1}-12"
    mes_dic = f"{anio_aud}-12"
    fila_dic_ant = df_conc[df_conc["Mes"] == mes_dic_anterior] if not df_conc.empty else df_conc
    fila_dic = df_conc[df_conc["Mes"] == mes_dic] if not df_conc.empty else df_conc
    if not fila_dic_ant.empty:
        saldo_inicial_real = float(fila_dic_ant.iloc[0]["Saldo Final"])
    else:
        # Sin saldo guardado justo para diciembre anterior: encadenamos
        # desde el último saldo real guardado antes de ese punto (mismo
        # criterio que Flujo de Efectivo) en vez de asumir $0. Se pide el
        # saldo inicial de enero de anio_aud (no de diciembre) porque
        # queremos el saldo de CIERRE de diciembre, es decir el de
        # arranque del mes siguiente.
        saldo_inicial_real, mes_ancla_aud = _saldo_inicial_encadenado(anio_aud, 1)
    saldo_final_real = float(fila_dic.iloc[0]["Saldo Final"]) if not fila_dic.empty else None

    if saldo_inicial_real is None:
        st.warning(f"⚠️ No hay ningún saldo real guardado antes de {mes_dic_anterior} — cargalo en Flujo de "
                   "Efectivo o Verificar Datos (el saldo real de cierre de esa cuenta al 31 de diciembre de "
                   f"2024, el primer mes con datos) para poder auditar {anio_aud} contra tu extracto real. "
                   "Mientras tanto se asume $0 como arranque.")
    elif fila_dic_ant.empty:
        st.caption(f"Saldo de arranque de {mes_dic_anterior} calculado encadenando desde el último saldo real "
                   f"guardado ({mes_ancla_aud}) más el flujo de los meses intermedios.")
    if saldo_final_real is None:
        st.warning(f"⚠️ Tampoco hay un saldo guardado para {mes_dic} — cargalo también para poder comparar el "
                   f"cierre de {anio_aud}.")

    datos = _ingresos_gastos_periodo(coincide_anio)
    fe = _flujo_efectivo_periodo(coincide_anio)

    st.markdown("##### Ingresos del año, por categoría")
    if datos["ingresos_por_categoria"]:
        df_ing = pd.DataFrame(sorted(datos["ingresos_por_categoria"].items(), key=lambda x: -x[1]),
                               columns=["Categoría", "Valor"])
        df_ing["Valor"] = df_ing["Valor"].map(fmt_moneda)
        st.dataframe(df_ing, hide_index=True, use_container_width=True, height=min(400, 45 + 35 * len(df_ing)))
    else:
        st.info(f"No hay ingresos cargados para {anio_aud}.")
    st.metric(f"Total Ingresos {anio_aud}", fmt_moneda(datos["total_ingresos"]))

    st.markdown("##### Gastos del año, por categoría")
    cat_sin_inversiones = {k: v for k, v in datos["gasto_por_categoria"].items() if k != "Inversiones"}
    if cat_sin_inversiones:
        df_gas = pd.DataFrame(sorted(cat_sin_inversiones.items(), key=lambda x: -x[1]),
                               columns=["Categoría", "Valor"])
        df_gas["Valor"] = df_gas["Valor"].map(fmt_moneda)
        st.dataframe(df_gas, hide_index=True, use_container_width=True, height=min(400, 45 + 35 * len(df_gas)))
    else:
        st.info(f"No hay gastos cargados para {anio_aud}.")
    total_gastos = datos["gasto_operativo"] + datos["descuentos_nomina"]
    st.metric(f"Total Gastos {anio_aud}", fmt_moneda(total_gastos))
    if datos["gasto_sin_categorizar"] > 0:
        st.warning(f"⚠️ {fmt_moneda(datos['gasto_sin_categorizar'])} de {anio_aud} sigue en la categoría "
                   "genérica 'Otros' — revisalo en 🔍 Salud de los Datos para que esta auditoría sea confiable.")
    if datos["descuentos_sin_categorizar"] > 100:
        st.warning(f"⚠️ {fmt_moneda(datos['descuentos_sin_categorizar'])} de los descuentos de nómina de "
                   f"{anio_aud} no tiene detalle categorizado en 'Colillas de Pago' — el total sí está bien "
                   "(viene del resumen de cada comprobante), pero falta desglosarlo ítem por ítem.")

    utilidad_neta = datos["total_ingresos"] - total_gastos
    st.markdown("##### Utilidad Neta del año")
    st.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(utilidad_neta)}</div>'
                f'<div class="stat-label">Utilidad Neta {anio_aud}</div></div>', unsafe_allow_html=True)

    st.divider()
    st.markdown("##### Reconciliación de caja del año (cuenta de ahorros)")
    st.caption("Mismas cuatro categorías que Flujo de Efectivo (Operación/Inversión/Financiación/"
               "Conciliación), acumuladas para el año completo — arrancando del saldo real de diciembre "
               "del año anterior.")
    flujo_operacion = datos["total_ingresos"] - datos["gasto_operativo"] - datos["descuentos_nomina"]
    flujo_inversion = -datos["gasto_inversiones"]
    flujo_financiacion = -fe["pago_deuda"]
    flujo_conciliacion = fe["flujo_conciliacion"]
    saldo_inicial_calc = saldo_inicial_real if saldo_inicial_real is not None else 0.0
    saldo_final_calculado = (saldo_inicial_calc + flujo_operacion + flujo_inversion + flujo_financiacion
                              + flujo_conciliacion)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Operación", fmt_moneda(flujo_operacion))
    c2.metric("Inversión", fmt_moneda(flujo_inversion))
    c3.metric("Financiación", fmt_moneda(flujo_financiacion))
    c4.metric("Conciliación", fmt_moneda(flujo_conciliacion))

    cc1, cc2 = st.columns(2)
    cc1.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(saldo_inicial_calc)}</div>'
                 f'<div class="stat-label">Saldo Inicial ({anio_aud - 1}-12-31)</div></div>',
                 unsafe_allow_html=True)
    cc2.markdown(f'<div class="stat-card"><div class="stat-num">{fmt_moneda(saldo_final_calculado)}</div>'
                 f'<div class="stat-label">Saldo Final Calculado ({anio_aud}-12-31)</div></div>',
                 unsafe_allow_html=True)

    if saldo_final_real is not None:
        diferencia = saldo_final_calculado - saldo_final_real
        if abs(diferencia) > 100:
            st.warning(f"⚠️ El saldo calculado no cuadra contra el saldo real de {mes_dic} — diferencia de "
                       f"{fmt_moneda(diferencia)}. Mirá el detalle mes a mes de abajo para ubicar en qué mes "
                       "se rompe la cadena.")
        else:
            st.success(f"✅ Cuadra contra el saldo real de {mes_dic} ({fmt_moneda(saldo_final_real)}).")

    with st.expander("Ver detalle mes a mes"):
        filas = []
        saldo_corrida = saldo_inicial_calc
        for m in range(1, 13):
            def coincide_mes(a, mm, _m=m):
                return a == anio_aud and mm == _m
            d = _ingresos_gastos_periodo(coincide_mes)
            f = _flujo_efectivo_periodo(coincide_mes)
            fo = d["total_ingresos"] - d["gasto_operativo"] - d["descuentos_nomina"]
            fi = -d["gasto_inversiones"]
            ff = -f["pago_deuda"]
            fc = f["flujo_conciliacion"]
            saldo_ini_mes = saldo_corrida
            saldo_fin_mes = saldo_ini_mes + fo + fi + ff + fc
            saldo_corrida = saldo_fin_mes
            mes_str = f"{anio_aud}-{m:02d}"
            fila_real = df_conc[df_conc["Mes"] == mes_str] if not df_conc.empty else df_conc
            saldo_real_mes = float(fila_real.iloc[0]["Saldo Final"]) if not fila_real.empty else None
            diff_mes = (saldo_fin_mes - saldo_real_mes) if saldo_real_mes is not None else None
            filas.append({
                "Mes": mes_str,
                "Ingresos": fmt_moneda(d["total_ingresos"]),
                "Gastos": fmt_moneda(d["gasto_operativo"] + d["descuentos_nomina"]),
                "Saldo Final Calculado": fmt_moneda(saldo_fin_mes),
                "Saldo Final Real": fmt_moneda(saldo_real_mes) if saldo_real_mes is not None else "-",
                "Diferencia": fmt_moneda(diff_mes) if diff_mes is not None else "-",
            })
        st.dataframe(pd.DataFrame(filas), hide_index=True, use_container_width=True, height=440)
        st.caption("El 'Saldo Final Calculado' es la cadena acumulada mes a mes desde el saldo real de "
                   "diciembre anterior — no depende de si guardaste un saldo manual ese mes en particular. "
                   "'Saldo Final Real' solo aparece si guardaste una conciliación para ese mes específico en "
                   "Flujo de Efectivo o Verificar Datos.")


def render_estados_financieros():
    st.title("🏢 Estados Financieros")
    st.caption("Tus finanzas vistas como las de una empresa: Estado de Resultados (ingresos vs. gastos de un "
               "período), Balance General (activos vs. pasivos de hoy), Flujo de Efectivo (cuánto entra y "
               "sale de tu cuenta cada mes) y Auditoría Anual (todo un año contra el saldo real de diciembre).")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "- **Estado de Resultados** — Ingresos menos gastos operativos del período que elijas → "
            "Utilidad Neta, con desglose por categoría de ambos lados. No incluye aportes a Inversiones (es "
            "una compra de activo, no un gasto) ni pagos de deuda (van en Flujo de Efectivo).\n"
            "- **Balance General** — Activos (efectivo + portafolio de inversiones) menos Pasivos (deudas y "
            "tarjetas) = Patrimonio Neto. Es una foto de hoy, no cambia con el mes.\n"
            "- **Flujo de Efectivo** — Saldo Inicial del mes (con cuánto arrancaste) más lo que entró y "
            "salió, separado en Operación/Inversión/Financiación/Conciliación, hasta el Saldo Final.\n"
            "- **Auditoría Anual** — el mismo Estado de Resultados y Flujo de Efectivo, pero acumulados para "
            "un año completo (ej. 2024 o 2025) y comparados contra el saldo real que tenías el 31 de "
            "diciembre — para verificar que todo lo cargado de ese año cuadra contra tus extractos, con un "
            "detalle mes a mes para ubicar dónde se rompe si no cuadra.")
    sub = st.radio("Sub-sección", ["Estado de Resultados", "Balance General", "Flujo de Efectivo",
                                    "Auditoría Anual"], horizontal=True, label_visibility="collapsed")
    st.divider()
    if sub == "Estado de Resultados":
        _render_estado_resultados()
    elif sub == "Balance General":
        _render_balance_general()
    elif sub == "Flujo de Efectivo":
        _render_flujo_efectivo()
    else:
        _render_auditoria_anual()


def render_analisis():
    st.title("📊 Análisis")
    st.caption("Cómo va tu mes visto desde distintos ángulos: por categoría, separando gasto esencial del que "
               "no lo es, la evolución mes a mes, el detalle de cada movimiento, o el balance de un mes puntual.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "6 vistas del mismo gasto:\n\n"
            "- **Categorías** — cuánto gastaste históricamente en cada categoría (devengado: cada compra en "
            "su propia fecha, no en el mes en que pagaste la tarjeta).\n"
            "- **Esenciales / No Esenciales** — separa el gasto de consumo entre esencial y no esencial, mes "
            "a mes. La clasificación de cada categoría se edita en la hoja 'Categorías Esenciales' del Sheet.\n"
            "- **Evolución** — la tendencia mes a mes de ingresos, gastos, deudas y disponible (vista "
            "**efectivo real**: la tarjeta cuenta en el mes en que se paga).\n"
            "- **Año vs. Año** — el mismo mes (enero, febrero...) comparado entre distintos años, para ver "
            "si vas mejor o peor que el año pasado en esa misma época.\n"
            "- **Movimientos** — el detalle de todos los movimientos con filtros, más un gráfico de top "
            "categorías.\n"
            "- **Balance Mensual** — el balance completo de un mes puntual, con un selector para ver el "
            "gasto de tarjeta como Consumo (mes de compra) o Efectivo real (mes de pago).")
    sub = st.radio("Sub-sección", ["Categorías", "Esenciales / No Esenciales", "Evolución", "Año vs. Año",
                                    "Movimientos", "Balance Mensual"], horizontal=True, label_visibility="collapsed")
    st.divider()
    if sub == "Categorías":
        render_categorias()
    elif sub == "Esenciales / No Esenciales":
        render_esenciales()
    elif sub == "Evolución":
        render_evolucion()
    elif sub == "Año vs. Año":
        render_anio_vs_anio()
    elif sub == "Movimientos":
        render_movimientos()
    else:
        render_balance_mensual()


def render_ingresos():
    st.title("💰 Ingresos")
    st.caption("Todo lo que entró: tus colillas de pago del hospital y cualquier otro ingreso ocasional "
               "registrado desde un extracto de cuenta o agregado a mano.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "- **Colillas de Pago** — el detalle de cada quincena (devengos y descuentos), con filtro por "
            "mes y un gráfico de tendencia. El expander '➕ Agregar una quincena manualmente' es la "
            "alternativa a subir el PDF: escribís la fecha, la quincena, y cada devengo/descuento con su "
            "categoría en una tabla editable. El expander '🗑️ Eliminar una quincena' borra una quincena "
            "puntual (resumen + devengos + descuentos) — pide confirmación antes de borrar.\n"
            "- **Otros Ingresos** — cualquier otro ingreso detectado en tus extractos de cuenta, con "
            "filtros de categoría/año/mes, un agrupamiento opcional de los rendimientos financieros por mes "
            "(son muchos abonos chicos) y un gráfico de ingresos por categoría. El expander '➕ Agregar un "
            "ingreso manualmente' de arriba lo escribe directo en la hoja, sin necesidad de un extracto.")
    sub = st.radio("Sub-sección", ["Colillas de Pago", "Otros Ingresos"], horizontal=True,
                   label_visibility="collapsed")
    st.divider()
    if sub == "Colillas de Pago":
        render_ingresos_colillas()
    else:
        render_ingresos_otros()


def render_egresos():
    st.title("💳 Egresos")
    st.caption("El detalle de cada hoja de egresos tal como está en el Sheet — Efectivo, Visa y Mastercard — "
               "con búsqueda y filtros de categoría/año/mes.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "- **Efectivo / Visa / Mastercard** — el detalle de cada uno, con búsqueda y filtros. Las "
            "tarjetas tienen además un selector Consumo (mes en que se compra) / Efectivo real (mes en que "
            "se paga), que también cambia los gráficos de tendencia y de categoría por mes. Todas tienen "
            "expanders '➕ Agregar...' para escribir a mano sin necesidad de subir un archivo: en Efectivo, "
            "un gasto; en Visa/Mastercard, tanto una compra suelta como el resumen del extracto del mes "
            "(cupo, saldo anterior, pago). En Visa/Mastercard también hay un '🔍 Ver un extracto puntual' que "
            "elegís de una lista y te muestra su resumen junto con sus compras — igual que elegir un mes en "
            "Colillas de Pago —, y un '🗑️ Eliminar un extracto' que borra el resumen de un período puntual "
            "junto con todas sus compras (pide confirmación antes).\n"
            "- **📊 Tendencia por Tarjeta** — compara Visa vs. Mastercard período a período, para ver cuál "
            "estás usando más.")
    sub = st.radio("Sub-sección", ["Efectivo", "Visa", "Mastercard", "📊 Tendencia por Tarjeta"],
                   horizontal=True, label_visibility="collapsed")
    st.divider()
    if sub == "Efectivo":
        render_egresos_efectivo()
    elif sub == "Visa":
        render_egresos_visa()
    elif sub == "Mastercard":
        render_egresos_mc()
    else:
        render_egresos_tendencia()


def _form_agregar_factura_electronica():
    with st.expander("➕ Agregar una factura electrónica manualmente"):
        with st.form("form_agregar_factura_electronica"):
            c1, c2 = st.columns(2)
            fecha_fe = c1.date_input("Fecha", value=date.today(), key="fecha_nueva_factura")
            valor_fe = c2.number_input("Valor (dejalo en 0 si no lo sabés)", min_value=0.0, step=1000.0,
                                        format="%.0f", key="valor_nueva_factura")
            emisor_fe = st.text_input("Emisor (razón social)", key="emisor_nueva_factura")
            c3, c4 = st.columns(2)
            nit_fe = c3.text_input("NIT Emisor", key="nit_nueva_factura")
            numero_fe = c4.text_input("Número de Documento", key="numero_nueva_factura")
            tipo_fe = st.text_input("Tipo de Documento", value="Factura Electrónica de Venta",
                                     key="tipo_nueva_factura")
            notas_fe = st.text_input("Notas (opcional)", key="notas_nueva_factura")
            guardar_fe = st.form_submit_button("💾 Guardar factura", type="primary")
        if guardar_fe:
            if not emisor_fe.strip() or not numero_fe.strip():
                st.error("Emisor y Número de Documento son obligatorios — son la clave que evita duplicados.")
            else:
                agregadas = db.agregar_facturas_electronicas([[
                    fecha_fe.isoformat(), nit_fe.strip(), emisor_fe.strip(), numero_fe.strip(),
                    tipo_fe.strip(), valor_fe if valor_fe > 0 else "", "", notas_fe.strip(),
                ]])
                if agregadas:
                    st.success("Factura agregada.")
                    st.rerun()
                else:
                    st.warning("Ya existía una factura con ese Emisor + Número de Documento — no se agregó de nuevo.")


def render_facturacion_electronica():
    st.title("🧾 Facturación Electrónica")
    st.caption("Registro de facturas y notas electrónicas (DIAN) recibidas por correo, año a año — no se "
               "concilia contra Egresos, es solo el archivo de los documentos.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Cada fila es un documento electrónico recibido por correo — es un registro documental, no un "
            "movimiento de plata. Valor y Fecha se leen del XML (UBL DIAN) adjunto al correo, no del texto "
            "del mensaje — es el dato exacto de la factura, no una aproximación.\n\n"
            "**La app no guarda el PDF/XML adjunto**, solo los datos que extrae de él — la columna 'Correo' "
            "de la tabla abre el mail original en Gmail para las facturas de la carga histórica, por si "
            "querés ver el archivo en sí. Un puñado de facturas no traía un adjunto legible y quedaron sin "
            "Valor — esas hay que revisarlas a mano abriendo el correo.")

    df = db.read_facturacion_electronica()
    if df.empty:
        st.info("Todavía no hay facturas electrónicas cargadas. Agregá una a mano abajo.")
    else:
        am = [_extraer_anio_mes(f) for f in df["Fecha"]]
        df = df.assign(_anio=[a for a, _ in am], _mes=[m for _, m in am])
        # La app no guarda el PDF/XML adjunto (solo los datos del correo) —
        # el ID del hilo de Gmail de la carga histórica quedó embebido en
        # Notas ("... (thread <id>) ..."), así que se puede armar un link
        # directo al correo original (ahí sí está el archivo adjunto) sin
        # tener que guardar el archivo en ningún lado.
        df = df.assign(Correo=df["Notas"].str.extract(r"thread (\w+)\)", expand=False)
                        .map(lambda tid: f"https://mail.google.com/mail/u/0/#all/{tid}" if pd.notna(tid) else None))

        con_valor = df[df["Valor"].apply(lambda v: isinstance(v, (int, float)) and v > 0)]

        c1, c2 = st.columns(2)
        c1.metric("Facturas cargadas", f"{len(df):,}")
        c2.metric("Con valor identificado", f"{len(con_valor):,}")

        resumen_anio = (con_valor.groupby("_anio").agg(Facturas=("Valor", "size"), Total=("Valor", "sum"))
                         .rename_axis("Año").sort_index(ascending=False).reset_index())
        resumen_anio["Total"] = resumen_anio["Total"].map(lambda v: f"${v:,.0f}")
        st.caption("Resumen por año (solo facturas con valor identificado):")
        st.dataframe(resumen_anio, hide_index=True, use_container_width=True)

        busqueda = st.text_input("🔍 Buscar por emisor", "", key="busqueda_facturacion")
        anios = ["(todos)"] + sorted({str(a) for a in df["_anio"] if a}, reverse=True)
        meses = ["(todos)"] + [f"{m:02d} - {MESES_NOMBRE[m]}" for m in range(1, 13)]
        c1, c2 = st.columns(2)
        anio_sel = c1.selectbox("Año", anios, key="anio_facturacion")
        mes_sel = c2.selectbox("Mes", meses, key="mes_facturacion")

        df_f = df
        if busqueda:
            df_f = df_f[df_f["Emisor"].str.contains(busqueda, case=False, na=False)]
        if anio_sel != "(todos)":
            df_f = df_f[df_f["_anio"] == int(anio_sel)]
        if mes_sel != "(todos)":
            df_f = df_f[df_f["_mes"] == int(mes_sel[:2])]

        cols_mostrar = ["Fecha", "Emisor", "NIT Emisor", "Número Documento", "Tipo Documento", "Valor", "Correo"]
        st.caption(f"{len(df_f):,} de {len(df):,} facturas — Valor y Fecha vienen del XML adjunto a cada "
                   "correo, no del texto del mensaje. La columna 'Correo' abre el mail original en Gmail "
                   "(ahí está el archivo; la app no lo guarda, solo los datos que extrajo de él).")
        orden = pd.to_datetime(df_f["Fecha"], format="%d/%m/%Y", errors="coerce").sort_values(ascending=False).index
        st.dataframe(df_f.loc[orden, cols_mostrar], hide_index=True, use_container_width=True, height=500,
                     column_config={"Correo": st.column_config.LinkColumn("Correo", display_text="Abrir ↗")})

    st.divider()
    _form_agregar_factura_electronica()


def render_declaraciones_renta():
    st.title("📑 Declaraciones de Renta")
    st.caption("Un registro por año — el desglose de la declaración y un link al PDF guardado en Google Drive.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Cada fila es una declaración de renta ya presentada, con los valores clave desglosados: "
            "patrimonio líquido, ingresos y renta líquida gravable, impuesto a cargo, retenciones y "
            "anticipos, y el saldo resultante.\n\n"
            "**La app no aloja el PDF** — subílo a mano a una carpeta de Google Drive y pegá acá el link "
            "para compartir (clic derecho sobre el archivo en Drive → Compartir → Copiar enlace); la tabla "
            "lo deja como un link que abre directo.\n\n"
            "Volver a guardar el mismo Año reemplaza esa fila, así que podés usarlo también para corregir "
            "una declaración ya cargada.")

    df = db.read_declaraciones_renta()
    if not df.empty:
        df_orden = df.sort_values("Año", ascending=False)
        cols_mostrar = ["Año", "Fecha de Presentación", "Patrimonio Líquido", "Ingresos Brutos",
                         "Renta Líquida Gravable", "Impuesto a Cargo", "Retenciones y Anticipos",
                         "Saldo (+ a pagar / - a favor)", "Link PDF", "Notas"]
        st.dataframe(df_orden[cols_mostrar], hide_index=True, use_container_width=True,
                     column_config={"Link PDF": st.column_config.LinkColumn("Link PDF", display_text="Abrir ↗")})

        if len(df) > 1:
            df_graf = df.sort_values("Año")
            fig = px.line(df_graf, x="Año", y=["Patrimonio Líquido", "Impuesto a Cargo"], markers=True)
            fig.update_layout(yaxis_title="COP", legend_title="")
            st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("Todavía no hay declaraciones cargadas. Agregá una abajo.")

    st.divider()
    with st.expander("➕ Agregar / editar una declaración", expanded=df.empty):
        with st.form("form_declaracion_renta"):
            c1, c2 = st.columns(2)
            anio_dr = c1.number_input("Año", min_value=2000, max_value=date.today().year,
                                       value=date.today().year - 1, step=1, key="anio_declaracion_renta")
            fecha_dr = c2.date_input("Fecha de Presentación", value=date.today(), key="fecha_declaracion_renta")
            c1, c2 = st.columns(2)
            patrimonio_dr = c1.number_input("Patrimonio Líquido", min_value=0.0, step=100000.0, format="%.0f",
                                             key="patrimonio_declaracion_renta")
            ingresos_dr = c2.number_input("Ingresos Brutos", min_value=0.0, step=100000.0, format="%.0f",
                                           key="ingresos_declaracion_renta")
            c1, c2 = st.columns(2)
            renta_dr = c1.number_input("Renta Líquida Gravable", min_value=0.0, step=100000.0, format="%.0f",
                                        key="renta_declaracion_renta")
            impuesto_dr = c2.number_input("Impuesto a Cargo", min_value=0.0, step=10000.0, format="%.0f",
                                           key="impuesto_declaracion_renta")
            c1, c2 = st.columns(2)
            retenciones_dr = c1.number_input("Retenciones y Anticipos", min_value=0.0, step=10000.0, format="%.0f",
                                              key="retenciones_declaracion_renta")
            saldo_dr = c2.number_input("Saldo (+ a pagar / - a favor)", step=10000.0, format="%.0f",
                                        key="saldo_declaracion_renta")
            link_dr = st.text_input("Link al PDF en Google Drive (opcional)", key="link_declaracion_renta")
            notas_dr = st.text_input("Notas (opcional)", key="notas_declaracion_renta")
            guardar_dr = st.form_submit_button("💾 Guardar declaración", type="primary")
        if guardar_dr:
            db.guardar_declaracion_renta([
                int(anio_dr), fecha_dr.isoformat(), patrimonio_dr, ingresos_dr, renta_dr, impuesto_dr,
                retenciones_dr, saldo_dr, link_dr.strip(), notas_dr.strip(),
            ])
            st.success(f"Declaración {int(anio_dr)} guardada.")
            st.rerun()


def _trm_en(serie_trm, fecha):
    """Último TRM (COP por USD) conocido en o antes de 'fecha' — si la
    serie no llega tan atrás, usa el primer dato disponible. 'serie_trm'
    es la salida de _descargar_historico_cierre('COP=X', ...)."""
    if serie_trm is None or serie_trm.empty:
        return None
    hasta = serie_trm[serie_trm.index <= pd.Timestamp(fecha)]
    return float(hasta.iloc[-1]) if not hasta.empty else float(serie_trm.iloc[0])


def _analisis_cambiario_dolares():
    """Para cada aporte/retiro a una plataforma en dólares (IBKR, Binance,
    Hapi), reconstruye cuántos dólares equivalió con la TRM del día de esa
    transferencia, y cuántos pesos valdrían esos mismos dólares hoy con la
    TRM de hoy — separa el efecto cambiario (cuánto se movió el dólar) del
    rendimiento de las posiciones en sí (que ya se mide en dólares con
    retorno bruto/capital propio/XIRR). None si no hay aportes o no se
    pudo descargar la TRM."""
    df_ap = db.read_aportes_inversion("dolares")
    if df_ap.empty:
        return None
    flujos = []
    for _, fila in df_ap.iterrows():
        fecha = pd.to_datetime(fila["Fecha"], dayfirst=True, errors="coerce")
        if pd.isna(fecha):
            continue
        monto = _num_o_cero(fila["Monto Transferido (COP)"])
        if monto:
            flujos.append((fecha.date(), monto, fila["Plataforma"]))
    if not flujos:
        return None

    fecha_inicio = min(f for f, _, _ in flujos)
    serie_trm = _descargar_historico_cierre("COP=X", fecha_inicio)
    trm_hoy = _descargar_trm()
    if serie_trm.empty or trm_hoy is None:
        return None

    filas = []
    for fecha, monto, plataforma in flujos:
        trm_fecha = _trm_en(serie_trm, fecha)
        if not trm_fecha:
            continue
        usd_equiv = monto / trm_fecha
        valor_hoy = usd_equiv * trm_hoy
        filas.append(dict(Fecha=fecha, Plataforma=plataforma, Monto_COP=monto, TRM_Fecha=trm_fecha,
                          USD_Equivalente=usd_equiv, TRM_Hoy=trm_hoy, Valor_Hoy_COP=valor_hoy,
                          Diferencia_Cambiaria=valor_hoy - monto))
    if not filas:
        return None
    df = pd.DataFrame(filas)
    total_cop = float(df["Monto_COP"].sum())
    total_usd = float(df["USD_Equivalente"].sum())
    total_hoy = float(df["Valor_Hoy_COP"].sum())
    return dict(df=df, total_cop=total_cop, total_usd=total_usd, total_hoy=total_hoy,
                diferencia=total_hoy - total_cop, trm_promedio=(total_cop / total_usd) if total_usd else None,
                trm_hoy=trm_hoy)


def _twr_moneda(moneda):
    """TWR (Time-Weighted Return) aproximado: encadena el rendimiento de
    cada sub-período entre fotos consecutivas de 'Historial de Valor de
    Cartera' (se guardan solas cada vez que actualizás precios o
    posiciones/aportes), neutralizando con Modified Dietz los aportes o
    retiros que hubo adentro de cada sub-período. A diferencia del XIRR
    (money-weighted, ya en 'Rentabilidad anualizada'), el TWR no le importa
    CUÁNDO metiste cada peso — mide solo qué tan bien rindieron las
    posiciones en cada tramo, que es lo que se compara contra un benchmark
    o contra otro inversor. Necesita al menos 2 fotos guardadas; None si no
    las hay. Para dólares, los aportes (en pesos transferidos) se
    convierten con la TRM histórica de cada fecha, igual que en
    _analisis_cambiario_dolares()."""
    df_snap = db.read_historial_valor_cartera()
    df_snap = df_snap[df_snap["Moneda"] == moneda] if not df_snap.empty else df_snap
    if df_snap.empty:
        return None
    tmp = df_snap.copy()
    tmp["Fecha"] = pd.to_datetime(tmp["Fecha"], dayfirst=True, errors="coerce")
    tmp = tmp.dropna(subset=["Fecha"]).sort_values("Fecha").drop_duplicates(subset="Fecha", keep="last")
    if len(tmp) < 2:
        return None

    df_ap = db.read_aportes_inversion(moneda)
    flujos = []
    if not df_ap.empty:
        serie_trm = None
        if moneda == "dolares":
            fechas_ap = pd.to_datetime(df_ap["Fecha"], dayfirst=True, errors="coerce").dropna()
            if not fechas_ap.empty:
                serie_trm = _descargar_historico_cierre("COP=X", fechas_ap.min().date())
        for _, fila in df_ap.iterrows():
            fecha = pd.to_datetime(fila["Fecha"], dayfirst=True, errors="coerce")
            if pd.isna(fecha):
                continue
            monto = _num_o_cero(fila["Monto Transferido (COP)"])
            if not monto:
                continue
            if moneda == "dolares":
                if serie_trm is None or serie_trm.empty:
                    continue
                trm_f = _trm_en(serie_trm, fecha.date())
                if not trm_f:
                    continue
                monto = monto / trm_f
            flujos.append((fecha, monto))

    fechas = tmp["Fecha"].tolist()
    valores = tmp["Valor Actual"].tolist()
    factor = 1.0
    n_sub = 0
    for i in range(1, len(fechas)):
        t0, t1 = fechas[i - 1], fechas[i]
        v0, v1 = valores[i - 1], valores[i]
        dias_sub = (t1 - t0).days
        if dias_sub <= 0:
            continue
        cfs = [(f, m) for f, m in flujos if t0 < f <= t1]
        suma_cf = sum(m for _, m in cfs)
        denom = v0 + sum(m * (t1 - f).days / dias_sub for f, m in cfs)
        if denom == 0:
            continue
        factor *= 1 + (v1 - v0 - suma_cf) / denom
        n_sub += 1
    if n_sub == 0:
        return None

    dias_totales = (fechas[-1] - fechas[0]).days
    twr_total = factor - 1
    twr_anual = (factor ** (365 / dias_totales) - 1) if dias_totales > 0 else None
    return dict(twr_total=twr_total, twr_anual=twr_anual, dias=dias_totales, n_subperiodos=n_sub,
                n_snapshots=len(tmp))


def _analizar_portafolio_moneda(moneda):
    """Separa las posiciones reales de la fila de margen/efectivo (la que
    no tiene 'Costo Total' propio, p. ej. 'IBKR - Efectivo/Margen') y arma
    tanto la vista bruta (ignora esa plata) como la de capital propio
    (la descuenta de ambos lados) — reusado por el Informe de Inversiones.
    None si la hoja todavía no tiene posiciones."""
    df_pos = db.read_posiciones_inversion(moneda)
    df_pos = df_pos[df_pos["Ticker / Fondo"].astype(str).str.strip() != ""].copy()
    if df_pos.empty:
        return None
    es_ajuste = df_pos["Costo Total"].fillna(0) == 0
    df_real = df_pos.loc[~es_ajuste].copy()
    ajuste = float(df_pos.loc[es_ajuste, "Valor Actual"].fillna(0).sum())
    costo = float(df_real["Costo Total"].fillna(0).sum())
    valor = float(df_real["Valor Actual"].fillna(0).sum())
    ganancia = valor - costo
    partes = df_real["Ticker / Fondo"].str.split(" - ", n=1)
    df_real["Plataforma"] = partes.str[0]
    df_real["Ticker"] = partes.str[-1]
    # El mismo ticker puede repetirse en más de una plataforma (p. ej. ADBE
    # en IBKR y en Hapi a la vez) — sin una etiqueta que los distinga, un
    # gráfico de barras por "Ticker" los apilaría en la misma fila en vez de
    # mostrarlos como posiciones separadas.
    df_real["Etiqueta"] = df_real["Ticker"] + " (" + df_real["Plataforma"] + ")"
    df_real["G/P %"] = df_real.apply(
        lambda r: (r["Ganancia/Pérdida"] / r["Costo Total"] * 100) if r["Costo Total"] else 0.0, axis=1)
    return dict(df=df_real, ajuste=ajuste, costo=costo, valor=valor, ganancia=ganancia,
                costo_propio=costo + ajuste, valor_propio=valor + ajuste)


def render_informe_inversiones():
    st.title("📈 Informe de Inversiones")
    st.caption("Resumen ejecutivo del portafolio en pesos y en dólares, por separado — posiciones, "
               "rentabilidad, composición y operaciones cerradas. Datos en vivo del Sheet.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Un informe de lectura corrida en vez de las herramientas interactivas de 📈 Inversiones. Para "
            "cada moneda separa las posiciones reales de cualquier fila de margen/efectivo sin costo propio "
            "(p. ej. un saldo negativo de margen en el bróker), y muestra tres cifras de rentabilidad "
            "distintas a propósito: el retorno bruto de las posiciones (ignora el margen — infla o subestima "
            "según el signo), el retorno sobre tu **capital propio** (descuenta el margen de ambos lados: lo "
            "que pusiste y lo que tenés hoy) y el **XIRR** — la rentabilidad anualizada real, ponderada por "
            "la fecha de cada aporte, que ya calcula 📈 Inversiones → Crecimiento y Rentabilidad. Si el "
            "capital propio y el XIRR no coinciden, es porque el XIRR pondera por tiempo (un aporte de hace "
            "un año no es lo mismo que uno de la semana pasada) y el capital propio es una foto de hoy.")

    for moneda, simbolo, nombre in (("pesos", "$", "Portafolio en pesos"),
                                     ("dolares", "US$", "Portafolio en dólares")):
        st.markdown(f"### {nombre}")
        info = _analizar_portafolio_moneda(moneda)
        if info is None:
            st.info("Todavía no hay posiciones cargadas acá.")
            continue

        def _fmt(v):
            return f"{simbolo}{v:,.2f}" if moneda == "dolares" else fmt_moneda(v)

        retorno_bruto = (info["ganancia"] / info["costo"] * 100) if info["costo"] else 0.0
        # Si el margen prestado supera el costo de las posiciones, "capital
        # propio" da 0 o negativo y dividir ahí ya no da un % que se pueda
        # leer como retorno (invertiría el signo o se dispararía) — se
        # avisa en vez de mostrar una cifra falsa.
        capital_propio_valido = info["costo_propio"] > 0
        retorno_propio = (info["ganancia"] / info["costo_propio"] * 100) if capital_propio_valido else None
        xirr = _rentabilidad_xirr(moneda)

        c1, c2 = st.columns(2)
        c1.metric("Valor de las posiciones", _fmt(info["valor"]))
        c2.metric("Ganancia/pérdida no realizada", _fmt(info["ganancia"]), f"{retorno_bruto:+.1f}% bruto")
        c3, c4 = st.columns(2)
        if capital_propio_valido:
            c3.metric("Capital propio hoy", _fmt(info["valor_propio"]),
                     f"{retorno_propio:+.1f}% sobre capital propio")
        else:
            c3.metric("Capital propio hoy", _fmt(info["valor_propio"]),
                     "apalancado más de 1:1 — % no es legible", delta_color="off")
        c4.metric("Rentabilidad anualizada (XIRR)", f"{xirr:+.1%}" if xirr is not None else "—")

        if abs(info["ajuste"]) > 1:
            motivo = "margen prestado por el bróker" if info["ajuste"] < 0 else "efectivo sin invertir"
            st.caption(f"Hay {_fmt(abs(info['ajuste']))} de {motivo} mezclados en el valor de la cuenta, fuera "
                       "de las posiciones. El retorno bruto de arriba lo ignora — compará mejor contra "
                       "'Capital propio', que sí lo descuenta de los dos lados (lo que pusiste y lo que tenés "
                       "hoy), o contra el XIRR, que además pondera por fecha.")

        st.markdown("**📐 Métricas de rendimiento**")
        twr = _twr_moneda(moneda)
        filas_metricas = [("Retorno simple (bruto)", f"{retorno_bruto:+.1f}%",
                           "Ganancia / costo de las posiciones — sin apalancamiento, sin importar fechas.")]
        if capital_propio_valido:
            filas_metricas.append(("Retorno sobre capital propio", f"{retorno_propio:+.1f}%",
                                   "Ídem, descontando el margen prestado — tu retorno real sobre tu plata."))
        filas_metricas.append(("XIRR (anualizado)", f"{xirr:+.1%}" if xirr is not None else "—",
                               "Money-weighted: pondera CUÁNDO metiste cada peso, no solo cuánto."))
        if twr:
            filas_metricas.append((
                "TWR (anualizado)", f"{twr['twr_anual']:+.1%}" if twr["twr_anual"] is not None else "—",
                f"Time-weighted: encadena {twr['n_subperiodos']} sub-período(s) entre fotos guardadas, "
                "ignorando cuándo aportaste — mide qué tan bien elegiste, no cuándo invertiste."))
            filas_metricas.append(("TWR del período (sin anualizar)", f"{twr['twr_total']:+.1%}",
                                   f"Retorno acumulado en los últimos {twr['dias']} días entre fotos."))
        else:
            filas_metricas.append(("TWR", "—",
                                   "Necesita al menos 2 fotos en 'Historial de Valor de Cartera' — se guardan "
                                   "solas cada vez que actualizás precios o posiciones/aportes."))
        st.dataframe(pd.DataFrame(filas_metricas, columns=["Métrica", "Valor", "Qué mide"]),
                    hide_index=True, use_container_width=True)
        if xirr is not None and twr and twr["twr_anual"] is not None and abs(xirr - twr["twr_anual"]) > 0.05:
            mejor_cuando = "elegiste mejor de lo que sugiere el timing de tus aportes" if twr["twr_anual"] > xirr \
                else "el timing de tus aportes te ayudó más de lo que sugiere la calidad de tus elecciones"
            st.caption(f"XIRR y TWR difieren bastante acá — probablemente {mejor_cuando}.")

        df_real = info["df"]
        mejor = df_real.loc[df_real["G/P %"].idxmax()]
        peor = df_real.loc[df_real["G/P %"].idxmin()]
        # Si hay una sola plataforma (típico en pesos), repetirla entre
        # paréntesis solo alarga el texto sin agregar información — y en
        # columnas angostas eso es justo lo que provoca que se corte a la
        # mitad de una palabra.
        multi_plataforma = df_real["Plataforma"].nunique() > 1

        def _etiqueta_posicion(fila):
            return f"{fila['Ticker']} ({fila['Plataforma']})" if multi_plataforma else fila["Ticker"]

        # Un ticker solo (sin espacios, p. ej. "PFGRUPSURA") puede no entrar
        # en 1/3 de fila y romperse a la mitad de la palabra — dos columnas
        # les da el doble de ancho; "Posiciones activas" (un número corto)
        # nunca tiene ese problema, así que va aparte.
        cm1, cm2 = st.columns(2)
        cm1.metric("Mejor posición", _etiqueta_posicion(mejor), f"{mejor['G/P %']:+.1f}%")
        cm2.metric("Peor posición", _etiqueta_posicion(peor), f"{peor['G/P %']:+.1f}%")
        st.caption(f"{len(df_real)} posiciones activas.")

        st.markdown("**Composición por posición (valor actual)**")
        df_comp = df_real.sort_values("Valor Actual", ascending=True)
        fig_comp = px.bar(df_comp, x="Valor Actual", y="Etiqueta", color="Plataforma", orientation="h")
        fig_comp.update_layout(margin=dict(l=0, r=0, t=10, b=0),
                               height=max(220, 28 * len(df_comp)), yaxis_title=None)
        st.plotly_chart(fig_comp, use_container_width=True, key=f"informe_inv_comp_{moneda}")

        st.markdown("**Ganancia/pérdida por posición**")
        df_gp = df_real.sort_values("G/P %", ascending=True)
        fig_gp = px.bar(df_gp, x="G/P %", y="Etiqueta",
                        color=(df_gp["G/P %"] >= 0).map({True: "Ganancia", False: "Pérdida"}),
                        color_discrete_map={"Ganancia": "#0ca30c", "Pérdida": "#d03b3b"}, orientation="h")
        fig_gp.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=max(220, 28 * len(df_gp)),
                             yaxis_title=None, legend_title=None)
        st.plotly_chart(fig_gp, use_container_width=True, key=f"informe_inv_gp_{moneda}")

        with st.expander(f"Ver las {len(df_real)} posiciones en detalle"):
            df_tabla = df_real[["Ticker", "Plataforma", "Tipo", "Cantidad", "Precio Compra Promedio",
                                "Costo Total", "Precio Actual", "Valor Actual", "Ganancia/Pérdida",
                                "G/P %"]].sort_values("Valor Actual", ascending=False)
            st.dataframe(df_tabla, hide_index=True, use_container_width=True)

        if moneda == "dolares":
            st.markdown("**💱 Efecto cambiario de los aportes (TRM)**")
            st.caption("Cada peso que mandaste a IBKR/Binance/Hapi se convirtió a dólares a la TRM de ese día. "
                      "Si el dólar cayó desde entonces (menos pesos por dólar), esos mismos dólares valen "
                      "menos pesos hoy que lo que costó comprarlos — un efecto aparte del rendimiento de las "
                      "posiciones en sí, que ya se mide en dólares arriba (retorno bruto, capital propio, "
                      "XIRR). Este cálculo es la TRM histórica día a día de Yahoo Finance, no la tasa exacta "
                      "de cada transferencia real.")
            fx = _analisis_cambiario_dolares()
            if fx is None:
                st.info("No pude calcular el efecto cambiario — puede ser que no haya aportes cargados acá "
                        "todavía, o que no se pudo descargar la TRM histórica de Yahoo Finance ahora mismo.")
            else:
                fx1, fx2, fx3 = st.columns(3)
                fx1.metric("Aportado (histórico)", fmt_moneda(fx["total_cop"]))
                fx2.metric("Equivalente en dólares al aportar", f"US${fx['total_usd']:,.2f}")
                fx3.metric("TRM promedio ponderado al aportar",
                          f"${fx['trm_promedio']:,.0f}" if fx["trm_promedio"] else "—")
                fx4, fx5, fx6 = st.columns(3)
                fx4.metric("TRM de hoy", f"${fx['trm_hoy']:,.0f}")
                fx5.metric("Esos mismos dólares valen hoy", fmt_moneda(fx["total_hoy"]))
                pct_fx = (fx["diferencia"] / fx["total_cop"] * 100) if fx["total_cop"] else 0.0
                fx6.metric("Efecto cambiario", fmt_moneda(fx["diferencia"]), f"{pct_fx:+.1f}%")

                if fx["diferencia"] < 0:
                    st.warning(f"⚠️ El dólar cayó frente al peso desde que hiciste estos aportes: en pesos de "
                              f"hoy, perdiste {fmt_moneda(abs(fx['diferencia']))} ({abs(pct_fx):.1f}%) solo "
                              "por el tipo de cambio, sin contar cómo les fue a las posiciones en sí.")
                else:
                    st.success(f"✅ El dólar subió frente al peso desde que hiciste estos aportes: en pesos "
                              f"de hoy, ganaste {fmt_moneda(fx['diferencia'])} ({pct_fx:.1f}%) solo por el "
                              "tipo de cambio, sin contar cómo les fue a las posiciones en sí.")

                if fx["trm_promedio"]:
                    st.markdown("**Retorno combinado en pesos (inversión ponderada por el tipo de cambio)**")
                    r_usd_pct = retorno_propio if capital_propio_valido else retorno_bruto
                    r_fx = fx["trm_hoy"] / fx["trm_promedio"] - 1
                    r_total = (1 + r_usd_pct / 100) * (1 + r_fx) - 1
                    rc1, rc2, rc3 = st.columns(3)
                    rc1.metric("Rendimiento en dólares", f"{r_usd_pct:+.1f}%")
                    rc2.metric("Efecto cambiario", f"{r_fx * 100:+.1f}%")
                    rc3.metric("Total combinado en pesos", f"{r_total * 100:+.1f}%")
                    st.caption("Los dos efectos se combinan multiplicando, no sumando — "
                              "(1 + rendimiento en dólares) × (1 + efecto cambiario) − 1 — porque el segundo "
                              "se aplica sobre el resultado del primero. Esta es la cifra que de verdad "
                              "importa si algún día convertís estos dólares de vuelta a pesos: no alcanza con "
                              "que la posición haya rendido bien en dólares si el dólar cayó lo suficiente "
                              "como para comerse esa ganancia (o al revés).")

                with st.expander("Ver el detalle por aporte"):
                    df_fx_tabla = fx["df"].copy()
                    df_fx_tabla["Fecha"] = df_fx_tabla["Fecha"].astype(str)
                    df_fx_tabla["Monto_COP"] = df_fx_tabla["Monto_COP"].map(fmt_moneda)
                    df_fx_tabla["TRM_Fecha"] = df_fx_tabla["TRM_Fecha"].map(lambda v: f"${v:,.0f}")
                    df_fx_tabla["USD_Equivalente"] = df_fx_tabla["USD_Equivalente"].map(lambda v: f"US${v:,.2f}")
                    df_fx_tabla["TRM_Hoy"] = df_fx_tabla["TRM_Hoy"].map(lambda v: f"${v:,.0f}")
                    df_fx_tabla["Valor_Hoy_COP"] = df_fx_tabla["Valor_Hoy_COP"].map(fmt_moneda)
                    df_fx_tabla["Diferencia_Cambiaria"] = df_fx_tabla["Diferencia_Cambiaria"].map(fmt_moneda)
                    df_fx_tabla = df_fx_tabla.rename(columns={
                        "Monto_COP": "Monto (COP)", "TRM_Fecha": "TRM ese día",
                        "USD_Equivalente": "USD equivalente", "TRM_Hoy": "TRM hoy",
                        "Valor_Hoy_COP": "Valor hoy (COP)", "Diferencia_Cambiaria": "Diferencia cambiaria"})
                    st.dataframe(df_fx_tabla, hide_index=True, use_container_width=True)

        st.divider()

    st.markdown("### Capital aportado")
    st.caption("Lo que efectivamente salió de tu cuenta hacia cada plataforma (siempre en pesos, sea cual "
               "sea la moneda de destino) — no incluye el margen prestado, ese nunca vino de tu banco.")
    partes_aportes = []
    for moneda in ("pesos", "dolares"):
        df_ap = db.read_aportes_inversion(moneda)
        if not df_ap.empty:
            partes_aportes.append(df_ap.assign(Moneda="Pesos" if moneda == "pesos" else "Dólares"))
    if partes_aportes:
        df_ap_todo = pd.concat(partes_aportes, ignore_index=True)
        por_plataforma = (df_ap_todo.groupby("Plataforma")["Monto Transferido (COP)"].sum()
                          .sort_values(ascending=False))
        total_aportado = por_plataforma.sum()
        ca1, ca2 = st.columns(2)
        ca1.metric("Total transferido (neto)", fmt_moneda(total_aportado))
        ca2.metric("Plataformas activas", f"{len(por_plataforma)}")
        fig_ap = px.bar(por_plataforma.reset_index(), x="Monto Transferido (COP)", y="Plataforma",
                        orientation="h")
        fig_ap.update_layout(margin=dict(l=0, r=0, t=10, b=0),
                             height=max(180, 40 * len(por_plataforma)), yaxis_title=None)
        st.plotly_chart(fig_ap, use_container_width=True, key="informe_inv_aportes")
    else:
        st.info("Todavía no hay aportes registrados.")

    st.markdown("### Operaciones cerradas")
    df_hist = db.read_historial_inversion()
    if not df_hist.empty:
        df_cierres = df_hist[df_hist["Operación"].isin(["SELL", "COVER"])].copy()
        if not df_cierres.empty:
            ganadoras = (df_cierres["Resultado Realizado"] > 0).sum()
            total_cierres = len(df_cierres)
            tasa_acierto = ganadoras / total_cierres * 100
            realizado_cop = df_cierres.loc[df_cierres["Moneda"] == "COP", "Resultado Realizado"].sum()
            realizado_usd = df_cierres.loc[df_cierres["Moneda"] == "USD", "Resultado Realizado"].sum()
            # Dos columnas de moneda porque st.metric interpreta un par de "$"
            # en el mismo valor como delimitador de LaTeX (p. ej. combinar
            # "$644.859" y "US$266,29" en un solo string) y termina
            # renderizando texto en cursiva en vez del signo de pesos.
            co1, co2 = st.columns(2)
            co1.metric("Operaciones cerradas", f"{total_cierres}")
            co2.metric("Tasa de acierto", f"{tasa_acierto:.0f}% ({ganadoras}/{total_cierres})")
            co3, co4 = st.columns(2)
            co3.metric("Resultado realizado (COP)", fmt_moneda(realizado_cop))
            co4.metric("Resultado realizado (USD)", f"US${realizado_usd:,.2f}")
            st.caption("Ya queda reflejado en el costo promedio de lo que sigue abierto arriba — no lo sumes "
                       "aparte a la ganancia no realizada.")
        else:
            st.info("Todavía no hay operaciones cerradas (solo compras abiertas).")
    else:
        st.info("Todavía no hay historial de operaciones importado.")


def render_informe_presupuesto():
    st.title("📊 Informe de Presupuesto, Ingresos y Gastos")
    st.caption("Resumen ejecutivo de tu mes/año: cuánto entró, cuánto salió, tasa de ahorro, gasto esencial "
               "vs. no esencial y cómo vas contra el presupuesto. Vista devengado (cada movimiento en su "
               "propia fecha), salvo la evolución mensual, que es efectivo real — igual que 🏠 Resumen.")
    with st.expander("ℹ️ ¿Qué hace esta sección?"):
        st.write(
            "Un informe de lectura corrida que junta en un solo lugar lo que hoy está repartido entre 🏠 "
            "Resumen, 📊 Análisis y 📋 Presupuesto: ingresos y gastos del alcance elegido, tasa de ahorro, "
            "gasto por categoría, esencial vs. no esencial, evolución de los últimos 12 meses y presupuesto "
            "vs. real del mes actualmente seleccionado en 📋 Presupuesto.")

    alcance = st.radio("Alcance del informe:", ["Total histórico", "Un año", "Un mes"], horizontal=True,
                       key="alcance_informe_presu")
    anio_sel = mes_sel = None
    hoy = date.today()
    anios_disp = list(range(2023, 2033))
    if alcance == "Un año":
        anio_sel = st.selectbox("Año", anios_disp,
                                index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0,
                                key="anio_informe_presu")
    elif alcance == "Un mes":
        c_a, c_m = st.columns(2)
        anio_sel = c_a.selectbox("Año", anios_disp,
                                 index=anios_disp.index(hoy.year) if hoy.year in anios_disp else 0,
                                 key="anio_informe_presu_mes")
        mes_sel = c_m.selectbox("Mes", list(range(1, 13)), index=hoy.month - 1,
                                format_func=lambda m: MESES_NOMBRE[m], key="mes_informe_presu_mes")

    def coincide(anio, mes):
        if alcance == "Total histórico":
            return True
        if anio is None:
            return False
        if alcance == "Un año":
            return anio == anio_sel
        return anio == anio_sel and mes == mes_sel

    datos = _ingresos_gastos_periodo(coincide)
    total_ingresos = datos["total_ingresos"]
    # El ahorro se discrimina por origen en vez de tratarlo como un solo
    # bulto: la categoría de gasto Inversiones, la categoría de gasto
    # Ahorro, y el ahorro vía nómina (descuento "Ahorro" + aporte al Fondo
    # de Empleados — no la cuota de su préstamo, esa es deuda). Sumar todo
    # de vuelta al balance es el mismo criterio que ya usa "Utilidad Neta"
    # en Estado de Resultados: tratar el ahorro como si fuera consumo (lo
    # que hacía esta cuenta antes) subestima cuánto ahorraste de verdad.
    gasto_inversiones = datos["gasto_inversiones"]
    ahorro_gasto = datos["gasto_ahorro"]
    ahorro_nomina = datos["descuento_ahorro"] + datos["descuento_fondo_empleados"]
    ahorro_total = gasto_inversiones + ahorro_gasto + ahorro_nomina
    deuda_nomina = datos["descuento_deuda_nomina"]

    gasto_consumo = datos["gasto_operativo"] + datos["descuentos_nomina"] - ahorro_nomina - deuda_nomina
    total_egresos = gasto_consumo + ahorro_total + deuda_nomina  # = gasto_real + descuentos_nomina
    disponible_efectivo = total_ingresos - total_egresos
    total_ahorrado = disponible_efectivo + ahorro_total
    tasa_ahorro = (total_ahorrado / total_ingresos * 100) if total_ingresos else 0.0

    # Streamlit renderiza markdown, y dos o más "$" en el mismo texto se
    # interpretan como un bloque de LaTeX (KaTeX) en vez de signos de pesos
    # literales — hay que escaparlos acá porque el párrafo combina varios
    # fmt_moneda() en una sola oración (a diferencia del resto de la app,
    # que nunca junta más de un monto por mensaje).
    def _money_md(v):
        return fmt_moneda(v).replace("$", r"\$")

    st.markdown(f"Con el alcance elegido entraron **{_money_md(total_ingresos)}**. De ahí, "
               f"**{_money_md(gasto_consumo)}** se consumió y **{_money_md(ahorro_total)}** se fue a ahorro e "
               f"inversión (no es gasto, sigue siendo tuyo) — entre lo que quedó líquido y lo ahorrado/"
               f"invertido, **no gastaste el {tasa_ahorro:.0f}%** de lo que ganaste.")

    df_deudas_inf = db.read_deudas_resumen()
    c1, c2 = st.columns(2)
    c1.metric("Ingresos totales", fmt_moneda(total_ingresos))
    c2.metric("Gasto de consumo", fmt_moneda(gasto_consumo))
    c3, c4 = st.columns(2)
    c3.metric("Disponible en efectivo", fmt_moneda(disponible_efectivo))
    c4.metric("Tasa de ahorro (efectivo + ahorro + inversión)", f"{tasa_ahorro:.0f}%")

    st.markdown("###### Ahorro e inversión, discriminado por origen")
    ca1, ca2 = st.columns(2)
    ca1.metric("Inversiones", fmt_moneda(gasto_inversiones))
    ca2.metric("Ahorro (efectivo/tarjeta)", fmt_moneda(ahorro_gasto))
    ca3, ca4 = st.columns(2)
    ca3.metric("Ahorro vía nómina", fmt_moneda(ahorro_nomina))
    ca4.metric("Deuda pendiente hoy", fmt_moneda(df_deudas_inf["Saldo Actual"].sum()) if not df_deudas_inf.empty
              else fmt_moneda(0))
    st.caption("'Tasa de ahorro' = (efectivo disponible + Inversiones + Ahorro, sea de tarjeta/efectivo o vía "
              "nómina) / ingresos. Si solo te interesa el efectivo disponible en la cuenta corriente, esa es "
              "'Disponible en efectivo' arriba.")
    _render_desglose_inversiones(datos)

    if deuda_nomina > 0:
        st.caption(f"Aparte, {fmt_moneda(deuda_nomina)} de este alcance se fue a cuotas de deuda descontadas "
                  "directo de la nómina (leasing habitacional / préstamo del Fondo de Empleados) — tampoco es "
                  "consumo, pero no se suma a la tasa de ahorro porque cada cuota mezcla capital e interés y "
                  "no hay forma de separarlos acá (el detalle de cada crédito está en 🏦 Deudas).")
    if datos["ingreso_cesantias"] > 0:
        st.caption(f"De los ingresos, {fmt_moneda(datos['ingreso_cesantias'])} son de la categoría Cesantías "
                  "(intereses o retiro del fondo) — ya están incluidos en 'Ingresos totales', discriminados "
                  "acá para que sepas cuánto de lo que entró no fue tu sueldo habitual.")

    st.divider()
    st.markdown("#### Ingresos")
    ci1, ci2 = st.columns(2)
    ci1.metric("Colillas de pago", fmt_moneda(datos["ingresos_colillas"]))
    ci2.metric("Otros ingresos", fmt_moneda(datos["otros_ingresos"]))
    if datos["ingresos_por_categoria"]:
        df_ing = pd.DataFrame(sorted(datos["ingresos_por_categoria"].items(), key=lambda x: -x[1]),
                              columns=["Categoría", "Valor"])
        fig_ing = px.bar(df_ing, x="Valor", y="Categoría", orientation="h")
        fig_ing.update_layout(yaxis=dict(autorange="reversed"), margin=dict(l=0, r=0, t=10, b=0),
                              height=max(180, 32 * len(df_ing)), yaxis_title=None)
        st.plotly_chart(fig_ing, use_container_width=True, key="informe_presu_ingresos")

    st.markdown("#### Gastos")
    cg1, cg2, cg3 = st.columns(3)
    cg1.metric("Gasto de consumo (sin inversiones ni ahorro)", fmt_moneda(datos["gasto_operativo"]))
    cg2.metric("Descuentos de nómina", fmt_moneda(datos["descuentos_nomina"]))
    cg3.metric("Aportes a inversión", fmt_moneda(datos["gasto_inversiones"]))
    cat_gasto = {k: v for k, v in datos["gasto_por_categoria"].items() if k not in ("Inversiones", "Ahorro")}
    if cat_gasto:
        df_gc = pd.DataFrame(sorted(cat_gasto.items(), key=lambda x: -x[1]), columns=["Categoría", "Valor"])
        fig_gc = px.bar(df_gc, x="Valor", y="Categoría", orientation="h")
        fig_gc.update_layout(yaxis=dict(autorange="reversed"), margin=dict(l=0, r=0, t=10, b=0),
                             height=max(220, 30 * len(df_gc)), yaxis_title=None)
        st.plotly_chart(fig_gc, use_container_width=True, key="informe_presu_gastos")
    if datos["gasto_sin_categorizar"] > 0:
        pct_otros = (datos["gasto_sin_categorizar"] / datos["gasto_real"] * 100) if datos["gasto_real"] else 0
        st.warning(f"⚠️ {fmt_moneda(datos['gasto_sin_categorizar'])} ({pct_otros:.0f}% del gasto) está en la "
                  "categoría genérica 'Otros' — revisalo en 📊 Análisis → Movimientos para categorizarlo mejor.")

    st.markdown("#### Esencial vs. no esencial")
    clasificacion_sheet = db.read_clasificacion_esencial()

    def _clasificar(categoria):
        return clasificacion_sheet.get(categoria) or clasificar_esencial(categoria)

    df_mov_inf = db.read_movimientos_recientes()
    df_gasto_inf = pd.DataFrame()
    if not df_mov_inf.empty:
        df_gasto_inf = df_mov_inf[
            (df_mov_inf["Valor"] < 0) & (df_mov_inf["Moneda"] == "COP")
            & (df_mov_inf["Presupuestar"] == "Sí") & (df_mov_inf["Fuente"] != "Colilla (descuento)")
        ].copy()
        df_gasto_inf["Valor"] = df_gasto_inf["Valor"].abs()
        am_inf = [_extraer_anio_mes(p) for p in df_gasto_inf["Periodo"]]
        mask_inf = [coincide(a, m) for a, m in am_inf]
        df_gasto_inf = df_gasto_inf.loc[mask_inf].copy()
        df_gasto_inf["Clasificación"] = df_gasto_inf["Categoría"].map(_clasificar)
        df_gasto_inf = df_gasto_inf[df_gasto_inf["Clasificación"] != "No consumo (ahorro/inversión)"]
    if not df_gasto_inf.empty:
        resumen_clas_inf = df_gasto_inf.groupby("Clasificación")["Valor"].sum()
        total_clas = resumen_clas_inf.sum()
        ce1, ce2, ce3 = st.columns(3)
        for col, clasif in zip((ce1, ce2, ce3), ("Esencial", "No esencial", "Sin clasificar")):
            v = resumen_clas_inf.get(clasif, 0.0)
            pct = (v / total_clas * 100) if total_clas else 0
            col.metric(clasif, fmt_moneda(v), f"{pct:.0f}% del gasto", delta_color="off")
        sin_clasificar_cats = sorted(df_gasto_inf.loc[df_gasto_inf["Clasificación"] == "Sin clasificar",
                                                       "Categoría"].unique())
        if sin_clasificar_cats:
            st.caption("Sin clasificar todavía: " + ", ".join(sin_clasificar_cats)
                      + " — editalo en la hoja 'Categorías Esenciales'.")
    else:
        st.caption("No hay gasto de consumo clasificable en este alcance.")

    st.divider()
    st.markdown("#### Evolución mensual (efectivo real, últimos 12 meses)")
    df_mes_inf = db.read_resumen_mensual()
    if not df_mes_inf.empty and "Ingresos ganados" in df_mes_inf.columns:
        df_mes_inf = df_mes_inf[df_mes_inf["Ingresos ganados"] != 0].tail(12).copy()
    if not df_mes_inf.empty and "Ingresos ganados" in df_mes_inf.columns:
        df_long_inf = df_mes_inf.melt(id_vars=["Mes"],
                                      value_vars=["Ingresos ganados", "Gastos personales", "Disponible del mes"],
                                      var_name="Concepto", value_name="Valor")
        fig_evo = px.bar(df_long_inf, x="Mes", y="Valor", color="Concepto", barmode="group")
        fig_evo.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=360)
        _fijar_eje_periodos(fig_evo, df_mes_inf["Mes"])
        st.plotly_chart(fig_evo, use_container_width=True, key="informe_presu_evolucion")

        df_mes_inf["Tasa de ahorro %"] = df_mes_inf.apply(
            lambda r: (r["Disponible del mes"] / r["Ingresos ganados"] * 100) if r["Ingresos ganados"] else 0.0,
            axis=1)
        fig_tasa = px.line(df_mes_inf, x="Mes", y="Tasa de ahorro %", markers=True)
        fig_tasa.update_layout(margin=dict(l=0, r=0, t=10, b=0), height=260)
        _fijar_eje_periodos(fig_tasa, df_mes_inf["Mes"])
        st.plotly_chart(fig_tasa, use_container_width=True, key="informe_presu_tasa_ahorro")
    else:
        st.info("Todavía no hay suficientes meses cargados para la evolución mensual.")

    st.divider()
    st.markdown("#### Presupuesto vs. real")
    p_inf = db.read_presupuesto()
    cats_con_meta = [c for c in p_inf["categorias"] if c["presupuesto"] > 0]
    if cats_con_meta:
        total_meta = sum(c["presupuesto"] for c in cats_con_meta)
        total_real_meta = sum(c["gasto_real"] for c in cats_con_meta)
        pct_usado = (total_real_meta / total_meta * 100) if total_meta else 0
        cp1, cp2, cp3 = st.columns(3)
        cp1.metric("Presupuestado (categorías con meta)", fmt_moneda(total_meta))
        cp2.metric("Gasto real de esas categorías", fmt_moneda(total_real_meta))
        cp3.metric("% del presupuesto usado", f"{pct_usado:.0f}%")
        pasadas = [c for c in cats_con_meta if c["gasto_real"] > c["presupuesto"]]
        if pasadas:
            # Ídem nota de más arriba: dos o más "$" en el mismo mensaje se
            # leerían como LaTeX, y acá se junta un monto por cada categoría
            # pasada de presupuesto.
            textos_pasadas = [f"{c['categoria']} ({_money_md(c['gasto_real'] - c['presupuesto'])} de más)"
                              for c in pasadas]
            st.warning("Te pasaste del presupuesto en: " + ", ".join(textos_pasadas))
        else:
            st.success("No te pasaste del presupuesto en ninguna categoría con meta definida.")
        st.caption("Corresponde al mes actualmente seleccionado en 📋 Presupuesto — cambialo ahí si querés "
                  "ver otro mes (no depende del alcance elegido arriba).")
    else:
        st.info("Todavía no definiste metas de presupuesto por categoría — hacelo en 📋 Presupuesto.")


# ===========================================================================
# UI — navegación por barra lateral
# ===========================================================================
with st.sidebar:
    st.markdown("## 💰 Presupuesto")
    st.caption("Los datos viven en Google Sheets — todo lo que aplicás en la app queda guardado ahí mismo.")
    st.divider()
    pagina = st.radio("Navegación", [
        "🏠 Resumen",
        "📊 Análisis",
        "📋 Presupuesto",
        "🏦 Deudas",
        "📈 Inversiones",
        "🏢 Estados Financieros",
        "✅ Verificar Datos",
        "🔍 Salud de los Datos",
        "💰 Ingresos",
        "💳 Egresos",
        "🧾 Facturación Electrónica",
        "📑 Declaraciones de Renta",
        "―――――――――",
        "📈 Informe de Inversiones",
        "📊 Informe de Presupuesto, Ingresos y Gastos",
        "――――――――― ",
        "📄 Cargar Colillas de Pago",
        "💳 Cargar Extractos de Tarjeta",
        "🏦 Cargar Cuenta de Ahorros",
    ], label_visibility="collapsed")
    st.divider()
    if st.button("🔄 Refrescar datos", use_container_width=True):
        db.clear_read_cache()
        st.rerun()
    if st.button("⬇️ Preparar descarga del Sheet (Excel)", use_container_width=True):
        with st.spinner("Armando el archivo..."):
            try:
                st.session_state["_excel_export"] = db.export_workbook_excel()
            except Exception as e:
                st.session_state.pop("_excel_export", None)
                st.error(f"No pude preparar la descarga: {e}")
    if "_excel_export" in st.session_state:
        st.download_button(
            "💾 Guardar Excel", data=st.session_state["_excel_export"],
            file_name=f"presupuesto_{date.today().isoformat()}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True)

PAGINAS = {
    "🏠 Resumen": render_resumen,
    "📊 Análisis": render_analisis,
    "📋 Presupuesto": render_presupuesto,
    "🏦 Deudas": render_deudas,
    "📈 Inversiones": render_inversiones,
    "🏢 Estados Financieros": render_estados_financieros,
    "✅ Verificar Datos": render_verificar_datos,
    "🔍 Salud de los Datos": render_salud_datos,
    "💰 Ingresos": render_ingresos,
    "💳 Egresos": render_egresos,
    "🧾 Facturación Electrónica": render_facturacion_electronica,
    "📑 Declaraciones de Renta": render_declaraciones_renta,
    "📈 Informe de Inversiones": render_informe_inversiones,
    "📊 Informe de Presupuesto, Ingresos y Gastos": render_informe_presupuesto,
    "📄 Cargar Colillas de Pago": render_cargar_colillas,
    "💳 Cargar Extractos de Tarjeta": render_cargar_extractos,
    "🏦 Cargar Cuenta de Ahorros": render_cargar_cuenta,
}
render_fn = PAGINAS.get(pagina)
if render_fn:
    render_fn()
